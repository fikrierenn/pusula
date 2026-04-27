# -*- coding: utf-8 -*-
"""
3Al2Öde Performans Raporu → Excel (GUI)
Gerekli: pip install openpyxl pyodbc
"""
import pyodbc, os, subprocess, threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# --- BAĞLANTI ---
def get_conn_str():
    """Makinede kurulu SQL Server ODBC sürücüsünü otomatik bul."""
    preferred = [
        "ODBC Driver 18 for SQL Server",
        "ODBC Driver 17 for SQL Server",
        "SQL Server Native Client 11.0",
        "SQL Server",
    ]
    installed = [d for d in pyodbc.drivers() if "sql" in d.lower()]
    driver = None
    for p in preferred:
        if p in installed:
            driver = p
            break
    if not driver and installed:
        driver = installed[0]
    if not driver:
        raise RuntimeError("SQL Server ODBC surucusu bulunamadi!\nKurulu suruculer: " + str(pyodbc.drivers()))
    return (
        f"DRIVER={{{driver}}};"
        "SERVER=192.168.40.201;"
        "DATABASE=EncoreMerkez;"
        "Trusted_Connection=yes;"
        "TrustServerCertificate=yes;"
    )

# --- ORTAK BASE CTE ---
BASE_CTE = """
;WITH base AS (
    SELECT
        CAST(s.Date AS date) AS Gun,
        st.Name AS Magaza,
        s.LineCount,
        s.GrossTotal,
        s.DiscountTotal,
        CASE
            WHEN fkt.Has3Al2Ode = 1 THEN '3Al2Ode'
            WHEN fkt.SalesId IS NOT NULL THEN 'Diger'
            ELSE 'Kampanyasiz'
        END AS KTip
    FROM dbo.Sales s
    INNER JOIN dbo.Stores st ON s.StoresId = st.Id
    LEFT JOIN (
        SELECT SalesId,
            MAX(CASE WHEN CampaignId IN (1,12) THEN 1 ELSE 0 END) AS Has3Al2Ode
        FROM dbo.SalesProductCampaigns
        GROUP BY SalesId
    ) fkt ON fkt.SalesId = s.Id
    WHERE s.DocumentsTypeId IN (1, 2)
"""

# --- METRIK BLOĞU ---
METRIC_BLOCK = """
    SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END) AS [3Al2Ode Fis],
    SUM(CASE WHEN KTip='3Al2Ode' THEN LineCount ELSE 0 END) AS [3Al2Ode Urun],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal ELSE 0 END) AS decimal(18,2)) AS [3Al2Ode Brut],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN DiscountTotal ELSE 0 END) AS decimal(18,2)) AS [3Al2Ode Indirim],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal-DiscountTotal ELSE 0 END) AS decimal(18,2)) AS [3Al2Ode Net],
    CAST(CASE WHEN SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='3Al2Ode' THEN LineCount ELSE 0 END)*1.0/SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(5,2)) AS [3Al2Ode Ort Sepet Adet],
    CAST(CASE WHEN SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal-DiscountTotal ELSE 0 END)*1.0/SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(10,2)) AS [3Al2Ode Ort Sepet Tutar],
    CAST(CASE WHEN SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='3Al2Ode' THEN DiscountTotal ELSE 0 END)*100.0/SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal ELSE 0 END)
        ELSE 0 END AS decimal(5,1)) AS [3Al2Ode Indirim Orani],
    SUM(CASE WHEN KTip='Diger' THEN 1 ELSE 0 END) AS [DigerKmp Fis],
    SUM(CASE WHEN KTip='Diger' THEN LineCount ELSE 0 END) AS [DigerKmp Urun],
    CAST(SUM(CASE WHEN KTip='Diger' THEN GrossTotal ELSE 0 END) AS decimal(18,2)) AS [DigerKmp Brut],
    CAST(SUM(CASE WHEN KTip='Diger' THEN DiscountTotal ELSE 0 END) AS decimal(18,2)) AS [DigerKmp Indirim],
    CAST(SUM(CASE WHEN KTip='Diger' THEN GrossTotal-DiscountTotal ELSE 0 END) AS decimal(18,2)) AS [DigerKmp Net],
    CAST(CASE WHEN SUM(CASE WHEN KTip='Diger' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='Diger' THEN LineCount ELSE 0 END)*1.0/SUM(CASE WHEN KTip='Diger' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(5,2)) AS [DigerKmp Ort Sepet Adet],
    CAST(CASE WHEN SUM(CASE WHEN KTip='Diger' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='Diger' THEN GrossTotal-DiscountTotal ELSE 0 END)*1.0/SUM(CASE WHEN KTip='Diger' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(10,2)) AS [DigerKmp Ort Sepet Tutar],
    CAST(CASE WHEN SUM(CASE WHEN KTip='Diger' THEN GrossTotal ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='Diger' THEN DiscountTotal ELSE 0 END)*100.0/SUM(CASE WHEN KTip='Diger' THEN GrossTotal ELSE 0 END)
        ELSE 0 END AS decimal(5,1)) AS [DigerKmp Indirim Orani],
    SUM(CASE WHEN KTip='Kampanyasiz' THEN 1 ELSE 0 END) AS [Kampanyasiz Fis],
    SUM(CASE WHEN KTip='Kampanyasiz' THEN LineCount ELSE 0 END) AS [Kampanyasiz Urun],
    CAST(SUM(CASE WHEN KTip='Kampanyasiz' THEN GrossTotal-DiscountTotal ELSE 0 END) AS decimal(18,2)) AS [Kampanyasiz Net],
    CAST(CASE WHEN SUM(CASE WHEN KTip='Kampanyasiz' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='Kampanyasiz' THEN LineCount ELSE 0 END)*1.0/SUM(CASE WHEN KTip='Kampanyasiz' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(5,2)) AS [Kampanyasiz Ort Sepet Adet],
    CAST(CASE WHEN SUM(CASE WHEN KTip='Kampanyasiz' THEN 1 ELSE 0 END)>0
        THEN SUM(CASE WHEN KTip='Kampanyasiz' THEN GrossTotal-DiscountTotal ELSE 0 END)*1.0/SUM(CASE WHEN KTip='Kampanyasiz' THEN 1 ELSE 0 END)
        ELSE 0 END AS decimal(10,2)) AS [Kampanyasiz Ort Sepet Tutar],
    COUNT(*) AS [Toplam Fis],
    SUM(LineCount) AS [Toplam Urun],
    CAST(SUM(GrossTotal) AS decimal(18,2)) AS [Toplam Brut],
    CAST(SUM(DiscountTotal) AS decimal(18,2)) AS [Toplam Indirim],
    CAST(SUM(GrossTotal-DiscountTotal) AS decimal(18,2)) AS [Toplam Net],
    CAST(SUM(LineCount)*1.0/COUNT(*) AS decimal(5,2)) AS [Toplam Ort Sepet Adet],
    CAST(SUM(GrossTotal-DiscountTotal)*1.0/COUNT(*) AS decimal(10,2)) AS [Toplam Ort Sepet Tutar],
    CAST(CASE WHEN SUM(GrossTotal)>0
        THEN SUM(DiscountTotal)*100.0/SUM(GrossTotal) ELSE 0 END AS decimal(5,1)) AS [Toplam Indirim Orani],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN 1 ELSE 0 END)*100.0/COUNT(*) AS decimal(5,1)) AS [3Al2Ode Fis Payi],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN LineCount ELSE 0 END)*100.0/SUM(LineCount) AS decimal(5,1)) AS [3Al2Ode Urun Payi],
    CAST(SUM(CASE WHEN KTip='3Al2Ode' THEN GrossTotal-DiscountTotal ELSE 0 END)*100.0/SUM(GrossTotal-DiscountTotal) AS decimal(5,1)) AS [3Al2Ode Ciro Payi],
    SortKey, IsTotal
"""

SELECT_COLS = """
    Tarih, Magaza,
    [3Al2Ode Fis],[3Al2Ode Urun],[3Al2Ode Brut],[3Al2Ode Indirim],[3Al2Ode Net],
    [3Al2Ode Ort Sepet Adet],[3Al2Ode Ort Sepet Tutar],[3Al2Ode Indirim Orani],
    [DigerKmp Fis],[DigerKmp Urun],[DigerKmp Brut],[DigerKmp Indirim],[DigerKmp Net],
    [DigerKmp Ort Sepet Adet],[DigerKmp Ort Sepet Tutar],[DigerKmp Indirim Orani],
    [Kampanyasiz Fis],[Kampanyasiz Urun],[Kampanyasiz Net],
    [Kampanyasiz Ort Sepet Adet],[Kampanyasiz Ort Sepet Tutar],
    [Toplam Fis],[Toplam Urun],[Toplam Brut],[Toplam Indirim],[Toplam Net],
    [Toplam Ort Sepet Adet],[Toplam Ort Sepet Tutar],[Toplam Indirim Orani],
    [3Al2Ode Fis Payi],[3Al2Ode Urun Payi],[3Al2Ode Ciro Payi]
"""

SQL_RS1 = BASE_CTE + """
      AND YEAR(CAST(s.Date AS date))=YEAR(GETDATE()) AND MONTH(CAST(s.Date AS date))=MONTH(GETDATE())
),
gun_detay AS (
    SELECT CONVERT(varchar, Gun, 104) AS Tarih, Magaza,
        Gun AS SortKey, 0 AS IsTotal, KTip, LineCount, GrossTotal, DiscountTotal
    FROM base
    UNION ALL
    SELECT CONVERT(varchar, Gun, 104), 'TOPLAM',
        Gun, 1, KTip, LineCount, GrossTotal, DiscountTotal
    FROM base
    UNION ALL
    SELECT 'AY TOPLAM', '',
        CAST('2099-12-31' AS date), 1, KTip, LineCount, GrossTotal, DiscountTotal
    FROM base
),
ham AS (
SELECT Tarih, Magaza,
""" + METRIC_BLOCK + """
FROM gun_detay GROUP BY Tarih, Magaza, SortKey, IsTotal
)
SELECT TOP 500 """ + SELECT_COLS + """
FROM ham ORDER BY SortKey DESC, IsTotal DESC, Magaza
"""

SQL_RS2 = BASE_CTE + """
      AND NOT(YEAR(CAST(s.Date AS date))=YEAR(GETDATE()) AND MONTH(CAST(s.Date AS date))=MONTH(GETDATE()))
      AND CAST(s.Date AS date) >= CONVERT(date, '01.01.2026', 104)
),
ay_detay AS (
    SELECT RIGHT('0'+CAST(MONTH(Gun) AS varchar),2)+'.'+CAST(YEAR(Gun) AS varchar) AS Tarih, Magaza,
        CAST(CAST(YEAR(Gun) AS varchar)+'-'+RIGHT('0'+CAST(MONTH(Gun) AS varchar),2)+'-01' AS date) AS SortKey, 0 AS IsTotal,
        KTip, LineCount, GrossTotal, DiscountTotal
    FROM base
    UNION ALL
    SELECT RIGHT('0'+CAST(MONTH(Gun) AS varchar),2)+'.'+CAST(YEAR(Gun) AS varchar), 'TOPLAM',
        CAST(CAST(YEAR(Gun) AS varchar)+'-'+RIGHT('0'+CAST(MONTH(Gun) AS varchar),2)+'-01' AS date), 1,
        KTip, LineCount, GrossTotal, DiscountTotal
    FROM base
),
ham2 AS (
SELECT Tarih, Magaza,
""" + METRIC_BLOCK + """
FROM ay_detay GROUP BY Tarih, Magaza, SortKey, IsTotal
)
SELECT TOP 100 """ + SELECT_COLS + """
FROM ham2 ORDER BY SortKey DESC, IsTotal DESC, Magaza
"""

# --- BAŞLIKLAR & FORMAT ---
HEADERS = [
    "Tarih", "Magaza",
    "3Al2Ode Fis", "3Al2Ode Urun", "3Al2Ode Brut", "3Al2Ode Indirim", "3Al2Ode Net",
    "3Al2Ode Ort Sepet Adet", "3Al2Ode Ort Sepet Tutar", "3Al2Ode Indirim Orani %",
    "DigerKmp Fis", "DigerKmp Urun", "DigerKmp Brut", "DigerKmp Indirim", "DigerKmp Net",
    "DigerKmp Ort Sepet Adet", "DigerKmp Ort Sepet Tutar", "DigerKmp Indirim Orani %",
    "Kampanyasiz Fis", "Kampanyasiz Urun", "Kampanyasiz Net",
    "Kampanyasiz Ort Sepet Adet", "Kampanyasiz Ort Sepet Tutar",
    "Toplam Fis", "Toplam Urun", "Toplam Brut", "Toplam Indirim", "Toplam Net",
    "Toplam Ort Sepet Adet", "Toplam Ort Sepet Tutar", "Toplam Indirim Orani %",
    "3Al2Ode Fis Payi %", "3Al2Ode Urun Payi %", "3Al2Ode Ciro Payi %",
]

COL_TYPES = [
    'txt','txt',
    'int','int','cur','cur','cur','dec','cur','pct',
    'int','int','cur','cur','cur','dec','cur','pct',
    'int','int','cur','dec','cur',
    'int','int','cur','cur','cur','dec','cur','pct',
    'pct','pct','pct',
]

FMT_MAP = {'int': '#,##0', 'cur': '#,##0.00', 'dec': '0.00', 'pct': '0.0', 'txt': '@'}
HEADER_FILL = PatternFill("solid", fgColor="D6EAF8")
TOPLAM_FILL = PatternFill("solid", fgColor="FFF9C4")
BOLD = Font(bold=True)


def write_sheet(ws, cursor):
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    row_idx = 2
    for db_row in cursor.fetchall():
        is_toplam = (db_row[1] or '').upper() in ('TOPLAM', '')
        for c in range(len(HEADERS)):
            val = db_row[c]
            if hasattr(val, 'as_tuple'):
                val = float(val)
            cell = ws.cell(row=row_idx, column=c+1, value=val)
            ct = COL_TYPES[c]
            if ct != 'txt':
                cell.number_format = FMT_MAP[ct]
            if is_toplam:
                cell.fill = TOPLAM_FILL
                cell.font = BOLD
        row_idx += 1
    widths = [12, 16] + [14]*32
    for c, w in enumerate(widths[:len(HEADERS)], 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
    ws.freeze_panes = 'A2'
    return row_idx - 2


# ========================================================
# GUI
# ========================================================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("3Al2Ode Performans Raporu")
        self.geometry("480x220")
        self.resizable(False, False)

        # --- Kayıt yeri ---
        frm_path = ttk.LabelFrame(self, text="Kayıt Yeri", padding=8)
        frm_path.pack(fill='x', padx=12, pady=(12, 4))

        default_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        self.var_path = tk.StringVar(value=default_dir)
        ttk.Entry(frm_path, textvariable=self.var_path, width=48).pack(side='left', fill='x', expand=True)
        ttk.Button(frm_path, text="...", width=3, command=self._browse).pack(side='left', padx=(4, 0))

        # --- Dosya adı ---
        frm_name = ttk.LabelFrame(self, text="Dosya Adi", padding=8)
        frm_name.pack(fill='x', padx=12, pady=4)

        self.var_fname = tk.StringVar(value="3al2ode_performans_raporu.xlsx")
        ttk.Entry(frm_name, textvariable=self.var_fname, width=52).pack(fill='x')

        # --- Excel'i aç checkbox ---
        self.var_open = tk.BooleanVar(value=True)
        ttk.Checkbutton(self, text="Olusturulduktan sonra Excel'i ac",
                        variable=self.var_open).pack(padx=12, pady=4, anchor='w')

        # --- Alt bar: buton + durum ---
        frm_bottom = ttk.Frame(self)
        frm_bottom.pack(fill='x', padx=12, pady=(4, 12))

        self.lbl_status = ttk.Label(frm_bottom, text="Hazir.", foreground="gray")
        self.lbl_status.pack(side='left')

        self.btn_go = ttk.Button(frm_bottom, text="Rapor Olustur", command=self._start)
        self.btn_go.pack(side='right')

    def _browse(self):
        d = filedialog.askdirectory(initialdir=self.var_path.get())
        if d:
            self.var_path.set(d)

    def _set_status(self, msg, color="gray"):
        self.lbl_status.config(text=msg, foreground=color)
        self.update_idletasks()

    def _start(self):
        self.btn_go.config(state='disabled')
        threading.Thread(target=self._run, daemon=True).start()

    def _run(self):
        try:
            self._set_status("SQL Server'a baglaniliyor...", "blue")
            conn = pyodbc.connect(get_conn_str(), timeout=60)
            conn.timeout = 120

            wb = Workbook()

            # Sheet 1
            ws1 = wb.active
            ws1.title = "Mevcut Ay (Gun Detay)"
            self._set_status("Mevcut ay sorgusu calisiyor...", "blue")
            cur1 = conn.execute(SQL_RS1)
            n1 = write_sheet(ws1, cur1)

            # Sheet 2
            ws2 = wb.create_sheet("Onceki Aylar")
            self._set_status("Onceki aylar sorgusu calisiyor...", "blue")
            cur2 = conn.execute(SQL_RS2)
            n2 = write_sheet(ws2, cur2)

            conn.close()

            # Kaydet
            out = os.path.join(self.var_path.get(), self.var_fname.get())
            wb.save(out)

            self._set_status(f"Tamam! {n1}+{n2} satir yazildi.", "green")

            # Excel'i aç
            if self.var_open.get():
                os.startfile(out)

        except Exception as e:
            self._set_status("HATA!", "red")
            messagebox.showerror("Hata", str(e))
        finally:
            self.btn_go.config(state='normal')


if __name__ == "__main__":
    App().mainloop()
