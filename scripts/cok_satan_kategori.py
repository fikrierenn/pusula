# -*- coding: utf-8 -*-
"""
KATEGORİ BAZLI ÇOK SATAN 50 × STOK × BULUNURLUK — emitter (Excel: kategori başına 1 sheet).
Girdi : cok_satan_evren.csv (scripts/cok_satan_bulunurluk.py çekirdeği)
Ek veri: geçen yıl sezon satışı (Ağu-Eki 2025) + geçen yıl aynı-tarih stok (JIT ayrımı için)
Çıktı : <OUT>/cok-satan-kategori-50.xlsx  + <OUT>/kategori_ozet.txt (konsol özeti)
Kural  : sema metrics.envanter_exclusions (sinav_okullari_SATIS · poset_ambalaj · sahaf · dergi)
         sema metrics.bulunurluk_osa (kapsam_gun_mevsimsellik · jit_ayrimi_YoY)
"""
import os, re, sys, csv, datetime, collections
import pyodbc
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
R = os.path.join(os.path.dirname(__file__), "..")
OUT = os.environ.get("COKSATAN_OUT", ".")
SRC = os.path.join(OUT, "cok_satan_evren.csv")
MIN_STOK = 3
TOPN = 50

ENV = {}
with open(os.path.join(R, ".env"), encoding="utf-8") as fh:
    for ln in fh:
        m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.+?)\s*$", ln)
        if m and not ln.lstrip().startswith("#"):
            ENV[m.group(1)] = m.group(2).strip().strip('"')
host = ENV["MSSQL_HOST"]; port = ENV.get("MSSQL_PORT", "1433")
if not re.fullmatch(r"[A-Za-z0-9._\-]+", host) or not re.fullmatch(r"\d+", port):
    sys.exit("Geçersiz host/port (.env)")
cn = pyodbc.connect("Driver={ODBC Driver 18 for SQL Server};Server=%s,%s;Database=DerinSISBkm;UID=%s;PWD=%s;"
                    "TrustServerCertificate=yes;Timeout=30"
                    % (host, port, ENV["MSSQL_USER"], ENV["MSSQL_PASSWORD"]), timeout=30)
cn.timeout = 1200
cur = cn.cursor()

rows = list(csv.DictReader(open(SRC, encoding="utf-8-sig")))
f = lambda r, k: float(r[k]) if r[k] not in ("", None) else 0.0
iv = lambda r, k: int(float(r[k])) if r[k] not in ("", None) else None

# --- dışlamalar (sema) ---
SERVIS = lambda r: (r["katana"] or "").startswith("Sınav Okul") or (r["kat3"] or "") == "Sınav Okulları"
POSET = lambda r: "Poşet" in (r["ad"] or "") and (r["kat3"] or "") == "Genel"
DERGI = lambda r: "SÜRELİ YAYIN" in (r["ad"] or "").upper() or (r["kat3"] or "") == "Dergi"
SAHAF = lambda r: (r["ad"] or "").startswith("SHF-")
raf = [r for r in rows if not (SERVIS(r) or POSET(r) or DERGI(r) or SAHAF(r))]
ids = [int(r["stkID"]) for r in raf]
print("Raf evreni: %d SKU" % len(ids), flush=True)


def chunks(x, n=800):
    for i in range(0, len(x), n):
        yield x[i:i + n]


# --- geçen yıl sezon (Ağu-Eki 2025) + geçen yıl aynı-tarih stok ---
today = datetime.date.today()
sez, yoy = {}, {}
for ch in chunks(ids):
    inl = ",".join(str(i) for i in ch)
    cur.execute("""
    SELECT h.ehstkID, -SUM(h.ehAdetN),
           SUM(CASE WHEN h.ehTip IN (4,100) THEN h.ehTutarN ELSE -h.ehTutarN END)
    FROM dbo.irsHrk h WITH(NOLOCK)
    WHERE h.ehTip IN (4,100,5,101) AND h.ehMekan IN (1,4477,4478) AND h.ehAltDepo=0
      AND h.ehTrhS>=DATEFROMPARTS(%d,8,1) AND h.ehTrhS<DATEFROMPARTS(%d,11,1)
      AND h.ehstkID IN (%s)
    GROUP BY h.ehstkID""" % (today.year - 1, today.year - 1, inl))
    for r in cur.fetchall():
        sez[int(r[0])] = (float(r[1]), float(r[2]))
    cur.execute("""
    SELECT h.ehstkID, SUM(CASE WHEN h.ehTrhS < DATEFROMPARTS(%d,%d,%d) THEN h.ehAdetN ELSE 0 END)
    FROM dbo.irsHrk h WITH(NOLOCK)
    WHERE h.ehMekan IN (1,4477,4478) AND h.ehAltDepo=0 AND h.ehstkID IN (%s)
    GROUP BY h.ehstkID""" % (today.year - 1, today.month, today.day + 1, inl))
    for r in cur.fetchall():
        yoy[int(r[0])] = float(r[1] or 0)
print("Sezon/YoY veri: %d / %d SKU" % (len(sez), len(yoy)), flush=True)

# --- açık sipariş (son 60 gün, sip.eTarih — eTarihS DEĞİL; adet = sipAyr.ehAdet) ---
sip = collections.defaultdict(float)
for ch in chunks(ids):
    cur.execute("""
    SELECT a.ehStkID, SUM(a.ehAdet)
    FROM dbo.sipAyr a WITH(NOLOCK) JOIN dbo.sip s WITH(NOLOCK) ON s.eID=a.ehID
    WHERE a.ehStkID IN (%s) AND s.eTarih >= DATEADD(DAY,-60,GETDATE()) AND s.eDurum <> 2
    GROUP BY a.ehStkID""" % ",".join(str(i) for i in ch))
    for r in cur.fetchall():
        sip[int(r[0])] = float(r[1] or 0)
cn.close()

MEK = [("fsm", "FSM"), ("ozl", "Özlüce"), ("ist", "İst.Yolu")]


def enrich(r):
    sid = int(r["stkID"])
    g3 = f(r, "adet3") / 90.0
    kaps = (f(r, "stok_sube") / g3) if g3 > 0 else None
    kuru_now = [n for k, n in MEK if f(r, "stok_" + k) < MIN_STOK and f(r, "satis_" + k) > 0]
    sad, sci = sez.get(sid, (0.0, 0.0))
    tot = f(r, "stok_sube") + f(r, "stok_depo")
    hazir = (tot / sad) if sad >= 100 else None
    s25 = max(0.0, yoy.get(sid, 0.0))
    tani = ""
    if hazir is not None and hazir < 0.25:
        tani = "JIT (gecen yil da dusuk)" if (s25 / sad) < 0.25 else "GERCEK GERILEME"
    birim = f(r, "ciro12") / max(1.0, f(r, "adet12"))
    return dict(r=r, g3=g3, kaps=kaps, kuru_now=kuru_now, sezon_ad=sad, sezon_ciro=sci,
                hazir=hazir, stok25=s25, tani=tani, birim=birim, siparis=sip.get(sid, 0.0),
                bagli=f(r, "stok_sube") * f(r, "alis"))


enr = [enrich(r) for r in raf]
by_kat = collections.defaultdict(list)
for e in enr:
    by_kat[(e["r"]["kat3"] or "Tanımsız")].append(e)

# kategoriler ciro azalan
kat_ciro = {k: sum(f(e["r"], "ciro12") for e in v) for k, v in by_kat.items()}
kats = [k for k, _ in sorted(kat_ciro.items(), key=lambda t: -t[1])]

HDR = ["#", "stkID", "Ürün", "Marka", "Tedarikçi", "adet12", "ciro12 ₺", "birim ₺", "adet3", "adet1",
       "stok FSM", "stok Özlüce", "stok İst.Yolu", "stok Şube", "stok Depo", "kapsam gün",
       "kuru ay FSM", "kuru ay Özlüce", "kuru ay İst.Yolu", "ŞİMDİ KURU",
       "sezon25 adet", "hazırlık %", "geçenyıl stok", "tanı", "sipariş 60g", "bağlı ₺"]
wb = openpyxl.Workbook()
wb.remove(wb.active)
BOLD = Font(bold=True, color="FFFFFF")
KIRMIZI = PatternFill("solid", fgColor="E30622")
UYARI = PatternFill("solid", fgColor="FDE8E8")
SARI = PatternFill("solid", fgColor="FFF7E0")

ozet_rows = []
for kat in kats:
    lst = sorted(by_kat[kat], key=lambda e: -f(e["r"], "ciro12"))
    if len(lst) < 3 or kat_ciro[kat] < 200000:
        continue
    top = lst[:TOPN]
    ws = wb.create_sheet(re.sub(r"[\[\]\*\?/\\:]", "-", kat)[:31])
    ws.append(HDR)
    for c in range(1, len(HDR) + 1):
        ws.cell(1, c).font = BOLD; ws.cell(1, c).fill = KIRMIZI
        ws.cell(1, c).alignment = Alignment(horizontal="center", wrap_text=True)
    for n, e in enumerate(top, 1):
        r = e["r"]
        ws.append([n, int(r["stkID"]), r["ad"], r["marka"], r["firma"],
                   f(r, "adet12"), round(f(r, "ciro12"), 2), round(e["birim"], 2), f(r, "adet3"), f(r, "adet1"),
                   f(r, "stok_fsm"), f(r, "stok_ozl"), f(r, "stok_ist"), f(r, "stok_sube"), f(r, "stok_depo"),
                   round(e["kaps"], 1) if e["kaps"] is not None else "",
                   iv(r, "kuru_fsm"), iv(r, "kuru_ozl"), iv(r, "kuru_ist"),
                   ", ".join(e["kuru_now"]),
                   e["sezon_ad"] or "", round(100 * e["hazir"], 1) if e["hazir"] is not None else "",
                   e["stok25"], e["tani"], e["siparis"], round(e["bagli"], 0)])
        row = ws.max_row
        if e["kuru_now"]:
            for c in range(1, len(HDR) + 1):
                ws.cell(row, c).fill = UYARI
        elif e["kaps"] is not None and e["kaps"] > 365:
            for c in range(1, len(HDR) + 1):
                ws.cell(row, c).fill = SARI
    ws.freeze_panes = "C2"
    for c, w in enumerate([4, 9, 46, 18, 22, 9, 12, 9, 8, 8, 9, 11, 11, 10, 10, 10,
                           8, 9, 9, 20, 11, 10, 12, 24, 10, 11], 1):
        ws.column_dimensions[get_column_letter(c)].width = w

    # kategori özeti
    n_all = len(lst)
    kuru = [e for e in lst if e["kuru_now"]]
    tam_kuru = [e for e in lst if all(f(e["r"], "stok_" + k) < MIN_STOK for k, _ in MEK) and f(e["r"], "adet3") > 0]
    over = [e for e in lst if e["kaps"] is not None and e["kaps"] > 365]
    ger = [e for e in lst if e["tani"] == "GERCEK GERILEME"]
    jit = [e for e in lst if e["tani"] == "JIT (gecen yil da dusuk)"]
    ka = [sum(x for x in (iv(e["r"], "kuru_fsm"), iv(e["r"], "kuru_ozl"), iv(e["r"], "kuru_ist")) if x is not None) / 3.0
          for e in lst if iv(e["r"], "kuru_ist") is not None]
    ozet_rows.append(dict(
        kat=kat, sku=n_all, ciro=kat_ciro[kat], adet=sum(f(e["r"], "adet12") for e in lst),
        top50_pay=100 * sum(f(e["r"], "ciro12") for e in top) / kat_ciro[kat],
        kuru=len(kuru), kuru_pay=100 * len(kuru) / n_all, kuru_ciro=sum(f(e["r"], "ciro12") for e in kuru),
        tam_kuru=len(tam_kuru), over=len(over), bagli=sum(e["bagli"] for e in over),
        ger=len(ger), ger_ciro=sum(e["sezon_ciro"] for e in ger), jit=len(jit),
        kuru_ay=(sum(ka) / len(ka)) if ka else 0.0,
        top50_kuru=len([e for e in top if e["kuru_now"]]),
        top50_ger=len([e for e in top if e["tani"] == "GERCEK GERILEME"])))

# özet sheet en başa
ws = wb.create_sheet("ÖZET", 0)
OH = ["Kategori", "SKU", "ciro12 ₺", "adet12", "top50 ciro payı %", "ŞİMDİ KURU SKU", "kuru %",
      "kuru SKU cirosu ₺", "3 şube kuru", "ort kuru ay/12", "kapsam>365g SKU", "bağlı para ₺",
      "sezon GERÇEK GERİLEME", "gerileme sezon cirosu ₺", "sezon JIT", "top50'de kuru", "top50'de gerileme"]
ws.append(OH)
for c in range(1, len(OH) + 1):
    ws.cell(1, c).font = BOLD; ws.cell(1, c).fill = KIRMIZI
    ws.cell(1, c).alignment = Alignment(horizontal="center", wrap_text=True)
for o in ozet_rows:
    ws.append([o["kat"], o["sku"], round(o["ciro"], 2), o["adet"], round(o["top50_pay"], 1),
               o["kuru"], round(o["kuru_pay"], 1), round(o["kuru_ciro"], 2), o["tam_kuru"],
               round(o["kuru_ay"], 2), o["over"], round(o["bagli"], 0), o["ger"],
               round(o["ger_ciro"], 2), o["jit"], o["top50_kuru"], o["top50_ger"]])
ws.freeze_panes = "B2"
for c, w in enumerate([22, 8, 15, 12, 12, 12, 9, 16, 11, 11, 12, 14, 14, 16, 10, 11, 12], 1):
    ws.column_dimensions[get_column_letter(c)].width = w

p = os.path.join(OUT, "cok-satan-kategori-50.xlsx")
wb.save(p)
print("OK -> %s (%d kategori sheet)" % (p, len(ozet_rows)), flush=True)

# --- markdown emitter (aynı çekirdek, ikinci çıktı) ---
MD = os.environ.get("COKSATAN_MD")
if MD:
    L = []
    L.append("# Kategori Bazlı Çok Satan 50 × Stok × Bulunurluk — %s\n" % today.strftime("%d.%m.%Y"))
    L.append("**Pencere:** son 12 tam ay (Ağustos %d kısmi → hariç) · **Evren:** şube perakende "
             "(`irsHrk` ehTip 4/100 − 5/101, mekan 1/4477/4478, `ehAltDepo=0`, `urnTip=0`), tutarlar **KDV hariç**.\n"
             % today.year)
    L.append("**Hariç (sema `metrics.envanter_exclusions`):** Sınav Okulları paket/hizmet SKU'ları · poşet (ambalaj) · "
             "süreli yayın · sahaf (`SHF-`).  **Raf evreni: %d SKU.**\n" % len(raf))
    L.append("**Kolonlar:** `kapsam` = şube stoğu ÷ son-3-ay günlük hız (⚠ mevsimsel üründe sezon dışı şişer) · "
             "`kuru ay` = son 12 ayın kaç ayında giriş bakiyesi <%d (FSM/Özlüce/İst.Yolu) · "
             "`ŞİMDİ KURU` = o şubede stok <%d ama son 12 ayda satış var · "
             "`hazırlık` = (şube+depo stok) ÷ geçen yıl Ağu-Eki satışı · "
             "`tanı` = GERİLEME (geçen yıl aynı tarihte stok vardı) / JIT (geçen yıl da düşüktü, normal mod).\n"
             % (MIN_STOK, MIN_STOK))
    L.append("Tam 50'lik listeler + tüm kolonlar: `cok-satan-kategori-50.xlsx` (kategori başına 1 sheet).\n")
    L.append("\n## Kategori özeti\n")
    L.append("| Kategori | SKU | ciro12 ₺ | top50 payı | ŞİMDİ KURU | kuru SKU cirosu | 3 şube kuru | "
             "ort kuru ay/12 | kapsam>365g | bağlı para ₺ | sezon GERİLEME |")
    L.append("|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|")
    for o in ozet_rows:
        L.append("| **%s** | %d | %s | %.0f%% | %d (%.0f%%) | %s | %d | %s | %d | %s | **%d** |" % (
            o["kat"], o["sku"], "{:,.0f}".format(o["ciro"]).replace(",", "."), o["top50_pay"],
            o["kuru"], o["kuru_pay"], "{:,.0f}".format(o["kuru_ciro"]).replace(",", "."),
            o["tam_kuru"], ("%.1f" % o["kuru_ay"]).replace(".", ","), o["over"],
            "{:,.0f}".format(o["bagli"]).replace(",", "."), o["ger"]))
    for kat in [o["kat"] for o in ozet_rows]:
        lst = sorted(by_kat[kat], key=lambda e: -f(e["r"], "ciro12"))[:TOPN]
        o = [x for x in ozet_rows if x["kat"] == kat][0]
        L.append("\n---\n\n## %s — top %d (ciro)\n" % (kat, len(lst)))
        L.append("%d SKU · %s ₺ ciro · top50 payı %%%.0f · şimdi kuru %d SKU (%%%.0f) · "
                 "ort kuru ay %.1f/12 · kapsam>365g %d SKU (%s ₺ bağlı) · sezon gerilemesi %d SKU\n" % (
                     o["sku"], "{:,.0f}".format(o["ciro"]).replace(",", "."), o["top50_pay"],
                     o["kuru"], o["kuru_pay"], o["kuru_ay"], o["over"],
                     "{:,.0f}".format(o["bagli"]).replace(",", "."), o["ger"]))
        L.append("| # | Ürün | Marka | adet12 | ciro12 ₺ | stok F/Ö/İ | depo | kapsam | kuru ay | ŞİMDİ KURU | "
                 "sezon25 | hazırlık | tanı | sip.60g |")
        L.append("|---:|---|---|---:|---:|---:|---:|---:|:---:|---|---:|---:|---|---:|")
        for n, e in enumerate(lst, 1):
            r = e["r"]
            L.append("| %d | %s | %s | %s | %s | %.0f/%.0f/%.0f | %.0f | %s | %s/%s/%s | %s | %s | %s | %s | %s |" % (
                n, (r["ad"] or "").replace("|", "/")[:52], (r["marka"] or "")[:18],
                "{:,.0f}".format(f(r, "adet12")).replace(",", "."),
                "{:,.0f}".format(f(r, "ciro12")).replace(",", "."),
                f(r, "stok_fsm"), f(r, "stok_ozl"), f(r, "stok_ist"), f(r, "stok_depo"),
                ("%.0f" % e["kaps"]) if e["kaps"] is not None else "–",
                iv(r, "kuru_fsm"), iv(r, "kuru_ozl"), iv(r, "kuru_ist"),
                ("**" + ", ".join(e["kuru_now"]) + "**") if e["kuru_now"] else "",
                ("%.0f" % e["sezon_ad"]) if e["sezon_ad"] >= 100 else "",
                ("%.0f%%" % (100 * e["hazir"])) if e["hazir"] is not None else "",
                ("**GERİLEME**" if e["tani"] == "GERCEK GERILEME" else ("JIT" if e["tani"] else "")),
                ("%.0f" % e["siparis"]) if e["siparis"] else ""))
    open(MD, "w", encoding="utf-8").write("\n".join(L) + "\n")
    print("MD -> %s" % MD, flush=True)

# konsol özeti
print("\n%-22s %5s %11s %7s %6s %9s %6s %8s %6s %6s" %
      ("KATEGORİ", "SKU", "ciro12 M₺", "top50%", "kuru", "kuruCiroM", "3şube", "kuru_ay", "over", "GERİL"))
for o in ozet_rows:
    print("%-22s %5d %11.1f %6.0f%% %5d %9.1f %6d %8.1f %6d %6d" %
          (o["kat"][:22], o["sku"], o["ciro"] / 1e6, o["top50_pay"], o["kuru"], o["kuru_ciro"] / 1e6,
           o["tam_kuru"], o["kuru_ay"], o["over"], o["ger"]))
