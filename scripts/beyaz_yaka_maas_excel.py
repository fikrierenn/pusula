# -*- coding: utf-8 -*-
"""Beyaz yaka maas listesi -> Excel (emitter).

Cekirdek (hesap/filtre/is mantigi) sorgular/2026-08-29-beyaz-yaka-maas-listesi.sql
dosyasindadir. Bu script SADECE bicimlendirir -- SUM/CASE/join YAPMAZ
(.claude/rules/emitter-ayrimi.md).

Girdi : cekirdek sorgunun JSON ciktisi (MCP zirve veya sqlcli --format json).
Cikti : xlsx (Liste + Departman ozeti sayfalari).

Kullanim:
    python scripts/beyaz_yaka_maas_excel.py \
        briefings/2026-08-29/beyaz-yaka-veri.json \
        briefings/2026-08-29/beyaz-yaka-maas-listesi.xlsx
"""
import json
import sys
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TL = '#,##0.00\\ "₺"'  # 1.234,56 TL
BASLIK_DOLGU = PatternFill("solid", fgColor="E30622")  # BKM kirmizisi
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=10)
TOPLAM_DOLGU = PatternFill("solid", fgColor="F2F2F2")
INCE = Side(style="thin", color="D9D9D9")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)

KOLONLAR = OrderedDict([
    ("Personelno",    ("Personel No", 13, None)),
    ("AdSoyad",       ("Ad Soyad", 28, None)),
    ("Departman",     ("Departman", 26, None)),
    ("Unvan",         ("Unvan", 32, None)),
    ("Firma",         ("Firma", 22, None)),
    ("GirisTarihi",   ("Giris Tarihi", 13, None)),
    ("KidemYil",      ("Kidem (yil)", 11, "0.0")),
    ("UcretOran",     ("Ucret Orani", 13, TL)),
    ("UcretSekli",    ("Ucret Sekli", 12, None)),
    ("AylikNet",      ("Aylik Net", 15, TL)),
    ("Cinsiyet",      ("Cinsiyet", 10, None)),
    ("OgrenimDurumu", ("Ogrenim", 22, None)),
])


def _basliklari_yaz(ws, basliklar):
    ws.append(basliklar)
    for h in ws[1]:
        h.fill = BASLIK_DOLGU
        h.font = BASLIK_YAZI
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = KENAR
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def liste_sayfasi(wb, satirlar):
    ws = wb.create_sheet("Liste")
    _basliklari_yaz(ws, [b for b, _, _ in KOLONLAR.values()])

    for r in satirlar:
        ws.append([r.get(k) for k in KOLONLAR])

    son = ws.max_row
    for idx, (_, genislik, bicim) in enumerate(KOLONLAR.values(), start=1):
        harf = get_column_letter(idx)
        ws.column_dimensions[harf].width = genislik
        if bicim:
            for hucre in ws[harf][1:]:
                hucre.number_format = bicim
        for hucre in ws[harf][1:]:
            hucre.border = KENAR

    # Toplam satiri: Excel formulu -- Python'da toplam hesaplanmaz (emitter kurali)
    ws.append([])
    toplam = ws.max_row + 1
    ws.cell(toplam, 1, "TOPLAM").font = Font(bold=True)
    ws.cell(toplam, 2, f"=COUNTA(B2:B{son}) & \" kisi\"").font = Font(bold=True)
    h = ws.cell(toplam, 10, f"=SUM(J2:J{son})")
    h.font = Font(bold=True)
    h.number_format = TL
    for c in range(1, len(KOLONLAR) + 1):
        ws.cell(toplam, c).fill = TOPLAM_DOLGU
        ws.cell(toplam, c).border = KENAR

    ws.auto_filter.ref = f"A1:{get_column_letter(len(KOLONLAR))}{son}"
    return ws


def ozet_sayfasi(wb, satirlar, liste_son_satir):
    """Departman ozeti. Kisi sayisi ve tutarlar Excel formuluyle Liste'den turer."""
    ws = wb.create_sheet("Departman Ozeti")
    _basliklari_yaz(ws, ["Departman", "Kisi", "Toplam Aylik Net", "Ortalama", "Pay %"])

    departmanlar = sorted({r["Departman"] for r in satirlar})
    ilk = 2
    for i, d in enumerate(departmanlar):
        s = ilk + i
        ws.cell(s, 1, d)
        ws.cell(s, 2, f'=COUNTIF(Liste!C2:C{liste_son_satir},A{s})')
        ws.cell(s, 3, f'=SUMIF(Liste!C2:C{liste_son_satir},A{s},Liste!J2:J{liste_son_satir})')
        ws.cell(s, 4, f'=IF(B{s}=0,"",C{s}/B{s})')
        ws.cell(s, 5, f'=IF(SUM($C$2:$C${ilk + len(departmanlar) - 1})=0,"",'
                      f'C{s}/SUM($C$2:$C${ilk + len(departmanlar) - 1}))')

    son = ilk + len(departmanlar) - 1
    toplam = son + 1
    ws.cell(toplam, 1, "TOPLAM").font = Font(bold=True)
    ws.cell(toplam, 2, f"=SUM(B{ilk}:B{son})").font = Font(bold=True)
    ws.cell(toplam, 3, f"=SUM(C{ilk}:C{son})").font = Font(bold=True)

    for harf, genislik, bicim in [("A", 30, None), ("B", 8, "0"),
                                  ("C", 18, TL), ("D", 16, TL), ("E", 10, "0.0%")]:
        ws.column_dimensions[harf].width = genislik
        for hucre in ws[harf][1:]:
            hucre.border = KENAR
            if bicim:
                hucre.number_format = bicim
    for c in range(1, 6):
        ws.cell(toplam, c).fill = TOPLAM_DOLGU
    return ws


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2

    girdi, cikti = Path(sys.argv[1]), Path(sys.argv[2])
    satirlar = json.loads(girdi.read_text(encoding="utf-8"))
    if not satirlar:
        print(f"HATA: {girdi} bos -- Excel uretilmedi.", file=sys.stderr)
        return 1

    eksik = [k for k in KOLONLAR if k not in satirlar[0]]
    if eksik:
        print(f"HATA: girdide eksik kolon(lar): {', '.join(eksik)}", file=sys.stderr)
        return 1

    wb = Workbook()
    wb.remove(wb.active)
    liste = liste_sayfasi(wb, satirlar)
    # Toplam satiri + bos satir listenin sonuna eklendi -> veri son satiri:
    ozet_sayfasi(wb, satirlar, len(satirlar) + 1)

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    print(f"OK: {cikti}  ({len(satirlar)} satir, {liste.max_row} Excel satiri)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
