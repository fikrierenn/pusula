# -*- coding: utf-8 -*-
"""Magaza bazli tek tablo -> Excel. Yan yana rakam, TOPLAM + FARK satiri, altta notlar.

Konvansiyon: as-of aktif = Igt <= T AND (Ict IS NULL OR Ict >= T)
  -> Zirve dbo.sp_PersonelKarsilastirma_Ozet ile BIREBIR (mutabakat 02.09.2026: 322 -> 335, +13).

Girdi : briefings/sezon-kadro-20260902/veri.json
Cikti : tek sayfali xlsx
Kullanim: python scripts/magaza_tablo_excel.py <veri.json> <cikti.xlsx>
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ADET = "#,##0"
BASLIK = PatternFill("solid", fgColor="E30622")
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=10)
GRI = PatternFill("solid", fgColor="F2F2F2")
FARK_DOLGU = PatternFill("solid", fgColor="FFF3CD")
INCE = Side(style="thin", color="D9D9D9")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
NOT_YAZI = Font(italic=True, size=9, color="666666")

# (ust baslik, alt baslik, genislik, format) — ust baslik ayni olan komsu ikili birlestirilir
KOLONLAR = [
    ("", "Magaza", 14, None),
    ("30.06", "2025", 8, "0"), ("30.06", "2026", 8, "0"),
    ("Sezon 31.08", "2025", 9, "0"), ("Sezon 31.08", "2026", 9, "0"),
    ("Kadro 31.08", "2025", 9, "0"), ("Kadro 31.08", "2026", 9, "0"),
    ("Toplam 31.08", "2025", 10, "0"), ("Toplam 31.08", "2026", 10, "0"),
    ("Fis", "2025", 11, ADET), ("Fis", "2026", 11, ADET),
    ("Kalem", "2025", 11, ADET), ("Kalem", "2026", 11, ADET),
    ("Urun adedi", "2025", 12, ADET), ("Urun adedi", "2026", 12, ADET),
    ("Net ciro (TL)", "2025", 15, ADET), ("Net ciro (TL)", "2026", 15, ADET),
]


def main(argv):
    if len(argv) != 3:
        print(__doc__)
        return 2
    veri = json.loads(Path(argv[1]).read_text(encoding="utf-8"))
    po = veri["patron_ozet"]
    kadro = {r["magaza"]: r for r in po["magaza"]}
    hacim = {r["magaza"]: r for r in veri["magaza_is_hacmi"]["satirlar"]}
    sira = ["IST. YOLU", "OZLUCE", "FSM", "HEYKEL", "SURA"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Magaza"

    # iki satirli baslik
    for i, (ust, alt, gen, _f) in enumerate(KOLONLAR, start=1):
        ws.cell(1, i, ust)
        ws.cell(2, i, alt)
        ws.column_dimensions[get_column_letter(i)].width = gen
    i = 2
    while i < len(KOLONLAR):
        if KOLONLAR[i - 1][0] == KOLONLAR[i][0]:
            ws.merge_cells(start_row=1, start_column=i, end_row=1, end_column=i + 1)
            i += 2
        else:
            i += 1
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    for satir in (1, 2):
        for h in ws[satir]:
            h.fill = BASLIK
            h.font = BASLIK_YAZI
            h.alignment = Alignment(horizontal="center", vertical="center")
            h.border = KENAR
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 16
    ws.freeze_panes = "B3"

    ilk = 3
    s = ilk
    for m in sira:
        k, h = kadro[m], hacim.get(m)
        ws.append([
            m,
            k["taban25"], k["taban26"],
            k["sez25"], k["sez26"],
            k["kad25"], k["kad26"],
            k["sez25"] + k["kad25"], k["sez26"] + k["kad26"],
            h["fis25"] if h else None, h["fis26"] if h else None,
            h["kalem25"] if h else None, h["kalem26"] if h else None,
            round(h["adet25"]) if h else None, round(h["adet26"]) if h else None,
            round(h["net25"]) if h else None, round(h["net26"]) if h else None,
        ])
        s += 1
    son = s - 1

    # TOPLAM (SUM)
    toplam_satir = s
    ws.append(["TOPLAM"] + [f"=SUM({get_column_letter(c)}{ilk}:{get_column_letter(c)}{son})"
                            for c in range(2, len(KOLONLAR) + 1)])
    for h in ws[toplam_satir]:
        h.fill = GRI
        h.font = Font(bold=True)

    # FARK (2026 - 2025) — her ikilinin ikinci kolonunda, ilkinde bos
    fark_satir = toplam_satir + 1
    fark = ["FARK"]
    for i in range(1, len(KOLONLAR)):
        kol = get_column_letter(1 + i)
        if KOLONLAR[i][1] == "2026":
            onceki = get_column_letter(i)
            fark.append(f"={kol}{toplam_satir}-{onceki}{toplam_satir}")
        else:
            fark.append(None)
    ws.append(fark)
    for h in ws[fark_satir]:
        h.fill = FARK_DOLGU
        h.font = Font(bold=True)

    for satir in ws.iter_rows(min_row=ilk, max_row=fark_satir):
        for h, (_u, _a, _g, fmt) in zip(satir, KOLONLAR):
            h.border = KENAR
            if fmt:
                h.number_format = fmt

    e25, e26 = sum(r["engelli25"] for r in po["magaza"]), sum(r["engelli26"] for r in po["magaza"])
    k25, k26 = sum(r["etkinlik25"] for r in po["magaza"]), sum(r["etkinlik26"] for r in po["magaza"])
    eng = " · ".join(f"{r['magaza']} {r['engelli25']} -> {r['engelli26']}"
                     for r in po["magaza"] if r["engelli25"] or r["engelli26"])
    etk = " · ".join(f"{r['magaza']} {r['etkinlik25']} -> {r['etkinlik26']}"
                     for r in po["magaza"] if r["etkinlik25"] or r["etkinlik26"])
    notlar = [
        "30.06 = sezon oncesi kadrolu taban · Sezon/Kadro/Toplam 31.08 = o tarihte fiilen calisan kisi "
        "· Fis/Kalem/Urun adedi/Net ciro = 1 Temmuz - 31 Agustos, perakende satis fisi, KDV haric.",
        f"Kadrolu icindeki ENGELLI (31.08): {eng} · magaza toplami {e25} -> {e26}. "
        "Ayrac 4857/30 engelli istihdam tesviki (+ Ozurlulukkodu izi); tesviksiz engelli bu sayida gorunmez.",
        f"Kadrolu icindeki ETKINLIK kadrosu (31.08): {etk} · magaza toplami {k25} -> {k26}.",
        po["konvansiyon"],
        "Heykel ve Sura EncoreMerkez POS raporlamasinda yok -> is hacmi kolonlari bos; "
        "TOPLAM/FARK is hacmi sutunlari uc magazanin toplamidir.",
    ]
    for i, txt in enumerate(notlar):
        ws.cell(fark_satir + 2 + i, 1, txt).font = NOT_YAZI

    cikti = Path(argv[2])
    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    print(f"yazildi: {cikti} (1 sayfa · {son - ilk + 1} magaza + TOPLAM + FARK · {len(KOLONLAR)} kolon)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
