# -*- coding: utf-8 -*-
"""Sunum ORTAK katmani: marka paleti + bicim yardimcilari + sablon durumu + cizim yardimcilari.

Sablon durumu (Presentation nesnesi + arka plan kaynak slaytlari) MODUL GLOBALIDIR; tek giris
noktasi `ac(tpl)`. Slayt modulleri yalniz buradaki yardimcilari cagirir, durumu bilmez.
Bolunme gerekcesi + harita: plans/40-sunum-deste-split.md (K-23).
"""
import copy as _copy
import os
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

try:   # Windows cp1254 konsolu: ok/uyari isaretleri UnicodeEncodeError veriyordu
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except (AttributeError, OSError) as _e:
    print("stdout utf-8 yapilamadi: %s" % _e)

ONCEKI, CARI = 2025, 2026        # kiyas yillari (verimlilik_ortak.py ile ayni)

IKON = Path(os.path.expanduser("~")) / ".claude" / "skills" / "bkm-sunum" / "assets"
ICOW, ICOR = str(IKON / "_icons"), str(IKON / "_icons_red")

RED = RGBColor(0xE3, 0x06, 0x22); DRED = RGBColor(0xA6, 0x00, 0x1A); ROSE = RGBColor(0xC0, 0x14, 0x2B)
CHAR = RGBColor(0x2B, 0x2B, 0x2B); GREY = RGBColor(0x6E, 0x6E, 0x6E); LGREY = RGBColor(0xF2, 0xF2, 0xF2)
MGREY = RGBColor(0x9A, 0x9A, 0x9A); WHITE = RGBColor(0xFF, 0xFF, 0xFF); INK = RGBColor(0x33, 0x33, 0x33)
YESIL = RGBColor(0xA6, 0x00, 0x1A)   # marka: yesil YASAK -> "artis" rakami koyu kirmizi
SET = [RED, ROSE, RGBColor(0x7A, 0x10, 0x20), RGBColor(0x4A, 0x4A, 0x4A), RGBColor(0x8C, 0x8C, 0x8C)]

# ---- sablon durumu: ac() doldurur
pr = None
TPL = None
TITLE_BG = CONTENT_BG = None


def ac(tpl):
    """Sablonu ac, ornek slaytlari dus, arka plan kaynaklarini hazirla. TEK giris noktasi."""
    global pr, TPL, TITLE_BG, CONTENT_BG
    TPL = tpl
    pr = Presentation(tpl)
    srcpr = Presentation(tpl)
    TITLE_BG, CONTENT_BG = srcpr.slides[0], srcpr.slides[1]
    sld = pr.slides._sldIdLst
    for sid in list(sld):
        try:
            pr.part.drop_rel(sid.get(qn('r:id')))
        except KeyError as e:      # iliski zaten yok — bilgi amacli, sessiz yutma yok
            print("  ⚠ şablon slayt ilişkisi bulunamadı, atlandı: %s" % e, flush=True)
        sld.remove(sid)
    return pr


def tr_title(metin):
    """Turkce baslik-buyutme. Python .title() I->i, İ->i̇ bozar (KIRTASİYE -> Kirtasiye)."""
    kucuk = {"I": "ı", "İ": "i", "Ş": "ş", "Ğ": "ğ", "Ü": "ü", "Ö": "ö", "Ç": "ç"}
    out = []
    for kelime in metin.split():
        if kelime in ("FSM", "POS", "ERP", "KDV", "İK"):   # gercek kisaltmalar aynen kalir
            out.append(kelime)                             # ("MAL" gibi 3 harfli kelimeler DEGIL)
            continue
        ilk, kalan = kelime[0], kelime[1:]
        out.append(ilk + "".join(kucuk.get(ch, ch.lower()) for ch in kalan))
    return " ".join(out)


def bin(x, ondalik=0):
    s = ("{:,.%df}" % ondalik).format(x).replace(",", "·").replace(".", ",").replace("·", ".")
    return s


def yuzde(x, ondalik=1):
    return ("%+." + str(ondalik) + "f%%").replace("%%", "%") % (x * 100)


def yzd(x):   # "+%34,9" — Turkce yazim (isaret, yuzde, virgul)
    return ("+" if x >= 0 else "−") + "%" + ("%.1f" % abs(x * 100)).replace(".", ",")


def L(name):
    """Sablon layout'u. K-14: bulunamazsa SESSIZCE bos layout'a dusmez — cikar.

    Eski davranis: slide_layouts[6] (bos) donuyordu -> baslik placeholder'i olmayan slayt,
    setph() de None donup sessizce gecince BASLIKSIZ deste uretiliyordu.
    """
    for l in pr.slide_masters[0].slide_layouts:
        if l.name == name:
            return l
    sys.exit("ŞABLON LAYOUT BULUNAMADI: '%s'. Şablon (%s) değişmiş — mevcut layout'lar: %s"
             % (name, TPL, ", ".join(l.name for l in pr.slide_masters[0].slide_layouts)))


def copy_bg(src, dst):
    csrc = src._element.find(qn('p:cSld')); bg = csrc.find(qn('p:bg'))
    if bg is None:
        return
    bg2 = _copy.deepcopy(bg)
    for blip in bg2.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        if not rid:
            continue
        nrid = dst.part.relate_to(src.part.related_part(rid), RT.IMAGE)
        blip.set(qn('r:embed'), nrid)
    dst._element.find(qn('p:cSld')).insert(0, bg2)


def add(name):
    sl = pr.slides.add_slide(L(name))
    copy_bg(TITLE_BG if name in ("Başlık Slaydı", "Bölüm Üst Bilgisi") else CONTENT_BG, sl)
    return sl


def setph(sl, idx, text):
    """Placeholder metni. Icerik slaytlarinda BASLIK KUTUSU SABITLENIR.

    ⚠ Yerlesim dersi (02.09.2026): sablonun "Yalnizca Baslik" layout'unda baslik placeholder'i
    y~0.3'ten y~2.5'e kadar uzuyor. Icerik y1.5'te basladigi icin kart/KPI kutulari basligin
    KUTUSUNUN icine giriyordu (gorsel olarak ust bant basligin altinda kaliyor). Cozum: icerik
    slaytlarinda baslik kutusu y0.30-1.30 arasina sabitlenir; icerik y1.45'ten sonra serbest.
    """
    for ph in sl.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            if idx == 0 and sl.slide_layout.name == "Yalnızca Başlık":
                ph.left, ph.top = Inches(0.6), Inches(0.30)
                ph.width, ph.height = Inches(12.05), Inches(1.00)
            return ph
    # K-14: placeholder yoksa basliksiz slayt uretilirdi (sessiz kayip) — cikar.
    sys.exit("PLACEHOLDER BULUNAMADI (idx=%d, layout '%s'): başlıksız slayt üretilmesin diye "
             "durduruldu. Metin: %s" % (idx, sl.slide_layout.name, text[:60]))


def tb(sl, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp=1.0):
    bx = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = bx.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, None, None, None)]
    first = True
    for txt, sz, bd, col in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.line_spacing = sp
        r = p.add_run(); r.text = txt; r.font.size = Pt(sz or 13); r.font.bold = bool(bd)
        r.font.name = "Calibri"; r.font.color.rgb = col or INK
    return bx


def rrect(sl, x, y, w, h, fill, line=None, rad=True, lw=1):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def circ(sl, x, y, d, bg, icon):
    c = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = bg; c.line.fill.background(); c.shadow.inherit = False
    ip = d * 0.55
    sl.shapes.add_picture(os.path.join(ICOW, icon + ".png"),
                          Inches(x + (d - ip) / 2), Inches(y + (d - ip) / 2), Inches(ip), Inches(ip))


def card(sl, x, y, w, h, top=None, ikon=None):
    """KART: KARE kenarli beyaz kutu + SOL dikey vurgu cubugu (+ opsiyonel ikon).

    ⚠ Tasarim gecmisi (02.09.2026): once yuvarlak kart + UST yatay serit denendi. Serit kare
    koseli oldugu icin yuvarlak kosenin disina tasti; ice alinca da kartin ustunde ayri bir
    cizgi gibi durdu ("kutular cizginin altina giriyor"). ÇOZUM: kart KARE kenar (rad=False),
    vurgu SOL dikeyde ve kartla tam ayni yukseklikte -> kose uyumsuzlugu MATEMATIKSEL OLARAK
    imkansiz. Yeni kart eklerken bu deseni bozma.
    """
    rrect(sl, x, y, w, h, WHITE, RGBColor(0xDC, 0xDC, 0xDC), rad=False)
    if top:
        rrect(sl, x, y, 0.075, h, top, rad=False)
    if ikon:
        sl.shapes.add_picture(os.path.join(ICOR, ikon + ".png"),
                              Inches(x + w - 0.52), Inches(y + 0.18), Inches(0.3), Inches(0.3))


def sig(sl):
    tb(sl, 9.3, 6.95, 3.4, 0.3, [("© Fikri Eren 2026", 9, False, MGREY)], align=PP_ALIGN.RIGHT)


def dipnot(sl, metin):
    """Kapsam notu — LOGONUN SAGINDA (x2.1, y6.2). Her rakamli slaytta ZORUNLU.

    ⚠ Yerlesim dersi (02.09.2026): dipnot y6.02'de tam genislikteydi ve slayt-ozel alt notlarla
    (y5.6-5.95, sarma sonrasi 6.1-6.5'e uzayan) UST USTE BINIYORDU. Artik logo bandi hizasinda,
    logodan sonra basliyor: slayt icerigi y6.1'e kadar serbest.
    """
    tb(sl, 2.1, 6.2, 10.45, 0.62, [(metin, 8, False, MGREY)], sp=1.05)


def kpi(sl, x, y, w, etiket, deger, alt, renk=RED, buyuk=34, ikon=None):
    """KPI kutusu.

    ⚠ Yerlesim dersi (02.09.2026): etiket 2 satira sardiginda (uzun etiket, ör. "KADRO · UC POS
    MAGAZASI (SEZONLUK DAHIL)") ikinci satir DEGERIN uzerine biniyordu. Bloklar arasi mesafe
    artirildi, kutu 1.75 -> 2.0 buyudu, etiket puntosu 10 -> 9. Etiketi kisa tut (<= 28 karakter):
    ikinci satir icin yer var ama tercih tek satir.
    """
    card(sl, x, y, w, 2.0, renk, ikon=ikon)
    tb(sl, x + 0.28, y + 0.18, w - 0.95, 0.42, [(etiket, 9, True, GREY)], sp=1.05)
    tb(sl, x + 0.28, y + 0.66, w - 0.45, 0.72, [(deger, buyuk, True, renk)])
    tb(sl, x + 0.28, y + 1.44, w - 0.45, 0.5, [(alt, 10, False, GREY)], sp=1.05)


def cift_bar(sl, x, y, w, h, kategoriler, s25, s26, etiket25="2025", etiket26="2026",
             yuzde_etiket=False, sifirdan=False):
    cd = CategoryChartData(); cd.categories = kategoriler
    cd.add_series(etiket25, s25); cd.add_series(etiket26, s26)
    ch = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                             Inches(w), Inches(h), cd).chart
    ch.has_title = False
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
    ch.font.size = Pt(10); ch.font.name = "Calibri"
    pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.number_format = '0.0"%"' if yuzde_etiket else '#,##0'
    pl.data_labels.font.size = Pt(9)
    ch.series[0].format.fill.solid(); ch.series[0].format.fill.fore_color.rgb = MGREY
    ch.series[1].format.fill.solid(); ch.series[1].format.fill.fore_color.rgb = RED
    if sifirdan:
        # ⚠ KESIK EKSEN YASAK: kadro grafiginde eksen 125'ten basliyordu, 139-153 farki
        # olcusuz buyuk gorunuyordu. Patron sunumunda "eksen kesik" itirazi acik hedef.
        ch.value_axis.minimum_scale = 0
    return ch
