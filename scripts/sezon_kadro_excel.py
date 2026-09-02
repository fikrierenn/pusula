# -*- coding: utf-8 -*-
"""Sezon kadrosu kohort analizi -> Excel (emitter).

Cekirdek (kohort tanimi / censoring / survival mantigi)
sorgular/2026-09-02-sezon-personel-kohort-magaza.sql dosyasindadir.
Bu script SADECE bicimlendirir -- SUM/CASE/join YAPMAZ
(.claude/rules/emitter-ayrimi.md). Oranlar Excel FORMULU olarak yazilir,
boylece hesap tek kaynakta (SQL) kalir ve kullanici hucreyi denetleyebilir.

Girdi : briefings/sezon-kadro-20260902/veri.json (cekirdek sorgu ciktisi)
Cikti : xlsx (Ozet · Sezonluk-Sube · Kadrolu-Sube · Sube-Reyon · Kadro-Gruplari · Magaza-Is-Hacmi ·
        Hafta-Sezonluk · Hafta-Kadrolu · Cozulme · Aylik-Alim · Kadro-Dagilim · Yontem)
        + native Excel grafikleri.

Kullanim:
    python scripts/sezon_kadro_excel.py \
        briefings/sezon-kadro-20260902/veri.json \
        briefings/sezon-kadro-20260902/sezon-kadro-seyri.xlsx
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

YUZDE = "0.0%"
BASLIK_DOLGU = PatternFill("solid", fgColor="E30622")   # BKM kirmizisi
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=10)
GRUP_DOLGU = PatternFill("solid", fgColor="F2F2F2")
INCE = Side(style="thin", color="D9D9D9")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
NOT_YAZI = Font(italic=True, size=9, color="666666")


def _basliklar(ws, kolonlar):
    """kolonlar: [(baslik, genislik, sayi_formati)]"""
    ws.append([k[0] for k in kolonlar])
    for i, (_, genislik, _fmt) in enumerate(kolonlar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = genislik
    for h in ws[1]:
        h.fill = BASLIK_DOLGU
        h.font = BASLIK_YAZI
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = KENAR
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _satir_bicim(ws, kolonlar):
    for satir in ws.iter_rows(min_row=2):
        for hucre, (_, _g, fmt) in zip(satir, kolonlar):
            hucre.border = KENAR
            if fmt:
                hucre.number_format = fmt


def _oran(pay_kol, bolen_kol, satir):
    """Bos bolene karsi guvenli Excel orani (kanit hucrede gorunur)."""
    return f'=IF({bolen_kol}{satir}=0,"",{pay_kol}{satir}/{bolen_kol}{satir})'


def sayfa_ozet(wb, veri):
    kolonlar = [
        ("Kadro tipi", 30, None), ("Kohort", 9, "0"), ("Alinan", 10, "0"),
        ("Bugun aktif", 12, "0"), ("Eylul oncesi kayip", 13, "0"),
        ("Eylul oncesi kayip %", 13, YUZDE),
        ("Risk 14g", 10, "0"), ("Kalan 14g", 10, "0"), ("14g tutunma", 12, YUZDE),
        ("Risk 30g", 10, "0"), ("Kalan 30g", 10, "0"), ("30g tutunma", 12, YUZDE),
        ("Ort. kalis (gun)", 13, "0"),
    ]
    ws = wb.create_sheet("Ozet")
    _basliklar(ws, kolonlar)
    for r in veri["ozet"]:
        s = ws.max_row + 1
        ws.append([
            r["kadro"], r["kohort"], r["alinan"], r["bugun_aktif"],
            r["eylul_oncesi_ayrilan"], _oran("E", "C", s),
            r["risk14"], r["kalan14"], _oran("H", "G", s),
            r["risk30"], r["kalan30"], _oran("K", "J", s),
            r["ort_kalis_gun"],
        ])
        if r["kadro"] == "SEZONLUK":
            for h in ws[s]:
                h.fill = GRUP_DOLGU
    _satir_bicim(ws, kolonlar)

    son = ws.max_row + 2
    ws.cell(son, 1, "Gri satirlar = gercek sezon personeli (Kadro='SEZONLUK'). "
                    "'Bugun aktif' yillar arasi kiyaslanamaz -- Yontem sayfasi.").font = NOT_YAZI

    # 14g tutunma: kadro tipi x kohort (sadece SEZONLUK + KADRO satirlari)
    gr = BarChart()
    gr.type = "col"
    gr.title = "14 gun tutunma - esit kidemde"
    gr.y_axis.numFmt = YUZDE
    gr.y_axis.title = "Tutunma"
    gr.height, gr.width = 8, 18
    gr.add_data(Reference(ws, min_col=9, min_row=1, max_row=7), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, max_col=2, min_row=2, max_row=7))
    ws.add_chart(gr, f"A{son + 2}")
    return ws


def sayfa_sezonluk_sube(wb, veri):
    kolonlar = [
        ("Grup", 18, None), ("Sube", 18, None), ("Kohort", 9, "0"),
        ("Alinan", 10, "0"), ("Bugun aktif", 12, "0"),
        ("Eylul oncesi kayip", 13, "0"), ("Eylul oncesi kayip %", 13, YUZDE),
        ("Risk 14g", 10, "0"), ("Kalan 14g", 10, "0"), ("14g tutunma", 12, YUZDE),
        ("Risk 30g", 10, "0"), ("Kalan 30g", 10, "0"), ("30g tutunma", 12, YUZDE),
    ]
    ws = wb.create_sheet("Sezonluk-Sube")
    _basliklar(ws, kolonlar)
    for r in veri["sezonluk_sube"]:
        s = ws.max_row + 1
        ws.append([
            r["grup"], r["sube"], r["kohort"], r["alinan"], r["bugun_aktif"],
            r["eylul_oncesi_ayrilan"], _oran("F", "D", s),
            r["risk14"], r["kalan14"], _oran("I", "H", s),
            r["risk30"], r["kalan30"], _oran("L", "K", s),
        ])
    _satir_bicim(ws, kolonlar)
    ws.auto_filter.ref = f"A1:M{ws.max_row}"
    return ws


def sayfa_kadrolu_sube(wb, veri):
    kolonlar = [
        ("Sube", 18, None), ("Kohort", 9, "0"), ("Alinan", 10, "0"),
        ("Bugun aktif", 12, "0"), ("Eylul oncesi kayip", 13, "0"),
        ("Eylul oncesi kayip %", 13, YUZDE),
        ("Risk 14g", 10, "0"), ("Kalan 14g", 10, "0"), ("14g tutunma", 12, YUZDE),
    ]
    ws = wb.create_sheet("Kadrolu-Sube")
    _basliklar(ws, kolonlar)
    for r in veri["kadrolu_sube"]:
        s = ws.max_row + 1
        ws.append([
            r["sube"], r["kohort"], r["alinan"], r["bugun_aktif"],
            r["eylul_oncesi_ayrilan"], _oran("E", "C", s),
            r["risk14"], r["kalan14"], _oran("H", "G", s),
        ])
    _satir_bicim(ws, kolonlar)
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    son = ws.max_row + 2
    ws.cell(son, 1, "Kadro <> SEZONLUK (kadrolu + stajyer/part-time/bos). "
                    "Bu yilin bozulmasi bu segmentte.").font = NOT_YAZI
    return ws


def sayfa_sube_reyon(wb, veri):
    kolonlar = [
        ("Sube", 18, None), ("Reyon / Birim", 22, None),
        ("2025 alinan", 12, "0"), ("2025 Eylul oncesi kayip", 14, "0"),
        ("2026 alinan", 12, "0"), ("2026 ayrilan", 12, "0"),
        ("2026 kayip %", 12, YUZDE),
    ]
    ws = wb.create_sheet("Sube-Reyon")
    _basliklar(ws, kolonlar)
    for r in veri["sezonluk_sube_reyon"]:
        s = ws.max_row + 1
        ws.append([
            r["sube"], r["departman"], r["a25"], r["eylul_oncesi25"],
            r["a26"], r["ayril26"], _oran("F", "E", s),
        ])
    _satir_bicim(ws, kolonlar)
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    son = ws.max_row + 2
    ws.cell(son, 1, "Yalniz SEZONLUK kadro. 1-3 kisilik satirlar desen isareti, olcum degil.").font = NOT_YAZI
    return ws


def sayfa_cozulme(wb, veri):
    kolonlar = [("Ay", 12, None), ("2024", 10, "0"), ("2025", 10, "0"), ("2026", 10, "0")]
    ws = wb.create_sheet("Cozulme")
    _basliklar(ws, kolonlar)
    aylar = []
    for r in veri["cozulme"]:
        if r["ay"] not in aylar:
            aylar.append(r["ay"])
    tablo = {(r["kohort"], r["ay"]): r["ayrilan"] for r in veri["cozulme"]}
    for ay in aylar:
        ws.append([ay, tablo.get((2024, ay)), tablo.get((2025, ay)), tablo.get((2026, ay))])
    _satir_bicim(ws, kolonlar)

    gr = LineChart()
    gr.title = "Sezon kohortunun cozulmesi (ayrilan kisi)"
    gr.y_axis.title = "Ayrilan"
    gr.height, gr.width = 8, 18
    gr.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=ws.max_row), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(gr, f"F2")
    son = ws.max_row + 2
    ws.cell(son, 1, "2026 Eylul ve sonrasi bos = henuz gerceklesmedi (kesim 02.09.2026).").font = NOT_YAZI
    return ws


def sayfa_aylik(wb, veri):
    kolonlar = [("Yil", 8, "0"), ("Ay", 8, "0"), ("Sezonluk alim", 14, "0")]
    ws = wb.create_sheet("Aylik-Alim")
    _basliklar(ws, kolonlar)
    for r in veri["aylik_alim"]:
        ws.append([r["yil"], r["ay"], r["sezonluk_alim"]])
    _satir_bicim(ws, kolonlar)

    gr = BarChart()
    gr.type = "col"
    gr.title = "SEZONLUK kadro alimi - ay bazli"
    gr.y_axis.title = "Kisi"
    gr.height, gr.width = 8, 20
    gr.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, max_col=2, min_row=2, max_row=ws.max_row))
    ws.add_chart(gr, "E2")
    return ws


def sayfa_kadro_gruplari(wb, veri):
    """Magaza kadrosu 4 ayri grup x 4 kesim (30.06 ve 31.08, iki yil).

    Toplam kolonlari Excel SUM formulu -> kullanici hucreyi acip denetleyebilir.
    """
    g = veri["kadro_gruplari"]
    GRUPLAR = ["SEZONLUK", "ENGELLI", "ETKINLIK", "DIGER KADROLU"]
    KESIM = [("d30Haz25", "30.06.25"), ("d30Haz26", "30.06.26"),
             ("d31Agu25", "31.08.25"), ("d31Agu26", "31.08.26")]

    kolonlar = [("Magaza", 14, None)]
    for _, etiket in KESIM:
        for grup in GRUPLAR:
            kolonlar.append((f"{etiket}\n{grup}", 11, "0"))
        kolonlar.append((f"{etiket}\nTOPLAM", 11, "0"))

    ws = wb.create_sheet("Kadro-Gruplari")
    _basliklar(ws, kolonlar)
    ws.row_dimensions[1].height = 40

    ilk = 2
    toplam_kol = []          # her kesimin TOPLAM kolon harfi
    for r in g["satirlar"]:
        s = ws.max_row + 1
        satir = [r["magaza"]]
        for anahtar, _e in KESIM:
            bas = len(satir) + 1                     # bu bloktaki ilk grup kolonu
            for grup in GRUPLAR:
                satir.append(r[grup][anahtar])
            b, e = get_column_letter(bas), get_column_letter(bas + 3)
            satir.append(f"=SUM({b}{s}:{e}{s})")
        ws.append(satir)
    son_veri = ws.max_row

    # kesim TOPLAM kolon harflerini bir kez hesapla (grafik + toplam satiri icin)
    for i, _k in enumerate(KESIM):
        toplam_kol.append(get_column_letter(2 + i * 5 + 4))

    t = son_veri + 1
    satir = ["MAGAZALAR"]
    for i in range(len(KESIM) * 5):
        kol = get_column_letter(2 + i)
        satir.append(f"=SUM({kol}{ilk}:{kol}{son_veri})")
    ws.append(satir)
    for hucre in ws[t]:
        hucre.fill = GRUP_DOLGU
        hucre.font = Font(bold=True)
    _satir_bicim(ws, kolonlar)

    gr = BarChart()
    gr.type = "col"
    gr.title = "Magaza kadrosu - kesim bazli toplam"
    gr.y_axis.title = "Kisi"
    gr.height, gr.width = 9, 20
    for kol in toplam_kol:
        idx = ws[f"{kol}1"].column
        gr.add_data(Reference(ws, min_col=idx, min_row=1, max_row=son_veri), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, min_row=ilk, max_row=son_veri))
    ws.add_chart(gr, f"A{t + 3}")

    n = t + 22
    ws.cell(n, 1, g["kapsam"]).font = NOT_YAZI
    for i, (grup, tanim) in enumerate(g["grup_tanimi"].items(), start=1):
        ws.cell(n + i, 1, f"{grup}: {tanim}").font = NOT_YAZI
    ws.cell(n + len(g["grup_tanimi"]) + 1, 1, g["not"]).font = NOT_YAZI
    return ws


def sayfa_is_hacmi(wb, veri):
    """Magaza is hacmi: fis + kalem + urun adedi + ciro, kadro ile karsilastirmali.

    Tum oranlar Excel FORMULU (degisim %, kisi-basi, sepet) -> hesap SQL'de kalir,
    kullanici hucreyi acip denetleyebilir.
    """
    h = veri["magaza_is_hacmi"]
    ADET = "#,##0"
    ADET2 = "#,##0.0"
    TL = '#,##0\\ "₺"'
    kolonlar = [
        ("Magaza", 14, None),
        ("Kadro 25", 9, "0"), ("Kadro 26", 9, "0"), ("Kadro Δ", 9, YUZDE),
        ("Fis 25", 11, ADET), ("Fis 26", 11, ADET), ("Fis Δ", 9, YUZDE),
        ("Kalem 25", 11, ADET), ("Kalem 26", 11, ADET), ("Kalem Δ", 9, YUZDE),
        ("Urun adedi 25", 13, ADET), ("Urun adedi 26", 13, ADET), ("Adet Δ", 9, YUZDE),
        ("Net ciro 25", 15, TL), ("Net ciro 26", 15, TL), ("Ciro Δ", 9, YUZDE),
        ("Adet/kisi 25", 12, ADET), ("Adet/kisi 26", 12, ADET), ("Adet/kisi Δ", 11, YUZDE),
        ("Ciro/kisi 25", 14, TL), ("Ciro/kisi 26", 14, TL),
        ("Sepet adet 25", 12, ADET2), ("Sepet adet 26", 12, ADET2),
        ("Sepet tutar 25", 13, TL), ("Sepet tutar 26", 13, TL),
    ]
    ws = wb.create_sheet("Magaza-Is-Hacmi")
    _basliklar(ws, kolonlar)

    def d(eski, yeni, s):
        return f'=IF({eski}{s}=0,"",{yeni}{s}/{eski}{s}-1)'

    ilk = 2
    for r in h["satirlar"]:
        s = ws.max_row + 1
        ws.append([
            r["magaza"],
            r["kadro25"], r["kadro26"], d("B", "C", s),
            r["fis25"], r["fis26"], d("E", "F", s),
            r["kalem25"], r["kalem26"], d("H", "I", s),
            r["adet25"], r["adet26"], d("K", "L", s),
            r["net25"], r["net26"], d("N", "O", s),
            f'=IF(B{s}=0,"",K{s}/B{s})', f'=IF(C{s}=0,"",L{s}/C{s})', d("Q", "R", s),
            f'=IF(B{s}=0,"",N{s}/B{s})', f'=IF(C{s}=0,"",O{s}/C{s})',
            f'=IF(E{s}=0,"",K{s}/E{s})', f'=IF(F{s}=0,"",L{s}/F{s})',
            f'=IF(E{s}=0,"",N{s}/E{s})', f'=IF(F{s}=0,"",O{s}/F{s})',
        ])
    son_veri = ws.max_row

    # TOPLAM satiri (Excel SUM -> kullanici dogrulayabilir)
    t = son_veri + 1
    ws.append([
        "UC MAGAZA",
        f"=SUM(B{ilk}:B{son_veri})", f"=SUM(C{ilk}:C{son_veri})", d("B", "C", t),
        f"=SUM(E{ilk}:E{son_veri})", f"=SUM(F{ilk}:F{son_veri})", d("E", "F", t),
        f"=SUM(H{ilk}:H{son_veri})", f"=SUM(I{ilk}:I{son_veri})", d("H", "I", t),
        f"=SUM(K{ilk}:K{son_veri})", f"=SUM(L{ilk}:L{son_veri})", d("K", "L", t),
        f"=SUM(N{ilk}:N{son_veri})", f"=SUM(O{ilk}:O{son_veri})", d("N", "O", t),
        f"=K{t}/B{t}", f"=L{t}/C{t}", d("Q", "R", t),
        f"=N{t}/B{t}", f"=O{t}/C{t}",
        f"=K{t}/E{t}", f"=L{t}/F{t}",
        f"=N{t}/E{t}", f"=O{t}/F{t}",
    ])
    for hucre in ws[t]:
        hucre.fill = GRUP_DOLGU
        hucre.font = Font(bold=True)
    _satir_bicim(ws, kolonlar)

    gr = BarChart()
    gr.type = "col"
    gr.title = "Urun adedi - magaza bazli (1 Tem - 31 Agu)"
    gr.y_axis.title = "Adet"
    gr.height, gr.width = 8, 18
    gr.add_data(Reference(ws, min_col=11, max_col=12, min_row=1, max_row=son_veri), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, min_row=ilk, max_row=son_veri))
    ws.add_chart(gr, f"A{t + 3}")

    gr2 = BarChart()
    gr2.type = "col"
    gr2.title = "Kisi basi urun adedi"
    gr2.y_axis.title = "Adet / kisi"
    gr2.height, gr2.width = 8, 18
    gr2.add_data(Reference(ws, min_col=17, max_col=18, min_row=1, max_row=son_veri), titles_from_data=True)
    gr2.set_categories(Reference(ws, min_col=1, min_row=ilk, max_row=son_veri))
    ws.add_chart(gr2, f"L{t + 3}")

    not_satir = t + 20
    ws.cell(not_satir, 1, h["pencere"]).font = NOT_YAZI
    ws.cell(not_satir + 1, 1, f"Kadro: {h['kadro_tanimi']} · Kaynak: {h['kaynak']}").font = NOT_YAZI
    ws.cell(not_satir + 2, 1, h["not"]).font = NOT_YAZI
    return ws


def sayfa_haftalik(wb, veri, segment):
    """Hafta hafta giris/cikis/sahada -- segment basina ayri sayfa (SEZONLUK vs KADROLU)."""
    h = veri["haftalik"]
    etiket = h["hafta_etiketi"]
    s25 = h["seriler"][f"2025_{segment}"]
    s26 = h["seriler"][f"2026_{segment}"]
    kolonlar = [
        ("Hafta", 8, "0"), ("Dilim basi", 12, None),
        ("2025 giren", 11, "0"), ("2025 cikan", 11, "0"), ("2025 sahada", 12, "0"),
        ("2026 giren", 11, "0"), ("2026 cikan", 11, "0"), ("2026 sahada", 12, "0"),
    ]
    ws = wb.create_sheet(f"Hafta-{'Sezonluk' if segment == 'SEZONLUK' else 'Kadrolu'}")
    _basliklar(ws, kolonlar)

    def al(seri, alan, i):
        d = seri[alan]
        return d[i] if i < len(d) else None

    for i, e in enumerate(etiket):
        ws.append([i + 1, e,
                   al(s25, "giris", i), al(s25, "cikis", i), al(s25, "sahada", i),
                   al(s26, "giris", i), al(s26, "cikis", i), al(s26, "sahada", i)])
    _satir_bicim(ws, kolonlar)
    n = ws.max_row

    gr = BarChart()
    gr.type = "col"
    gr.title = f"{segment} -- hafta hafta giris / cikis"
    gr.y_axis.title = "Kisi"
    gr.x_axis.title = "Sezon haftasi (1 Tem = 1. hafta)"
    gr.height, gr.width = 8, 22
    for kol in (3, 4, 6, 7):
        gr.add_data(Reference(ws, min_col=kol, min_row=1, max_row=n), titles_from_data=True)
    gr.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(gr, "J2")

    ln = LineChart()
    ln.title = f"{segment} -- gun sonu sahadaki kisi (kumulatif)"
    ln.y_axis.title = "Kisi"
    ln.height, ln.width = 8, 22
    for kol in (5, 8):
        ln.add_data(Reference(ws, min_col=kol, min_row=1, max_row=n), titles_from_data=True)
    ln.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(ln, "J20")

    son = n + 2
    ws.cell(son, 1, h["hafta_tanimi"]).font = NOT_YAZI
    ws.cell(son + 1, 1, h["not"]).font = NOT_YAZI
    ws.cell(son + 2, 1, f"{segment}: 2025 kohort {s25['n']} kisi (cerceve disi cikis "
                        f"{s25['cerceve_disi_cikis']}) · 2026 kohort {s26['n']} kisi "
                        f"(kesim 10. hafta)").font = NOT_YAZI
    return ws


def sayfa_kadro_dagilim(wb, veri):
    kolonlar = [
        ("Kadro tipi", 20, None), ("Toplam kayit", 13, "0"), ("Aktif", 10, "0"),
        ("Ilk giris", 13, None), ("Son giris", 13, None),
    ]
    ws = wb.create_sheet("Kadro-Dagilim")
    _basliklar(ws, kolonlar)
    for r in veri["kadro_dagilim"]:
        ws.append([r["kadro"], r["toplam"], r["aktif"], r["ilk_giris"], r["son_giris"]])
    _satir_bicim(ws, kolonlar)
    son = ws.max_row + 2
    ws.cell(son, 1, "SEZONLUK bayragi 01.08.2023'ten itibaren kullaniliyor -- 2023 oncesi kohort bu alanla izlenemez.").font = NOT_YAZI
    return ws


def sayfa_yontem(wb, veri, kisi_var=False):
    ws = wb.create_sheet("Yontem")
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 130
    ws["B1"] = f"{veri['meta']['baslik']} -- yontem ve sinirlar"
    ws["B1"].font = Font(bold=True, size=12)
    ws["B2"] = (f"Kesim {veri['meta']['kesim']} · kohort {veri['meta']['kohort_penceresi']} · "
                f"sezon bayragi {veri['meta']['sezon_bayragi']}")
    ws["B2"].font = NOT_YAZI
    satir = 4
    for i, not_ in enumerate(veri["yontem"], start=1):
        ws.cell(satir, 1, i).font = Font(bold=True)
        h = ws.cell(satir, 2, not_)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[satir].height = 30
        satir += 1
    satir += 1
    if kisi_var:
        uyari = ws.cell(satir - 1, 2,
                        "⚠ KISISEL VERI: bu dosyada Liste-* ve Gun-Gun-* sayfalari ad-soyad + personel no "
                        "icerir (KVKK). Yetkili kisi disinda paylasilmaz; patron/kurul sunumu icin "
                        "isimsiz surum kullanilmali.")
        uyari.font = Font(bold=True, size=10, color="A2382B")
        uyari.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[satir - 1].height = 30
        satir += 1
    ws.cell(satir, 2, f"Kaynak: {veri['meta']['kaynak']}").font = NOT_YAZI
    ws.cell(satir + 1, 2, f"Cekirdek SQL: {veri['meta']['cekirdek_sql']}").font = NOT_YAZI
    return ws


def _kisi_sayfalari(wb, kisi_json):
    """Kisi-duzeyi sayfalari ayni workbook'a ekler (kardes emitter'i yeniden kullanir).

    ⚠ KVKK: bu sayfalar ad-soyad + personel no icerir -> dosya paylasim kisitli olur.
    """
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from sezon_personel_listesi_excel import sayfa_gun_gun, sayfa_liste, _gun

    satirlar = json.loads(Path(kisi_json).read_text(encoding="utf-8"))
    yillar = sorted({r["Kohort"] for r in satirlar})
    son_gun = max(_gun(r["GirisTarihi"]) for r in satirlar)
    for yil in yillar:
        bitis = son_gun if yil == max(yillar) else date(yil, 9, 30)
        sayfa_gun_gun(wb, satirlar, yil, bitis)
    sayfa_liste(wb, satirlar)
    return len(satirlar)


def main(argv):
    if len(argv) not in (3, 4):
        print(__doc__)
        return 2
    girdi, cikti = Path(argv[1]), Path(argv[2])
    kisi_json = argv[3] if len(argv) == 4 else None
    veri = json.loads(girdi.read_text(encoding="utf-8"))

    wb = Workbook()
    wb.remove(wb.active)
    sayfa_ozet(wb, veri)
    sayfa_sezonluk_sube(wb, veri)
    sayfa_kadrolu_sube(wb, veri)
    sayfa_sube_reyon(wb, veri)
    sayfa_kadro_gruplari(wb, veri)
    sayfa_is_hacmi(wb, veri)
    sayfa_haftalik(wb, veri, "SEZONLUK")
    sayfa_haftalik(wb, veri, "KADROLU")
    sayfa_cozulme(wb, veri)
    sayfa_aylik(wb, veri)
    sayfa_kadro_dagilim(wb, veri)
    kisi_say = _kisi_sayfalari(wb, kisi_json) if kisi_json else 0
    sayfa_yontem(wb, veri, kisi_var=bool(kisi_json))

    cikti.parent.mkdir(parents=True, exist_ok=True)
    wb.save(cikti)
    print(f"yazildi: {cikti} ({len(wb.sheetnames)} sayfa: {', '.join(wb.sheetnames)})")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
