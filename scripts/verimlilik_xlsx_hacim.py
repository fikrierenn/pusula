# -*- coding: utf-8 -*-
"""Excel: is hacmi sayfalari — Magaza · Kategori · Aylik · Yillar · Oca-Agu."""
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from verimlilik_ortak import CARI, ONCEKI
from verimlilik_xlsx_ortak import (ADET, ADET1, BASLIK, BASLIK_YAZI, BOLUM_YAZI, GRI,
                                   INCE, KAT, KENAR, NOT_YAZI, TL, VURGU, YESIL_YAZI,
                                   YUZDE, _basliklar, _notlar, _yaz)

# ------------------------------------------------------------------ Magaza
def sayfa_magaza(wb, veri):
    ws = wb.create_sheet("Magaza")
    gun = veri["meta"]["gun"]
    kolonlar = [
        ("Magaza", 13, None),
        ("Kadro 2025", 9, ADET), ("Kadro 2026", 9, ADET), ("Kadro Δ%", 9, YUZDE),
        ("Urun adedi 2025", 13, ADET), ("Urun adedi 2026", 13, ADET), ("Adet Δ%", 9, YUZDE),
        ("Adet/kisi 2025", 11, ADET), ("Adet/kisi 2026", 11, ADET), ("Adet/kisi Δ%", 11, YUZDE),
        ("Adet/kisi/gun 2025", 12, ADET1), ("Adet/kisi/gun 2026", 12, ADET1),
        ("Ciro 2025 (KDV haric)", 15, TL), ("Ciro 2026 (KDV haric)", 15, TL), ("Ciro Δ%", 9, YUZDE),
        ("Ciro/kisi 2025", 13, TL), ("Ciro/kisi 2026", 13, TL), ("Ciro/kisi Δ%", 11, YUZDE),
        ("Is / kadro (kac kat)", 11, KAT),
    ]
    ws.cell(1, 1, "Magaza bazinda kisi basi is — okul-hizali pencere (%d gun, Sinav haric)" % gun).font = Font(bold=True, size=12)
    _basliklar(ws, kolonlar, satir=3)

    s = 4
    ilk = s
    for m in veri["magaza"]:
        ws.cell(s, 1, m["ad"]).border = KENAR
        ham = {2: m["kadro25"], 3: m["kadro26"], 5: m["adet25"], 6: m["adet26"],
               13: m["kdvharic25"], 14: m["kdvharic26"]}
        for kol, v in ham.items():
            c = ws.cell(s, kol, v)
            c.number_format = kolonlar[kol - 1][2]
            c.border = KENAR
        s += 1
    son = s - 1

    # TOPLAM satiri — ham kolonlar SUM
    ws.cell(s, 1, "TOPLAM").font = Font(bold=True)
    ws.cell(s, 1).fill = GRI
    ws.cell(s, 1).border = KENAR
    for kol in (2, 3, 5, 6, 13, 14):
        c = ws.cell(s, kol, "=SUM(%s%d:%s%d)" % (get_column_letter(kol), ilk, get_column_letter(kol), son))
        c.number_format = kolonlar[kol - 1][2]
        c.border = KENAR
        c.fill = GRI
        c.font = Font(bold=True)
    toplam = s

    # tureme kolonlari (formul) — hem magaza satirlari hem TOPLAM
    for r in list(range(ilk, son + 1)) + [toplam]:
        kalin = Font(bold=True) if r == toplam else None
        dolgu = GRI if r == toplam else None
        turemeler = {
            4: "=C%d/B%d-1" % (r, r),
            7: "=F%d/E%d-1" % (r, r),
            8: "=E%d/B%d" % (r, r),
            9: "=F%d/C%d" % (r, r),
            10: "=I%d/H%d-1" % (r, r),
            11: "=E%d/B%d/%d" % (r, r, gun),
            12: "=F%d/C%d/%d" % (r, r, gun),
            15: "=N%d/M%d-1" % (r, r),
            16: "=M%d/B%d" % (r, r),
            17: "=N%d/C%d" % (r, r),
            18: "=Q%d/P%d-1" % (r, r),
            19: "=G%d/D%d" % (r, r),
        }
        for kol, f in turemeler.items():
            c = ws.cell(r, kol, f)
            c.number_format = kolonlar[kol - 1][2]
            c.border = KENAR
            if dolgu:
                c.fill = dolgu
            if kalin:
                c.font = kalin
            if kol in (10, 18, 19):
                c.font = Font(bold=True, color="1F7A4D")

    # grafik: kisi basi urun adedi, magaza bazinda 2025 vs 2026
    g = BarChart()
    g.type = "col"
    g.title = "Kisi basi urun adedi — 2025 vs 2026"
    g.y_axis.title = "adet / kisi"
    g.height, g.width = 8, 16
    g.add_data(Reference(ws, min_col=8, max_col=9, min_row=3, max_row=son), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=ilk, max_row=son))
    ws.add_chart(g, "A%d" % (toplam + 3))

    s = toplam + 22
    _notlar(ws, [
        "Kadro = 31.08 itibariyla o magazada fiilen calisan TUM personel (sezonluk + kadrolu).",
        "'Is / kadro' = urun adedi buyumesi / kadro buyumesi. 1,0x'in uzeri: is kadrodan hizli buyudu.",
        "Ist. Yolu kadrosu en cok buyuyen magaza (50 -> 61) ama ise ragmen kisi basi adedi de artti.",
    ], s)
    return ws

# ------------------------------------------------------------------ Kategori
def sayfa_kategori(wb, veri):
    """Hangi kategori ne kadar buyudu — bolum kadro artisiyla ESLESTIRME sayfasi."""
    ws = wb.create_sheet("Kategori")
    # kategori -> hangi bolumun (reyonun) isi (BKM reyon/kategori eslesmesi)
    ESLES = {
        "Hazırlık Kitapları": "YARDIMCI KİTAP", "Kırtasiye": "KIRTASİYE", "Kitap": "KÜLTÜR",
        "Çocuk Kitabı": "ÇOCUK", "Oyuncak": "OYUNCAK", "Akademi": "AKADEMİ",
        "Hediyelik": "KIRTASİYE / OYUNCAK", "Gıda": "KAFE / GIDA",
    }
    kadro_delta = {b["bolum"]: b["kadrolu26"] - b["kadrolu25"] for b in veri["bolum"]}

    kolonlar = [("Kategori", 19, None), ("İlgili bölüm", 18, None),
                ("Adet 2025", 12, ADET), ("Adet 2026", 12, ADET), ("Adet Δ", 9, YUZDE),
                ("Ciro 2025 (M ₺)", 13, ADET1), ("Ciro 2026 (M ₺)", 13, ADET1), ("Ciro Δ", 9, YUZDE),
                ("Oca-Ağu adet 2025", 14, ADET), ("Oca-Ağu adet 2026", 14, ADET), ("Oca-Ağu adet Δ", 12, YUZDE),
                ("Oca-Ağu ciro Δ", 12, YUZDE),
                ("Bölüm kadro Δ", 11, "+0;-0;0")]
    ws.cell(1, 1, "Kategori büyümesi (okul-hizalı pencere VE Ocak-Ağustos) + o kategoriye bakan bölümün "
                  "kadro değişimi").font = Font(bold=True, size=12)
    _basliklar(ws, kolonlar, satir=3)

    s = 4
    ilk = s
    for k in veri["kategori"]:
        bolum = ESLES.get(k["kategori"], "—")
        ws.cell(s, 1, k["kategori"]).border = KENAR
        ws.cell(s, 2, bolum).border = KENAR
        for kol, v_ in ((3, k["adet25"]), (4, k["adet26"])):
            c = ws.cell(s, kol, v_); c.number_format = ADET; c.border = KENAR
        for kol, v_ in ((6, k["ciro25"] / 1e6), (7, k["ciro26"] / 1e6)):
            c = ws.cell(s, kol, v_); c.number_format = ADET1; c.border = KENAR
        for kol, f in ((5, "=D%d/C%d-1" % (s, s)), (8, "=G%d/F%d-1" % (s, s))):
            c = ws.cell(s, kol, f); c.number_format = YUZDE; c.border = KENAR
            c.font = Font(bold=True)
        for kol, v_ in ((9, k.get("oa_adet25")), (10, k.get("oa_adet26"))):
            c = ws.cell(s, kol, v_ if v_ is not None else "—")
            c.number_format = ADET
            c.border = KENAR
        for kol, f in ((11, "=IF(OR(I%d=\"—\",J%d=\"—\"),\"\",J%d/I%d-1)" % (s, s, s, s)),
                       (12, "=IF(OR(I%d=\"—\",J%d=\"—\"),\"\",%s)" % (s, s, "0"))):
            c = ws.cell(s, kol, f)
            c.number_format = YUZDE
            c.border = KENAR
        if k.get("oa_ciro25"):
            c = ws.cell(s, 12, (k["oa_ciro26"] / k["oa_ciro25"]) - 1)
            c.number_format = YUZDE
            c.border = KENAR
        kd = kadro_delta.get(bolum)
        c = ws.cell(s, 13, kd if kd is not None else "—")
        c.number_format = "+0;-0;0"
        c.border = KENAR
        if kd:
            c.fill = VURGU
        s += 1
    son = s - 1

    ws.cell(s, 1, "TOPLAM").font = Font(bold=True)
    ws.cell(s, 1).fill = GRI; ws.cell(s, 1).border = KENAR
    ws.cell(s, 2).fill = GRI; ws.cell(s, 2).border = KENAR
    for kol in (3, 4, 6, 7, 9, 10):
        c = ws.cell(s, kol, "=SUM(%s%d:%s%d)" % (get_column_letter(kol), ilk, get_column_letter(kol), son))
        c.number_format = ADET if kol in (3, 4, 9, 10) else ADET1
        c.border = KENAR; c.fill = GRI; c.font = Font(bold=True)
    for kol, f in ((5, "=D%d/C%d-1" % (s, s)), (8, "=G%d/F%d-1" % (s, s)),
                   (11, "=J%d/I%d-1" % (s, s))):
        c = ws.cell(s, kol, f); c.number_format = YUZDE; c.border = KENAR
        c.fill = GRI; c.font = Font(bold=True)

    s += 2
    _notlar(ws, [
        "IKI PENCERE: sol blok OKUL-HIZALI pencere (sezon kiyasi), sag blok OCAK-AGUSTOS kumulatif (yil geneli).",
        "ESLESME: kategori (urun) -> o urune bakan reyon (bolum). Kadro Δ o BOLUMUN kadrolu degisimidir.",
        "OKUNACAK: kadro artisi en hizli buyuyen kategorilere gitti (Hazirlik Kitaplari, Kirtasiye, Akademi);",
        "   en yavas buyuyen kategoride (Cocuk Kitabi) kadro AZALTILDI. Yani alim rastgele degil.",
        "Kapsam: uc POS magazasi, okul-hizali pencere, Sinav haric, KDV dahil, iadeler dusulmus.",
        "Toplam satiri yalniz bu tablodaki kategorilerin toplami (adet>=2000 esigi altindaki kuyruk haric).",
    ], s)
    return ws



# ------------------------------------------------------------------ Aylik (takvim ayi)
def sayfa_aylik(wb, veri):
    """Haz/Tem/Agu ay ay — Agustos'un neden zayif gorundugu (okul kaymasi) burada gorunur."""
    ws = wb.create_sheet("Aylik")
    kolonlar = [("Ay", 12, None), ("Adet 2025", 13, ADET), ("Adet 2026", 13, ADET), ("Adet Δ", 10, YUZDE),
                ("Ciro 2025 (M ₺)", 14, ADET1), ("Ciro 2026 (M ₺)", 14, ADET1), ("Ciro Δ", 10, YUZDE)]
    ws.cell(1, 1, "Aylık seyir — takvim ayı (üç POS mağazası, Sınav hariç, KDV dahil)").font = Font(bold=True, size=12)
    _basliklar(ws, kolonlar, satir=3)

    s = 4
    ilk = s
    for a in veri.get("aylik", []):
        ws.cell(s, 1, a["ad"]).border = KENAR
        for kol, v_ in ((2, a["adet25"]), (3, a["adet26"])):
            c = ws.cell(s, kol, v_); c.number_format = ADET; c.border = KENAR
        for kol, v_ in ((5, a["ciro25"] / 1e6), (6, a["ciro26"] / 1e6)):
            c = ws.cell(s, kol, v_); c.number_format = ADET1; c.border = KENAR
        for kol, f in ((4, "=C%d/B%d-1" % (s, s)), (7, "=F%d/E%d-1" % (s, s))):
            c = ws.cell(s, kol, f); c.number_format = YUZDE; c.border = KENAR; c.font = Font(bold=True)
        s += 1
    son = s - 1

    ws.cell(s, 1, "TOPLAM").font = Font(bold=True); ws.cell(s, 1).fill = GRI; ws.cell(s, 1).border = KENAR
    for kol in (2, 3, 5, 6):
        c = ws.cell(s, kol, "=SUM(%s%d:%s%d)" % (get_column_letter(kol), ilk, get_column_letter(kol), son))
        c.number_format = ADET if kol in (2, 3) else ADET1
        c.border = KENAR; c.fill = GRI; c.font = Font(bold=True)
    for kol, f in ((4, "=C%d/B%d-1" % (s, s)), (7, "=F%d/E%d-1" % (s, s))):
        c = ws.cell(s, kol, f); c.number_format = YUZDE; c.border = KENAR; c.fill = GRI; c.font = Font(bold=True)

    # --- KAYMA DUZELTMESI blogu
    k = veri.get("kayma")
    if k:
        s += 2
        ws.cell(s, 1, "OKUL KAYMASI DÜZELTMESİ").font = BOLUM_YAZI
        s += 1
        h = k["hizali_buyume"]
        satirlar = [
            ("Hizalı pencere (aynı talep gününe denk gelen günler)",
             "%s  vs  %s  (%d gün)" % (h["pencere_2025"], h["pencere_2026"], h["gun"])),
            ("Hizalı pencerede büyüme",
             "adet +%%%.1f · ciro +%%%.1f" % (h["adet"] * 100, h["ciro"] * 100)),
            ("Ağustos 2025 (gerçek)",
             "%s adet · %.1f M TL" % ("{:,.0f}".format(k["y25_agu_tam"]["adet"]).replace(",", "."),
                                      k["y25_agu_tam"]["ciro"] / 1e6)),
            ("Ağustos 2026 (gerçek)",
             "%s adet · %.1f M TL  (adet +%%%.1f)"
             % ("{:,.0f}".format(k["y26_agu_tam"]["adet"]).replace(",", "."),
                k["y26_agu_tam"]["ciro"] / 1e6,
                (k["y26_agu_tam"]["adet"] / k["y25_agu_tam"]["adet"] - 1) * 100)),
            ("Ağustos 2026 — KAYMA OLMASAYDI (tahmin)",
             "%s adet · %.1f M TL"
             % ("{:,.0f}".format(k["agustos_kaymasiz_tahmin"]["adet"]).replace(",", "."),
                k["agustos_kaymasiz_tahmin"]["ciro"] / 1e6)),
            ("EYLÜL'E KAYAN (tahmin)",
             "%s adet · %.1f M TL"
             % ("{:,.0f}".format(k["eylule_kayan"]["adet"]).replace(",", "."),
                k["eylule_kayan"]["ciro"] / 1e6)),
            ("Okul öncesi dalga — 2025 gerçekleşen (%s)" % k["eylul_dalga_beklentisi"]["pencere_2025"],
             "%s adet · %.1f M TL"
             % ("{:,.0f}".format(k["eylul_dalga_beklentisi"]["adet_2025"]).replace(",", "."),
                k["eylul_dalga_beklentisi"]["ciro_2025"] / 1e6)),
            ("Okul öncesi dalga — 2026 beklenen (%s)" % k["eylul_dalga_beklentisi"]["pencere_2026"],
             "%s adet · %.1f M TL"
             % ("{:,.0f}".format(k["eylul_dalga_beklentisi"]["adet_2026_tahmin"]).replace(",", "."),
                k["eylul_dalga_beklentisi"]["ciro_2026_tahmin"] / 1e6)),
        ]
        for etiket, deger in satirlar:
            a = ws.cell(s, 1, etiket); a.border = KENAR
            b = ws.cell(s, 2, deger); b.border = KENAR
            ws.merge_cells(start_row=s, start_column=2, end_row=s, end_column=7)
            if "KAYMA OLMASAYDI" in etiket or "KAYAN" in etiket or "beklenen" in etiket:
                a.font = Font(bold=True); b.font = Font(bold=True, color="A6001A")
                a.fill = VURGU; b.fill = VURGU
            s += 1
        s += 1
        ws.cell(s, 1, "Yöntem: " + k["yontem"]).font = NOT_YAZI
        ws.cell(s, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=7)
        ws.row_dimensions[s].height = 42
        s += 1

    s += 2
    _notlar(ws, [
        "AGUSTOS NEDEN ZAYIF GORUNUYOR: okullar 2025'te 8 Eylul, 2026'da 14 Eylul acildi (6 gun kayma).",
        "   2025'in son-Agustos alis dalgasi 2026'da EYLUL'e kaydi -> takvim ayi kiyasinda Agustos dusuk cikar.",
        "   Temmuz +%36,3 adet, Agustos +%5,5 adet: fark talep kaybi degil, TAKVIM.",
        "Dogru kiyas okul-acilisina hizali penceredir (Ozet sayfasi): adet +%34,9 · ciro +%70,7.",
        "Bu sayfa 'ay ay ne oldu' sorusunun cevabidir; kadro kiyasinda hizali pencere kullanilir.",
    ], s)
    return ws


# ------------------------------------------------------------------ Yillar
def sayfa_yillar(wb, veri):
    ws = wb.create_sheet("Yillar")
    kolonlar = [("Yil", 8, None), ("Kadrolu (31.08)", 13, ADET), ("Urun adedi (Oca-Agu)", 16, ADET),
                ("Adet / kisi", 12, ADET), ("Onceki yila gore", 13, YUZDE)]
    ws.cell(1, 1, "Kisi basi is — 4 yillik trend (uc POS magazasi, Ocak-Agustos kumulatif, Sinav haric)").font = Font(bold=True, size=12)
    _basliklar(ws, kolonlar, satir=3)

    s = 4
    ilk = s
    for y in veri["yillar"]:
        ws.cell(s, 1, y["yil"]).border = KENAR
        for kol, v in ((2, y["kadrolu"]), (3, y["adet"])):
            c = ws.cell(s, kol, v)
            c.number_format = ADET
            c.border = KENAR
        c = ws.cell(s, 4, "=C%d/B%d" % (s, s))
        c.number_format = ADET
        c.border = KENAR
        c.font = Font(bold=True)
        if s > ilk:
            d = ws.cell(s, 5, "=D%d/D%d-1" % (s, s - 1))
            d.number_format = YUZDE
            d.border = KENAR
        else:
            ws.cell(s, 5, "—").border = KENAR
        s += 1
    son = s - 1

    g = LineChart()
    g.title = "Kisi basi urun adedi — 4 yillik trend"
    g.y_axis.title = "adet / kisi"
    g.height, g.width = 8, 16
    g.add_data(Reference(ws, min_col=4, min_row=3, max_row=son), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=ilk, max_row=son))
    ws.add_chart(g, "G3")

    s += 1
    _notlar(ws, [
        "2024 ATLAMASI: kadrolu 60 -> 91, adet yalniz +%11 -> kisi basi is -%27. Kadro sismesi 2024'te oldu.",
        "AMA 2023 verimliligi 'norm' DEGIL: 2023'te FSM kasada 0 kisi, Ozluce kasada 1 kisi vardi (eksik kadroyla calisma).",
        "2024 -> 2026: kisi basi is +%28,8 toparlanma. Bu yil kadro +15 kisi buyurken kisi basi is de artti.",
        "Bu sayfada kadro yalniz KADROLU (sezonluk haric) — yillar arasi sezonluk tahliye zamanlamasi kiyasi bozuyor.",
    ], s)
    return ws


# ------------------------------------------------------------------ Oca-Agu (itiraz cevabi)
def sayfa_oca_agu(wb, veri):
    ws = wb.create_sheet("Oca-Agu")
    oa = veri["ocak_agustos"]
    kolonlar = [("Kanal", 24, None), ("Adet 2025", 14, ADET), ("Adet 2026", 14, ADET), ("Adet Δ%", 10, YUZDE),
                ("Ciro 2025 (KDV dahil)", 17, TL), ("Ciro 2026 (KDV dahil)", 17, TL), ("Ciro Δ%", 10, YUZDE)]
    ws.cell(1, 1, "\"Buyume kurumsaldan geldi\" itirazinin cevabi — Ocak-Agustos, uc POS magazasi").font = Font(bold=True, size=12)
    _basliklar(ws, kolonlar, satir=3)

    s = 4
    for etiket, blok in (("Magaza (perakende raf)", oa["magaza"]), ("Sinav Okullari (kurumsal)", oa["sinav"])):
        ws.cell(s, 1, etiket).border = KENAR
        for kol, v in ((2, blok["adet25"]), (3, blok["adet26"]),
                       (5, blok["kdvdahil25"]), (6, blok["kdvdahil26"])):
            c = ws.cell(s, kol, v)
            c.number_format = kolonlar[kol - 1][2]
            c.border = KENAR
        for kol, f in ((4, "=C%d/B%d-1" % (s, s)), (7, "=F%d/E%d-1" % (s, s))):
            c = ws.cell(s, kol, f)
            c.number_format = YUZDE
            c.border = KENAR
        s += 1
    ilk, son = 4, s - 1

    ws.cell(s, 1, "TOPLAM").font = Font(bold=True)
    ws.cell(s, 1).fill = GRI
    ws.cell(s, 1).border = KENAR
    for kol in (2, 3, 5, 6):
        c = ws.cell(s, kol, "=SUM(%s%d:%s%d)" % (get_column_letter(kol), ilk, get_column_letter(kol), son))
        c.number_format = kolonlar[kol - 1][2]
        c.border = KENAR
        c.fill = GRI
        c.font = Font(bold=True)
    for kol, f in ((4, "=C%d/B%d-1" % (s, s)), (7, "=F%d/E%d-1" % (s, s))):
        c = ws.cell(s, kol, f)
        c.number_format = YUZDE
        c.border = KENAR
        c.fill = GRI
        c.font = Font(bold=True)
    s += 2

    _notlar(ws, [
        "Sinav Okullari KUCULDU (adet -%35, ciro -%22). Buyumenin tamami magaza rafindan geldi.",
        "Magaza tarafi kurumsal dususu de kapatti: toplam yine buyudu.",
        "Sinav = Kategori3 'Sinav Okullari' + 'Sinav Kiyafet'; ayni POS belgesi icinde geldigi icin AYIKLANMASI zorunlu.",
    ], s)
    return ws
