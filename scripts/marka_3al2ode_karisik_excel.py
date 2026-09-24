# -*- coding: utf-8 -*-
"""3 AL 2 ÖDE tanımlı kitabı OLAN ama 3 AL 2 ÖDE DIŞINDA da tanımı bulunan markalar -> Excel.

GMY isteği 21.09.2026: "3 al 2 öde kampanyasında kitabı tanımlı olan ama aynı zamanda
3 al 2 öde dışında tanım bulunan markaların listesini almak lazım".

TANIMIN YERİ (ölçüldü 21.09.2026):
  · Kampanya üyeliği ÜRÜN KARTINDA: dbo.urn.kod4ID -> dbo.urnkod4.kod4Ad
    kod4ID = 1 -> '3 AL 2 ÖDE' (430.749 ürün, urnTip=0). Kartta alan adı "Reyon"
    görünür ama içerik kampanyadır (sema codes:urn.kod4ID).
  · Marka seviyesinde flag VAR ama BOŞ: dbo.urnMrk.mrkUcalikiOde -> 10.306 markanın
    10.306'sında NULL. Kullanılmıyor; marka bazlı tanım ürün kartından TÜRETİLİR.
  · Kasa tarafı: EncoreMerkez.dbo.Campaign Id=1 '3AL2ÖDE' (3 al -> 2 öde), IsActive=0,
    son kampanyalı satış kalemi 06.02.2026. Etiket ile kasa tutuyor ama nüfus küçük
    (Oca-Şub 2026'da 17 kalem, 17'si de kod4Ad='3 AL 2 ÖDE').

KAPSAM KARARLARI (raporun Kapsam sayfasında da yazılı):
  · KİTAP = bkm.UrunBilgi.KatAna LIKE '%Kitap%' (DİKKAT: 'Kitap Aksesuarları' HARİÇ --
    adında Kitap geçiyor ama kitap değil) + 'Eğitim - Sınavlara Hazırlık - Okula Yardımcı'.
    --tum-urun ile kategori süzgeci kalkar (kırtasiye/oyuncak/hediyelik dahil).
  · "3 al 2 öde dışında TANIM" = kod4ID NOT IN (0,1). Tanımsız (kod4ID=0) ayrı kolonda
    raporlanır, "diğer tanım" SAYILMAZ -- tanımsızlık bir tanım değildir.
  · urnTip=0 (normal ürün; 1=gider/hizmet, 2=demirbaş hariç).
  · ÜRÜN DURUMU (GMY isteği 21.09.2026 "pasif veya tükendi olanları devre dışı bırak"):
    VARSAYILAN olarak yalnız AKTİF ürün -> dbo.urn.kod1ID = 1 (lookup dbo.urnKod1:
    0=Pasif, 1=Aktif, 2=Tükendi). --tum-durum bayrağı süzgeci kaldırır.
    Ölçüldü: katalogun yalnız %39'u aktif (aktif 327.083 / pasif 355.400 / tükendi 156.925).
    DİKKAT: `urn.urnDurum` AYRI bir kolondur (kayıt-düzeyi bayrak, 813K'sı 1) ve ürün
    durumu DEĞİLDİR; 'pasif/tükendi' sorusu kod1ID'den yanıtlanır. Kesişim ölçüldü:
    kod1=Aktif olup urnDurum=0 olan 112 ürün var, bunlar dahil edilir (kod1 esas alınır).

pyodbc kullanılıyor (pymssql DEĞİL): DerinSIS varchar kolonları CP1254, pymssql
Türkçe'yi bozuyor (coding-discipline § rapor scripti).

Kullanım:
    python scripts/marka_3al2ode_karisik_excel.py [--tum-urun] [--cikti yol.xlsx]
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import os
import re
import sys

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def env_oku(yol: str) -> dict:
    """`.env` dosyasını okur. Kimlik YALNIZ burada durur (dört sözleşme: tek kimlik yolu)."""
    if not os.path.exists(yol):
        sys.exit("`.env` bulunamadi: " + yol)
    env = {}
    with open(yol, encoding="utf-8") as f:
        for ln in f:
            if ln.lstrip().startswith("#"):
                continue
            m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.+?)\s*$", ln)
            if m:
                env[m.group(1)] = m.group(2).strip().strip('"')
    return env


def baglan(env: dict) -> pyodbc.Connection:
    host, port = env["MSSQL_HOST"], env.get("MSSQL_PORT", "1433")
    # ODBC connection-string'e env degeri gomulurken whitelist guard (injection)
    if not re.fullmatch(r"[A-Za-z0-9._\-]+", host) or not re.fullmatch(r"\d+", port):
        sys.exit("Gecersiz MSSQL_HOST/MSSQL_PORT (.env)")
    cn = pyodbc.connect(
        "Driver={ODBC Driver 18 for SQL Server};"
        f"Server={host},{port};Database=DerinSISBkm;"
        f"UID={env['MSSQL_USER']};PWD={env['MSSQL_PASSWORD']};"
        "TrustServerCertificate=yes;Timeout=30",
        timeout=30,
    )
    cn.timeout = 900
    return cn


# Kategori süzgeci parametre DEĞİL, iki sabit metin: SQL'e kullanıcı girdisi gömülmez.
KITAP_SUZGEC = (
    "AND (b.KatAna LIKE N'%Kitap%' "
    "OR b.KatAna = N'Eğitim - Sınavlara Hazırlık - Okula Yardımcı') "
    "AND b.KatAna <> N'Kitap Aksesuarları' "
)
TUM_SUZGEC = ""

# Ürün durumu: kod1ID 0=Pasif · 1=Aktif · 2=Tükendi (lookup dbo.urnKod1).
AKTIF_SUZGEC = "AND u.kod1ID = 1 "
TUM_DURUM_SUZGEC = ""

SQL_OZET = """
WITH kapsam AS (
    SELECT u.urnMrkID, u.kod4ID
    FROM dbo.urn u WITH (NOLOCK)
    JOIN bkm.UrunBilgi b ON b.stkID = u.stkID
    WHERE u.urnTip = 0 {DURUM}{SUZGEC}
),
marka AS (
    SELECT k.urnMrkID,
           SUM(CASE WHEN k.kod4ID = 1 THEN 1 ELSE 0 END)          AS UcAlIkiOde,
           SUM(CASE WHEN k.kod4ID NOT IN (0,1) THEN 1 ELSE 0 END) AS DigerTanimli,
           SUM(CASE WHEN k.kod4ID = 0 THEN 1 ELSE 0 END)          AS Tanimsiz,
           COUNT(*)                                               AS Toplam
    FROM kapsam k
    GROUP BY k.urnMrkID
)
SELECT m.urnMrkID                          AS MarkaID,
       ISNULL(mk.mrkAd, '(marka yok)')     AS Marka,
       m.UcAlIkiOde,
       m.DigerTanimli,
       m.Tanimsiz,
       m.Toplam,
       CONVERT(decimal(5,1), 100.0 * m.DigerTanimli / NULLIF(m.Toplam,0)) AS DigerTanimliYuzde,
       STUFF((
           SELECT ', ' + x.kod4Ad + ' (' + CONVERT(varchar(12), x.Adet) + ')'
           FROM (
               SELECT k2.kod4Ad, COUNT(*) AS Adet
               FROM kapsam c2
               JOIN dbo.urnkod4 k2 ON k2.kod4ID = c2.kod4ID
               WHERE c2.urnMrkID = m.urnMrkID AND c2.kod4ID NOT IN (0,1)
               GROUP BY k2.kod4Ad
           ) x
           ORDER BY x.Adet DESC
           FOR XML PATH(''), TYPE).value('.', 'nvarchar(max)'), 1, 2, '') AS DigerTanimlar
FROM marka m
LEFT JOIN dbo.urnMrk mk ON mk.mrkID = m.urnMrkID
WHERE m.UcAlIkiOde > 0 AND m.DigerTanimli > 0
ORDER BY m.DigerTanimli DESC, m.Toplam DESC
"""

SQL_DETAY = """
WITH kapsam AS (
    SELECT u.urnMrkID, u.kod4ID
    FROM dbo.urn u WITH (NOLOCK)
    JOIN bkm.UrunBilgi b ON b.stkID = u.stkID
    WHERE u.urnTip = 0 {DURUM}{SUZGEC}
),
karisik AS (
    SELECT k.urnMrkID
    FROM kapsam k
    GROUP BY k.urnMrkID
    HAVING SUM(CASE WHEN k.kod4ID = 1 THEN 1 ELSE 0 END) > 0
       AND SUM(CASE WHEN k.kod4ID NOT IN (0,1) THEN 1 ELSE 0 END) > 0
)
SELECT ISNULL(mk.mrkAd, '(marka yok)') AS Marka,
       k4.kod4Ad                       AS Tanim,
       COUNT(*)                        AS Adet
FROM kapsam c
JOIN karisik kr ON kr.urnMrkID = c.urnMrkID
JOIN dbo.urnkod4 k4 ON k4.kod4ID = c.kod4ID
LEFT JOIN dbo.urnMrk mk ON mk.mrkID = c.urnMrkID
GROUP BY mk.mrkAd, k4.kod4Ad
ORDER BY mk.mrkAd, COUNT(*) DESC
"""

BASLIK_DOLGU = PatternFill("solid", fgColor="E30622")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def sayfa_yaz(ws, basliklar, satirlar, genislik) -> None:
    ws.append(basliklar)
    for h in range(1, len(basliklar) + 1):
        c = ws.cell(row=1, column=h)
        c.fill = BASLIK_DOLGU
        c.font = BASLIK_YAZI
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for s in satirlar:
        ws.append(list(s))
    for i, g in enumerate(genislik, start=1):
        ws.column_dimensions[get_column_letter(i)].width = g
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(basliklar))}{len(satirlar) + 1}"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--tum-urun", action="store_true",
                    help="kategori süzgecini kaldır (kırtasiye/oyuncak/hediyelik dahil)")
    ap.add_argument("--tum-durum", action="store_true",
                    help="pasif + tükendi ürünleri de dahil et (varsayılan: yalnız AKTİF)")
    ap.add_argument("--cikti", default=None)
    a = ap.parse_args()

    suzgec = TUM_SUZGEC if a.tum_urun else KITAP_SUZGEC
    durum = TUM_DURUM_SUZGEC if a.tum_durum else AKTIF_SUZGEC
    kapsam_ad = (
        ("TÜM ÜRÜNLER" if a.tum_urun else "YALNIZ KİTAP KATEGORİLERİ")
        + (" · TÜM DURUMLAR (pasif+tükendi dahil)" if a.tum_durum
           else " · YALNIZ AKTİF ÜRÜN")
    )

    env = env_oku(os.path.join(KOK, ".env"))
    cn = baglan(env)
    try:
        with cn.cursor() as cur:
            cur.execute(SQL_OZET.format(SUZGEC=suzgec, DURUM=durum))
            ozet = cur.fetchall()
            cur.execute(SQL_DETAY.format(SUZGEC=suzgec, DURUM=durum))
            detay = cur.fetchall()
    finally:
        cn.close()

    if not ozet:
        print("UYARI: sonuç BOŞ — nüfus sıfır. 'ihlal yok' DEĞİL, 'bakamadım' demektir.")
        return 2

    wb = Workbook()
    ws = wb.active
    ws.title = "Karışık Markalar"
    sayfa_yaz(
        ws,
        ["Marka ID", "Marka", "3 AL 2 ÖDE", "Diğer Tanımlı", "Tanımsız", "Toplam",
         "Diğer Tanımlı %", "Diğer Tanımlar (adet)"],
        [tuple(r) for r in ozet],
        [10, 42, 12, 13, 11, 10, 14, 70],
    )

    ws2 = wb.create_sheet("Marka x Tanım")
    sayfa_yaz(ws2, ["Marka", "Tanım", "Adet"], [tuple(r) for r in detay], [42, 28, 10])

    ws3 = wb.create_sheet("Kapsam ve Yöntem")
    notlar = [
        ("Üretim", dt.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Kapsam", kapsam_ad),
        ("Tanımın yeri", "dbo.urn.kod4ID -> dbo.urnkod4.kod4Ad (kod4ID=1 -> '3 AL 2 ÖDE')"),
        ("Kart alan adı", "DerinSIS ürün kartında 'Reyon' yazar, içerik kampanyadır"),
        ("Marka flag'i",
         "dbo.urnMrk.mrkUcalikiOde kolonu VAR ama 10.306/10.306 NULL -> kullanılmıyor"),
        ("Ürün süzgeci", "urnTip=0 (gider/hizmet ve demirbaş hariç)"),
        ("Ürün durumu",
         ("TÜM DURUMLAR — pasif ve tükendi DAHİL" if a.tum_durum else
          "YALNIZ AKTİF (dbo.urn.kod1ID=1). Pasif (0) ve Tükendi (2) HARİÇ. "
          "Katalogun yalnız %39'u aktif.")),
        ("Durum kolonu uyarısı",
         "urn.urnDurum AYRI bir kayıt bayrağıdır, ürün durumu değildir. "
         "kod1=Aktif olup urnDurum=0 olan 112 ürün dahildir (kod1 esas)."),
        ("Kitap tanımı",
         "KatAna LIKE '%Kitap%' + 'Eğitim - Sınavlara Hazırlık - Okula Yardımcı'; "
         "'Kitap Aksesuarları' HARİÇ (adında Kitap geçiyor ama kitap değil)"),
        ("Diğer tanım", "kod4ID NOT IN (0,1). Tanımsız (0) ayrı kolon — tanım sayılmaz"),
        ("Kasa karşılığı",
         "EncoreMerkez Campaign Id=1 '3AL2ÖDE' (3 al 2 öde), IsActive=0; "
         "son kampanyalı satış kalemi 06.02.2026"),
        ("SINIR",
         "Bu liste ürün kartının ANLIK etiketidir, geçmiş değil. Geçmiş için "
         "bkm.urnkod4log. Katalog 06.05.2026 ve 12.05.2026'da toplu yeniden "
         "etiketlendi — o tarihten öncesiyle kıyas kırılır."),
        ("Arşiv SQL", "sorgular/2026-09-21-marka-3al2ode-karisik-tanim.sql"),
    ]
    sayfa_yaz(ws3, ["Başlık", "Açıklama"], notlar, [22, 110])
    for r in range(2, len(notlar) + 2):
        ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    cikti = a.cikti or os.path.join(
        os.environ.get("TEMP", "."),
        f"marka-3al2ode-karisik-{'tum' if a.tum_urun else 'kitap'}-"
        f"{'tumdurum' if a.tum_durum else 'aktif'}-"
        f"{dt.datetime.now():%Y%m%d-%H%M}.xlsx",
    )
    wb.save(cikti)
    print("Kapsam       : " + kapsam_ad)
    print(f"Karışık marka: {len(ozet)}")
    print(f"Detay satırı : {len(detay)}")
    print("Dosya        : " + cikti)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
