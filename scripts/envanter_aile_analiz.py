# -*- coding: utf-8 -*-
"""
ENVANTER FAZLA — AİLE-DÜZELTMELİ (İkame Confound) analiz raporu.

Problem: Per-SKU fazla-stok, İKAME edilebilir emtia ailelerinde (A4 kağıt, puzzle, markör)
  YANILTICI şişer — talep rafta ne varsa ona kayar → tek SKU "61 ay stok" görünür ama AİLE sağlıklı.
  Farklılaşmış üründe (karakter oyuncağı, belirli kitap) ise per-SKU fazla GERÇEKTİR.
Çözüm: Fazlayı İKİ grain'de ölç — (a) per-SKU, (b) AİLE (leaf kategori). Fark = "phantom"
  (aile-içi ikamenin per-SKU'yu şişirdiği kısım). Doğru serbest-kalabilir para = Σ aile_excess.

Leaf grain = en-derin-boş-olmayan kategori: COALESCE(Kat3, Kat2, Kat1).
  Kağıtta Kat3 doludur ("Fotokopi Kağıtları" ≠ "Gramajlı" ≠ "Yedek") → ikame sadece gerçek eşdeğerler arası.
  Makas/Cetvel'de Kat3 boş → Kat2 leaf olur.
Kapsam: Kat3ID IN (10,12,16) = Hediyelik/Kırtasiye/Oyuncak (kitap hariç); stok>0 VEYA son12 satış>0.

Model:
  stok    = şube (stokSon_vw mekan 1/4477/4478) + depo (stok_adres_palet_vw adrsAlanTipID 0/1)
  satış12 = irsHrk ehTip 1,4,100 (şube+POS) son 12 ay
  hedef   = 3-AY kapsama = satış12 / 4   (aynı kural per-SKU + aile → phantom target-seçiminden bağımsız)
  sku_excess = max(0, stok − hedef);   aile_excess = max(0, aile_stok − aile_satış12/4)
  ₺        = adet × birim_maliyet (fatAyr son alış faturası > devir ehTip=99 fallback)
Sınıf: aile_ay≤6 & phantom_oran≥%50 → İKAME (aile sağlıklı, per-SKU yanıltıcı; gerçek = aile_excess)
       değilse & aile_excess anlamlı → GERÇEK FAZLA (aile kendisi şişkin — makas/cetvel buraya düşer)
       değilse → SAĞLIKLI

Çıktı: briefings/envanter-aile-analiz-YYYYMMDD.xlsx  (Özet · Aile Analizi · Ürün Detay)
Kullanım: python scripts/envanter_aile_analiz.py
"""
import os, sys, re, datetime
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import pyodbc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEDEF_AY = 3          # 3-ay kapsama hedefi (satış12 / (12/3) = satış12/4)
IKAME_AILE_AY = 6.0   # aile ≤ bu ay → aile sağlıklı (ikame adayı)
IKAME_PHANTOM = 0.50  # phantom/sku_excess ≥ bu → ikame (aile per-SKU fazlayı yutuyor)
MAT_AILE = 25000      # aile_excess ₺ < bu → "SAĞLIKLI" (anlamsız küçük)

# ---- .env + bağlantı (pyodbc — Türkçe varchar CP1254 doğru decode) ----
ENV = {}
with open(os.path.join(os.path.dirname(__file__), "..", ".env"), encoding="utf-8") as f:
    for ln in f:
        m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.+?)\s*$", ln)
        if m and not ln.lstrip().startswith("#"):
            ENV[m.group(1)] = m.group(2).strip().strip('"')
host = ENV["MSSQL_HOST"]; port = ENV.get("MSSQL_PORT", "1433")
if not re.fullmatch(r"[A-Za-z0-9._\-]+", host) or not re.fullmatch(r"\d+", port):
    sys.exit("Geçersiz host/port (.env)")
cn = pyodbc.connect(
    f"Driver={{ODBC Driver 18 for SQL Server}};Server={host},{port};Database=DerinSISBkm;"
    f"UID={ENV['MSSQL_USER']};PWD={ENV['MSSQL_PASSWORD']};TrustServerCertificate=yes;Timeout=30", timeout=30)
cn.timeout = 900
cur = cn.cursor()

t = datetime.date.today().replace(day=1)
_y12 = t.year - 1
T_S12 = f"{_y12:04d}{t.month:02d}01"   # son 12 ay başı
print(f"Son-12-ay satış penceresi: [{T_S12}, bugün) · leaf = COALESCE(Kat3,Kat2,Kat1)", flush=True)

# ---- 1) UNIVERSE + kategori ağacı (Kat3ID 10/12/16) ----
print("Master + kategori ağacı...", flush=True)
master = {}   # sid -> dict(ad, katana, kat1, kat2, kat3, leaf, marka)
cur.execute("""
    SELECT u.stkID, u.stkAd, u.KatAna, u.Kat1, u.Kat2, u.Kat3, u.mrkAd
    FROM bkm.UrunBilgi u WITH(NOLOCK)
    WHERE u.Kat3ID IN (10,12,16)""")
for sid, ad, ka, k1, k2, k3, mrk in cur.fetchall():
    k1 = (k1 or "").strip(); k2 = (k2 or "").strip(); k3 = (k3 or "").strip()
    leaf = k3 or k2 or k1 or "(sınıfsız)"
    master[sid] = dict(ad=(ad or "").strip(), katana=(ka or "").strip(), kat1=k1,
                       kat2=k2, kat3=k3, leaf=leaf, marka=(mrk or "").strip())
print(f"  {len(master)} ürün (Hediyelik/Kırtasiye/Oyuncak)", flush=True)

# geçici tablo: universe stkID'ler (join — hızlı, injection yok)
cur.execute("CREATE TABLE #u (stkID int PRIMARY KEY)")
cur.executemany("INSERT INTO #u VALUES (?)", [(int(s),) for s in master.keys()])

# ---- 2) STOK: şube (stokSon) + depo (WMS) ----
print("Stok (şube + depo)...", flush=True)
stok = {}
cur.execute("""
    SELECT z.stkID, SUM(z.stok) FROM (
        SELECT ehstkID stkID, stok FROM dbo.stokSon_vw WITH(NOLOCK) WHERE ehMekan IN (1,4477,4478)
        UNION ALL SELECT stkID, Stok FROM depo.stok_adres_palet_vw WITH(NOLOCK) WHERE adrsAlanTipID IN (0,1)
    ) z JOIN #u ON #u.stkID=z.stkID GROUP BY z.stkID""")
for sid, s in cur.fetchall():
    stok[sid] = int(s or 0)

# ---- 3) SATIŞ 12 AY: irsHrk ehTip 1,4,100 (şube+POS) ----
print("Satış (son 12 ay)...", flush=True)
satis = {}
cur.execute(f"""
    SELECT h.ehstkID, CONVERT(int, SUM(-h.ehAdetN))
    FROM dbo.irsHrk h WITH(NOLOCK) JOIN #u ON #u.stkID=h.ehstkID
    WHERE h.ehTip IN (1,4,100) AND h.ehTrhS>='{T_S12}'
    GROUP BY h.ehstkID""")
for sid, q in cur.fetchall():
    satis[sid] = max(0, int(q or 0))

# ---- 4) MALIYET: son alış faturası > devir(99) fallback ----
print("Maliyet (fatAyr son fatura > devir99)...", flush=True)
maliyet = {}
cur.execute("""
    SELECT stkID, birim FROM (
      SELECT fa.ehstkID stkID, CONVERT(decimal(18,4), SUM(fa.ehTutarN)/NULLIF(SUM(fa.ehAdetN),0)) birim,
             ROW_NUMBER() OVER (PARTITION BY fa.ehstkID ORDER BY MAX(f.eTarih) DESC, f.eID DESC) rn
      FROM dbo.fatAyr fa WITH(NOLOCK) JOIN dbo.fat f WITH(NOLOCK) ON fa.ehID=f.eID
      JOIN #u ON #u.stkID=fa.ehstkID
      WHERE f.eTip=0 AND f.eDurum<>2 AND fa.ehAdetN>0
      GROUP BY fa.ehstkID, f.eID) t WHERE rn=1""")
for sid, b in cur.fetchall():
    maliyet[sid] = float(b or 0)
cur.execute("""
    SELECT h.ehstkID, CONVERT(decimal(18,4), SUM(h.ehTutarN)/NULLIF(SUM(h.ehAdetN),0))
    FROM dbo.irsHrk h WITH(NOLOCK) JOIN #u ON #u.stkID=h.ehstkID
    WHERE h.ehTip=99 AND h.ehTrhS>='20210531' AND h.ehTrhS<'20210601'
    GROUP BY h.ehstkID HAVING SUM(h.ehAdetN)<>0""")
for sid, b in cur.fetchall():
    maliyet.setdefault(sid, float(b or 0))
cur.execute("DROP TABLE #u")
cur.close(); cn.close()

# ================= HESAP (Python) =================
print("Hesap...", flush=True)
DIV = 12.0 / HEDEF_AY   # satış12 / DIV = HEDEF_AY-aylık kapsama hedefi

# per-SKU (sadece stok>0 VEYA satış>0 olan universe)
sku = {}   # sid -> dict
for sid, m in master.items():
    st = stok.get(sid, 0); s12 = satis.get(sid, 0)
    if st <= 0 and s12 <= 0:
        continue
    c = maliyet.get(sid, 0.0)
    hedef = s12 / DIV
    exc = max(0, st - hedef)      # adet
    sku[sid] = dict(stok=st, s12=s12, cost=c, hedef=hedef, exc_ad=exc, exc_tl=exc * c,
                    stok_tl=st * c, leaf=m["leaf"])

# aile (leaf) topla
aile = {}   # leaf -> dict
for sid, x in sku.items():
    a = aile.setdefault(x["leaf"], dict(sku=0, stok=0, s12=0, stok_tl=0.0, exc_tl_sku=0.0,
                                        katana="", kat1="", kat2="", kat3=""))
    a["sku"] += 1; a["stok"] += x["stok"]; a["s12"] += x["s12"]
    a["stok_tl"] += x["stok_tl"]; a["exc_tl_sku"] += x["exc_tl"]
    if not a["katana"]:
        m = master[sid]; a.update(katana=m["katana"], kat1=m["kat1"], kat2=m["kat2"], kat3=m["kat3"])

# aile metrikleri + sınıf
for leaf, a in aile.items():
    a["avg_cost"] = (a["stok_tl"] / a["stok"]) if a["stok"] > 0 else 0.0
    a["aile_ay"] = (a["stok"] * 12.0 / a["s12"]) if a["s12"] > 0 else 999.0
    a["hedef"] = a["s12"] / DIV
    a["exc_ad"] = max(0, a["stok"] - a["hedef"])          # aile fazla adet
    a["exc_tl"] = a["exc_ad"] * a["avg_cost"]              # aile fazla ₺ (gerçek serbest-kalabilir)
    a["phantom_tl"] = max(0.0, a["exc_tl_sku"] - a["exc_tl"])   # per-SKU'nun şişirdiği (ikame)
    a["phantom_oran"] = (a["phantom_tl"] / a["exc_tl_sku"]) if a["exc_tl_sku"] > 0 else 0.0
    if a["exc_tl_sku"] < 1:
        a["sinif"] = "SAĞLIKLI"
    elif a["aile_ay"] <= IKAME_AILE_AY and a["phantom_oran"] >= IKAME_PHANTOM:
        a["sinif"] = "İKAME"
    elif a["exc_tl"] >= MAT_AILE:
        a["sinif"] = "GERÇEK FAZLA"
    else:
        a["sinif"] = "SAĞLIKLI"

# toplamlar
T_stok_tl = sum(a["stok_tl"] for a in aile.values())
T_exc_sku = sum(a["exc_tl_sku"] for a in aile.values())    # naif per-SKU fazla
T_exc_aile = sum(a["exc_tl"] for a in aile.values())        # aile-düzeltmeli (gerçek)
T_phantom = T_exc_sku - T_exc_aile
sinif_tl = {}
for a in aile.values():
    d = sinif_tl.setdefault(a["sinif"], dict(sku=0, aile=0, stok_tl=0.0, exc_sku=0.0, exc_aile=0.0, phantom=0.0))
    d["aile"] += 1; d["sku"] += a["sku"]; d["stok_tl"] += a["stok_tl"]
    d["exc_sku"] += a["exc_tl_sku"]; d["exc_aile"] += a["exc_tl"]; d["phantom"] += a["phantom_tl"]

print(f"  {len(sku)} aktif ürün · {len(aile)} aile", flush=True)
print(f"  Envanter değeri: {T_stok_tl:,.0f} ₺", flush=True)
print(f"  Naif per-SKU fazla: {T_exc_sku:,.0f} ₺  →  Aile-düzeltmeli: {T_exc_aile:,.0f} ₺  (phantom {T_phantom:,.0f} ₺)", flush=True)

# ================= EXCEL =================
print("Excel yazılıyor...", flush=True)
wb = openpyxl.Workbook()

# stiller
F_H = Font(bold=True, color="FFFFFF", size=10)
FILL_H = PatternFill("solid", fgColor="1F3A5F")
FILL_KPI = PatternFill("solid", fgColor="EEF3F8")
FILL_IKAME = PatternFill("solid", fgColor="FDECEA")     # ikame = kırmızımsı (dikkat: phantom)
FILL_GERCEK = PatternFill("solid", fgColor="FCF3D6")    # gerçek fazla = sarımsı
FILL_SAGLIK = PatternFill("solid", fgColor="E9F6EC")    # sağlıklı = yeşilimsi
THIN = Side(style="thin", color="D0D5DD")
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def hdr(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row, j, c); cell.font = F_H; cell.fill = FILL_H; cell.border = BORD
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 1)
def tl(v): return round(float(v or 0))
SINIF_FILL = {"İKAME": FILL_IKAME, "GERÇEK FAZLA": FILL_GERCEK, "SAĞLIKLI": FILL_SAGLIK}

# ---------- Sheet 1: Özet ----------
ws = wb.active; ws.title = "Özet"
ws["A1"] = "ENVANTER FAZLA — AİLE-DÜZELTMELİ ANALİZ"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = (f"Kapsam: Hediyelik/Kırtasiye/Oyuncak · {len(sku):,} aktif ürün · {len(aile):,} ürün ailesi (leaf kategori) · "
            f"{datetime.date.today():%d.%m.%Y}").replace(",", ".")
ws["A2"].font = Font(italic=True, color="667085")
ws["A4"] = "İKAME CONFOUND: emtia ailelerinde (kağıt, puzzle, markör) talep rafta ne varsa ona kayar →"
ws["A5"] = "per-SKU fazla YANILTICI şişer. Doğru serbest-kalabilir para = AİLE düzeyinde fazla (aile-içi ikame netlenir)."
for r in (4, 5): ws[f"A{r}"].font = Font(color="475467")

kpi = [
    ("Toplam envanter değeri (maliyet)", T_stok_tl, "FFFFFF"),
    ("Naif per-SKU fazla (İKAME confound VAR — şişkin)", T_exc_sku, "FDECEA"),
    ("AİLE-DÜZELTMELİ fazla (gerçek serbest-kalabilir)", T_exc_aile, "E9F6EC"),
    ("Phantom (ikame kayması — gerçek değil)", T_phantom, "FCF3D6"),
]
r0 = 7
for i, (lbl, val, col) in enumerate(kpi):
    r = r0 + i
    ws.cell(r, 1, lbl).font = Font(bold=True, size=11)
    c = ws.cell(r, 2, tl(val)); c.number_format = '#,##0 "₺"'; c.font = Font(bold=True, size=12)
    c.fill = PatternFill("solid", fgColor=col); c.alignment = Alignment(horizontal="right")
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=col)
ws.cell(r0 + len(kpi) + 1, 1, "Düzeltme oranı (phantom ÷ naif)")
cc = ws.cell(r0 + len(kpi) + 1, 2, T_phantom / T_exc_sku if T_exc_sku else 0)
cc.number_format = "0.0%"; cc.alignment = Alignment(horizontal="right")

# sınıf dağılımı
rS = r0 + len(kpi) + 4
ws.cell(rS - 1, 1, "SINIF DAĞILIMI").font = Font(bold=True, size=12)
scols = ["Sınıf", "Aile", "SKU", "Envanter ₺", "Naif per-SKU fazla ₺", "Aile-düzeltmeli fazla ₺", "Phantom ₺"]
for j, c in enumerate(scols, 1):
    cell = ws.cell(rS, j, c); cell.font = F_H; cell.fill = FILL_H; cell.border = BORD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
ws.row_dimensions[rS].height = 28
order = ["GERÇEK FAZLA", "İKAME", "SAĞLIKLI"]
for i, sn in enumerate([s for s in order if s in sinif_tl], 1):
    d = sinif_tl[sn]; r = rS + i
    vals = [sn, d["aile"], d["sku"], tl(d["stok_tl"]), tl(d["exc_sku"]), tl(d["exc_aile"]), tl(d["phantom"])]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(r, j, v); cell.border = BORD; cell.fill = SINIF_FILL.get(sn, FILL_KPI)
        if j >= 4: cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal="right")
        if j == 1: cell.font = Font(bold=True)
ws.column_dimensions["A"].width = 46
for col in "BCDEFG": ws.column_dimensions[col].width = 18

# yorum kutusu
rY = rS + len([s for s in order if s in sinif_tl]) + 3
notlar = [
    "NASIL OKUNUR:",
    "• GERÇEK FAZLA = ailenin KENDİSİ şişkin (aile_ay yüksek). Per-SKU fazla gerçek — makas/cetvel/silgi burada. Aksiyon: sipariş durdur, likidite et.",
    "• İKAME = aile SAĞLIKLI (aile_ay≤6) ama per-SKU şişkin. Fazla YANILTICI — talep marka/desen arası kayıyor (kağıt, puzzle). Aksiyon: aile düzeyinde yönet, SKU çeşidini azalt (fazla stok değil).",
    "• SAĞLIKLI = anlamlı fazla yok.",
    "• Hedef = 3-ay kapsama (satış12÷4), per-SKU + aile AYNI kural → phantom hedef-seçiminden bağımsız.",
    "• Maliyet = son alış faturası (fatAyr) > 2021 devir fallback. KDV hariç.",
]
for i, n in enumerate(notlar):
    cell = ws.cell(rY + i, 1, n)
    cell.font = Font(bold=(i == 0), color="475467", size=10); cell.alignment = Alignment(wrap_text=False)

# ---------- Sheet 2: Aile Analizi ----------
ws2 = wb.create_sheet("Aile Analizi")
A_HEAD = ["Sınıf", "Ana Kategori", "Departman (Kat1)", "Ürün Ailesi (leaf)",
          "SKU", "Toplam Stok (adet)", "Son 12 Ay Satış (adet)", "Aile Kaç Ayda Tükenir",
          "Envanter Değeri ₺", "Naif per-SKU Fazla ₺", "Aile-Düzeltmeli Fazla ₺",
          "Phantom ₺ (ikame)", "Phantom %"]
hdr(ws2, A_HEAD)
arows = sorted(aile.items(), key=lambda kv: -kv[1]["exc_tl_sku"])
r = 2
for leaf, a in arows:
    if a["exc_tl_sku"] < 1 and a["stok_tl"] < 1:
        continue
    vals = [a["sinif"], a["katana"], a["kat1"], leaf,
            a["sku"], a["stok"], a["s12"],
            round(a["aile_ay"], 1) if a["aile_ay"] < 999 else None,
            tl(a["stok_tl"]), tl(a["exc_tl_sku"]), tl(a["exc_tl"]),
            tl(a["phantom_tl"]), a["phantom_oran"]]
    fill = SINIF_FILL.get(a["sinif"], FILL_KPI)
    for j, v in enumerate(vals, 1):
        cell = ws2.cell(r, j, v); cell.border = BORD
        if j == 1: cell.fill = fill; cell.font = Font(bold=True, size=9)
        if j in (5, 6, 7): cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal="right")
        if j == 8: cell.number_format = '0.0'; cell.alignment = Alignment(horizontal="right")
        if j in (9, 10, 11, 12): cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal="right")
        if j == 13: cell.number_format = '0.0%'; cell.alignment = Alignment(horizontal="right")
    r += 1
widths2 = [13, 14, 22, 30, 7, 13, 14, 12, 16, 16, 16, 15, 10]
for j, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ---------- Sheet 3: Ürün Detay ----------
ws3 = wb.create_sheet("Ürün Detay")
D_HEAD = ["Ürün Kodu", "Ürün Adı", "Marka", "Ana Kategori", "Departman", "Ürün Ailesi (leaf)", "Aile Sınıfı",
          "Stok (adet)", "Son 12 Ay Satış", "Hedef Stok (3 ay)", "Fazla (adet)",
          "Birim Maliyet ₺", "Fazla Bağlı Para ₺", "Aile Kaç Ayda Tükenir"]
hdr(ws3, D_HEAD)
drows = sorted(sku.items(), key=lambda kv: -kv[1]["exc_tl"])
r = 2
for sid, x in drows:
    if x["exc_tl"] < 1:
        continue
    m = master[sid]; a = aile[x["leaf"]]
    vals = [sid, m["ad"], m["marka"], m["katana"], m["kat1"], x["leaf"], a["sinif"],
            x["stok"], x["s12"], round(x["hedef"], 1), round(x["exc_ad"]),
            round(x["cost"], 2), tl(x["exc_tl"]),
            round(a["aile_ay"], 1) if a["aile_ay"] < 999 else None]
    for j, v in enumerate(vals, 1):
        cell = ws3.cell(r, j, v); cell.border = BORD
        if j == 7: cell.fill = SINIF_FILL.get(a["sinif"], FILL_KPI); cell.font = Font(bold=True, size=9)
        if j in (8, 9, 10, 11): cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal="right")
        if j == 12: cell.number_format = '#,##0.00'; cell.alignment = Alignment(horizontal="right")
        if j == 13: cell.number_format = '#,##0'; cell.alignment = Alignment(horizontal="right")
        if j == 14: cell.number_format = '0.0'; cell.alignment = Alignment(horizontal="right")
    r += 1
widths3 = [10, 40, 16, 13, 20, 26, 12, 10, 12, 13, 10, 13, 15, 13]
for j, w in enumerate(widths3, 1):
    ws3.column_dimensions[get_column_letter(j)].width = w

# kaydet
R = os.path.join(os.path.dirname(__file__), "..")
_slug = datetime.date.today().strftime("%Y%m%d")
OUT = os.path.join(R, "briefings", f"envanter-aile-analiz-{_slug}.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"✓ Yazıldı: {OUT}", flush=True)
