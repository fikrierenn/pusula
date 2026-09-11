# -*- coding: utf-8 -*-
"""Ağu-Eyl 2026 'geçmişi olduğu halde aşırı alınan' dilim → Excel (kategori/tedarikçi/ürün).

Tanım: irsHrk alım (tip 0 Alış + 10 Yerel Alım − 2 Alış İade), 01.08–10.09.2026.
Kohort = tabanda EŞLEŞEN (okul ders kitabı hariç) + ESKİ ürün (IlkGiris < 1 yıl önce) +
(365g hiç satmamış  VEYA  arz günü > 365). Cold-start (yeni ürün) DIŞARIDA — meşru lansman.
Tedarikçi: irs.eFirma → frm.frmAd. Grup/ilişkili taraf işaretlenir (bridges SatinIliskiliTaraf).
Arz günü = alınan adet ÷ (Satis365/365). Sansür: satış sağdan sansürlü → arz günü ÜST SINIR.

ÖLÇÜM (10.09.2026): net alım 151,2M ₺ / 34.140 çeşit. Cold-start ve okul kitabı ayıklandıktan
sonra hesap sorulacak dilim = >365 ESKİ 17,4M + satışsız ESKİ 4,54M ≈ **21,9M ₺ (%14,5)**.
Arşiv SQL + tam bulgu: `sorgular/2026-09-10-agu-eyl-alim-arz-gunu.sql`.

⚠ KOK sabiti mutlak yol taşıyor (scratchpad'den taşındı) — başka makinede çalışması için
dosya konumundan türetiliyor.

⚠ Alıcıya atıf YAPILAMAZ: kararı kimin verdiği veride izli değil. Bu bir ürün/tedarikçi
kohortudur, kişi karnesi değil (satinalma-danisman adil-atıf şartı).
"""
from __future__ import annotations
import io, os, re, sys
import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ILISKILI = {9525, 22100, 56, 38093, 4841, 23842, 58, 9339, 4694, 7950, 50582}  # grup-içi frmID

def env_oku(yol):
    env = {}
    with open(yol, encoding="utf-8") as f:
        for ln in f:
            if ln.lstrip().startswith("#"):
                continue
            m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.+?)\s*$", ln)
            if m:
                env[m.group(1)] = m.group(2).strip().strip('"')
    return env

env = env_oku(os.path.join(KOK, ".env"))
host, port = env["MSSQL_HOST"], env.get("MSSQL_PORT", "1433")
assert re.fullmatch(r"[A-Za-z0-9._\-]+", host) and re.fullmatch(r"\d+", port)
cn = pyodbc.connect(
    "Driver={ODBC Driver 18 for SQL Server};"
    f"Server={host},{port};Database=DerinSISBkm;"
    f"UID={env['MSSQL_USER']};PWD={env['MSSQL_PASSWORD']};"
    "TrustServerCertificate=yes;Timeout=30", timeout=900)

SQL = """
DECLARE @kesim date = (SELECT MAX(Kesim) FROM DerinSISBkm.bkm.SatisAnaliziTaban);
WITH alim AS (
    SELECT h.ehstkID AS stkID, SUM(h.ehAdetN) AS adet, SUM(h.ehTutarN) AS tutar
    FROM DerinSISBkm.dbo.irsHrk h
    WHERE h.ehTrhS >= '20260801' AND h.ehTrhS < '20260911' AND h.ehTip IN (0,10,2)
    GROUP BY h.ehstkID HAVING SUM(h.ehAdetN) > 0
),
sup_raw AS (
    SELECT h.ehstkID AS stkID, i.eFirma AS frmID,
           -- ⚠ SIFIR SENTINEL (12.09.2026): eFirma=0 = karşı taraf YOK (iç işlem).
           -- frm'de frmID=0 satırı VAR, adı 'GENEL' → ISNULL koruması çalışmaz,
           -- tedarikçi listesine sahte bir firma girer.
           CASE WHEN i.eFirma = 0 THEN N'(tedarikçi yok — iç işlem)'
                ELSE ISNULL(f.frmAd, '(firma yok)') END AS frmAd,
           SUM(h.ehAdetN) AS adet
    FROM DerinSISBkm.dbo.irsHrk h
    JOIN DerinSISBkm.dbo.irs i ON i.eID = h.ehID
    LEFT JOIN DerinSISBkm.dbo.frm f ON f.frmID = i.eFirma
    WHERE h.ehTrhS >= '20260801' AND h.ehTrhS < '20260911' AND h.ehTip IN (0,10,2)
    GROUP BY h.ehstkID, i.eFirma, f.frmAd
),
sup AS (
    SELECT stkID, frmID, frmAd,
           ROW_NUMBER() OVER (PARTITION BY stkID ORDER BY adet DESC) AS rn
    FROM sup_raw
),
t AS (
    SELECT stkID, stkAd, Kategori3, Kategori1, SatisToplam, IlkGiris, BirimMaliyet
    FROM DerinSISBkm.bkm.SatisAnaliziTaban
    WHERE Kesim = @kesim AND SezonYil = 2025
)
SELECT
    a.stkID,
    LEFT(t.stkAd, 70)                              AS Urun,
    t.Kategori1, t.Kategori3,
    s.frmID, s.frmAd                               AS Tedarikci,
    CONVERT(int, a.adet)                           AS AlinanAdet,
    CONVERT(bigint, a.tutar)                       AS AlinanTutar,
    CONVERT(int, t.SatisToplam)                    AS Satis365,
    CONVERT(int, CASE WHEN t.SatisToplam > 0
        THEN a.adet / (t.SatisToplam/365.0) END)   AS ArzGun,
    CONVERT(date, t.IlkGiris)                      AS IlkGiris,
    CASE WHEN ISNULL(t.SatisToplam,0)=0 THEN N'satışsız (365g 0)' ELSE N'>365 gün arz' END AS Kova
FROM alim a
JOIN t ON t.stkID = a.stkID
LEFT JOIN sup s ON s.stkID = a.stkID AND s.rn = 1
WHERE t.IlkGiris < DATEADD(YEAR, -1, @kesim)                       -- ESKİ (cold-start hariç)
  AND ( ISNULL(t.SatisToplam,0) = 0                               -- satışsız
     OR a.adet / NULLIF(t.SatisToplam/365.0, 0) > 365 )           -- >1 yıl arz
ORDER BY a.tutar DESC
"""

cur = cn.cursor()
cur.execute(SQL)
cols = [c[0] for c in cur.description]
rows = [dict(zip(cols, r)) for r in cur.fetchall()]
cn.close()

if len(rows) < 100:
    print(f"KOSAMADI: yalnizca {len(rows)} satir — filtre/kesim kontrol", file=sys.stderr)
    raise SystemExit(2)

for r in rows:
    r["Grup"] = "İlişkili taraf (grup içi)" if r["frmID"] in ILISKILI else "Dış tedarikçi"

toplam_tutar = sum(r["AlinanTutar"] or 0 for r in rows)
toplam_adet = sum(r["AlinanAdet"] or 0 for r in rows)

def pivot(anahtar):
    d = {}
    for r in rows:
        k = r[anahtar] or "(boş)"
        e = d.setdefault(k, {"cesit": 0, "adet": 0, "tutar": 0})
        e["cesit"] += 1
        e["adet"] += r["AlinanAdet"] or 0
        e["tutar"] += r["AlinanTutar"] or 0
    return sorted(d.items(), key=lambda x: -x[1]["tutar"])

BASLIK_DOLGU = PatternFill("solid", fgColor="1F3864")
def basyap(ws, basliklar):
    ws.append(basliklar)
    for c in range(1, len(basliklar)+1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = BASLIK_DOLGU
        cell.alignment = Alignment(wrap_text=True, vertical="center")

wb = Workbook()

# 1) Kategori özet
ws = wb.active; ws.title = "1-Kategori"
basyap(ws, ["Kategori3", "Çeşit", "Adet", "Tutar ₺", "Pay %"])
for k, v in pivot("Kategori3"):
    ws.append([k, v["cesit"], v["adet"], v["tutar"], round(100*v["tutar"]/toplam_tutar, 1)])
ws.append(["TOPLAM", len(rows), toplam_adet, toplam_tutar, 100.0])
ws.cell(ws.max_row, 1).font = Font(bold=True)

# 2) Tedarikçi özet (+ grup ayrımı)
ws2 = wb.create_sheet("2-Tedarikci")
basyap(ws2, ["Tedarikçi", "Grup", "Çeşit", "Adet", "Tutar ₺", "Pay %"])
tp = {}
for r in rows:
    k = r["Tedarikci"] or "(firma yok)"
    e = tp.setdefault(k, {"grup": r["Grup"], "cesit": 0, "adet": 0, "tutar": 0})
    e["cesit"] += 1; e["adet"] += r["AlinanAdet"] or 0; e["tutar"] += r["AlinanTutar"] or 0
for k, v in sorted(tp.items(), key=lambda x: -x[1]["tutar"]):
    ws2.append([k, v["grup"], v["cesit"], v["adet"], v["tutar"], round(100*v["tutar"]/toplam_tutar, 1)])
# grup vs dış özet en üste not
gr = {"İlişkili taraf (grup içi)": 0, "Dış tedarikçi": 0}
for r in rows:
    gr[r["Grup"]] += r["AlinanTutar"] or 0

# 3) Ürün tek tek
ws3 = wb.create_sheet("3-Urun tek tek")
basyap(ws3, ["#", "stkID", "Ürün", "Kategori", "Tedarikçi", "Grup", "Alınan adet",
             "Tutar ₺", "Satış 365g", "Arz günü", "İlk giriş", "Kova"])
for n, r in enumerate(rows, 1):
    ws3.append([n, r["stkID"], r["Urun"], r["Kategori3"], r["Tedarikci"], r["Grup"],
                r["AlinanAdet"], r["AlinanTutar"], r["Satis365"], r["ArzGun"],
                r["IlkGiris"], r["Kova"]])
gen = [5, 9, 46, 14, 40, 22, 11, 13, 10, 9, 12, 16]
for i, w in enumerate(gen, 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = "C2"

# 0) Özet/yöntem sayfası — en başa
wsz = wb.create_sheet("0-OZET", 0)
wsz.column_dimensions["A"].width = 48; wsz.column_dimensions["B"].width = 22; wsz.column_dimensions["C"].width = 70
def s(a, b="", c=""): wsz.append([a, b, c])
s("AĞU-EYL 2026 — GEÇMİŞİ OLDUĞU HALDE AŞIRI ALINAN", "", "01.08.2026 – 10.09.2026 · irsHrk tip 0+10−2")
s("")
s("Toplam çeşit", len(rows))
s("Toplam adet", toplam_adet)
s("Toplam tutar ₺", toplam_tutar)
s("  — İlişkili taraf (grup içi)", round(gr["İlişkili taraf (grup içi)"]), "ODAK/POİNT vb. grup şirketi — dış alım değil, transfer benzeri")
s("  — Dış tedarikçi", round(gr["Dış tedarikçi"]))
s("")
s("KOHORT TANIMI", "", "")
s("dahil", "tabanda eşleşen", "okul ders kitabı (tabanda yok) HARİÇ — kurumsal kesin talep")
s("dahil", "ESKİ ürün", "IlkGiris < 1 yıl önce — cold-start/yeni lansman HARİÇ (meşru)")
s("dahil", "satışsız VEYA >365g", "365g hiç satmamış ya da >1 yıllık arz alınmış")
s("")
s("YÖNTEM / SINIR", "", "")
s("arz günü", "alınan ÷ (Satis365/365)", "kaç günlük stok satın alındı")
s("ÜST SINIR", "sağdan sansür", "stok bitince satış kesilir → talep düşük görünür → arz günü ABARTILI; gerçek fazlalık DAHA AZ")
s("payda kararsız", "yavaş devir", "düşük hızlı üründe her parti 'aşırı' görünür — MOQ olabilir, tek başına suç değil")
s("alıcı yok", "veride izli değil", "kişiye atıf YAPILAMAZ — süreç düzeyinde okunur")
for c in (1,):
    wsz.cell(1, c).font = Font(bold=True)

cikti = os.path.join(KOK, "raporlar", "asiri-alim-agu-eyl-2026.xlsx")
os.makedirs(os.path.dirname(cikti), exist_ok=True)
wb.save(cikti)
print(f"cesit {len(rows)} · adet {toplam_adet} · tutar {toplam_tutar:,.0f} TL")
print(f"  iliskili taraf: {gr['İlişkili taraf (grup içi)']:,.0f} · dis: {gr['Dış tedarikçi']:,.0f}")
print(f"Yazildi: {cikti}")
