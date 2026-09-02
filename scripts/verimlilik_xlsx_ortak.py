# -*- coding: utf-8 -*-
"""Excel emitter ORTAK katmani: bicim sabitleri + baslik/satir/not yardimcilari.

Emitter HESAP YAPMAZ (emitter-ayrimi kurali) — cekirdegin urettigi JSON'u bicime doker.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ADET = "#,##0"
ADET1 = "#,##0.0"
TL = "#,##0"
YUZDE = "+0.0%;-0.0%"
KAT = "0.0\"x\""

BASLIK = PatternFill("solid", fgColor="E30622")
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=10)
GRI = PatternFill("solid", fgColor="F2F2F2")
VURGU = PatternFill("solid", fgColor="FFF3CD")
YESIL_YAZI = Font(bold=True, color="1F7A4D")
INCE = Side(style="thin", color="D9D9D9")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)
NOT_YAZI = Font(italic=True, size=9, color="666666")
BOLUM_YAZI = Font(bold=True, size=10)


def _basliklar(ws, kolonlar, satir=1):
    """kolonlar = [(baslik, genislik, format), ...]"""
    for i, (ad, gen, _f) in enumerate(kolonlar, start=1):
        h = ws.cell(satir, i, ad)
        h.fill = BASLIK
        h.font = BASLIK_YAZI
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = KENAR
        ws.column_dimensions[get_column_letter(i)].width = gen
    ws.row_dimensions[satir].height = 30
    ws.freeze_panes = ws.cell(satir + 1, 2)


def _yaz(ws, satir, kolonlar, degerler):
    for i, ((_ad, _gen, fmt), deger) in enumerate(zip(kolonlar, degerler), start=1):
        c = ws.cell(satir, i, deger)
        c.border = KENAR
        if fmt:
            c.number_format = fmt
        if i == 1:
            c.alignment = Alignment(horizontal="left")
    return satir + 1


def _notlar(ws, notlar, satir, kol=1):
    for n in notlar:
        c = ws.cell(satir, kol, n)
        c.font = NOT_YAZI
        c.alignment = Alignment(wrap_text=False)
        satir += 1
    return satir
