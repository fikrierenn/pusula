"""
PERSONEL KARTI ANOMALİ RAPORU
Personel müşteri kartı (DerinCrm.Customer.Group1Id=2) başkalarının alışverişinde
okutuluyor mu? Sinyal: kart çok farklı KASİYER elinde okutulmuş (kendi alışverişi
1-2 kasiyer olur) + yüksek personel indirimi. Suistimal/puan-indirim kaçağı tespiti.
Çıktı: briefings/personel-kart-anomali.xlsx
Kullanım: python personel_kart_anomali.py [--gun 90]
"""
import sys, os, json, argparse
from pathlib import Path
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
import pymssql
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

R = Path(__file__).resolve().parent.parent
KIRMIZI = "E30622"


def cfg():
    e = R / ".env"
    if e.exists():
        for l in e.read_text(encoding="utf-8").splitlines():
            if "=" in l and not l.strip().startswith("#"):
                k, v = l.split("=", 1); os.environ.setdefault(k.strip(), v.strip())
    h = os.environ.get("MSSQL_HOST")
    if h:
        return dict(server=h, user=os.environ.get("MSSQL_USER", "sa"), password=os.environ.get("MSSQL_PASSWORD", ""), database=os.environ.get("MSSQL_DATABASE", "master"))
    return json.loads((R / ".secrets" / "db.json").read_text(encoding="utf-8"))


def get_data(gun):
    c = cfg()
    conn = pymssql.connect(server=c["server"], user=c["user"], password=c["password"], database=c.get("database", "master"), charset="UTF-8", login_timeout=20, timeout=180)
    cur = conn.cursor(as_dict=True)
    cur.execute("""SELECT c.Id, CAST(c.Name AS nvarchar(40)) Personel, CAST(c.PhoneNumber AS nvarchar(15)) Tel,
        COUNT(*) Fis, COUNT(DISTINCT s.UsersId) Kasiyer, COUNT(DISTINCT s.PosId) Kasa,
        COUNT(DISTINCT CAST(s.Date AS date)) Gun,
        CAST(SUM(s.GrossTotal-s.DiscountTotal) AS decimal(18,0)) Ciro,
        CAST(SUM(s.DiscountTotal) AS decimal(18,0)) Indirim
      FROM EncoreMerkez.dbo.Sales s WITH(NOLOCK)
      JOIN DerinCrm.dbo.Customer c WITH(NOLOCK) ON c.Id=s.CustomersId AND c.Group1Id=2
      WHERE s.DocumentsTypeId=1 AND s.Date>=DATEADD(DAY,-%d,CAST(GETDATE() AS date))
      GROUP BY c.Id, CAST(c.Name AS nvarchar(40)), CAST(c.PhoneNumber AS nvarchar(15))
      HAVING COUNT(*)>=3""" % gun)
    rows = cur.fetchall()
    cur.close(); conn.close()
    return rows


def build(gun):
    data = get_data(gun)
    out = []
    for r in data:
        fis = int(r["Fis"]); kas = int(r["Kasiyer"]); ciro = float(r["Ciro"] or 0); ind = float(r["Indirim"] or 0)
        gross = ciro + ind
        indp = 100 * ind / gross if gross else 0
        # anomali skoru: farklı kasiyer ana sinyal (kart paylaşımı); indirim oranı destekleyici
        if kas >= 8:
            durum, sev = "YÜKSEK — kart paylaşımı şüphesi", 3
        elif kas >= 5:
            durum, sev = "ORTA — çok kasiyer", 2
        elif kas >= 3 and indp >= 25:
            durum, sev = "İZLE — indirim yoğun", 1
        else:
            durum, sev = "normal", 0
        out.append(dict(id=r["Id"], ad=r["Personel"], tel=r["Tel"] or "", fis=fis, kas=kas, kasa=int(r["Kasa"]),
                        gun=int(r["Gun"]), ciro=round(ciro), ind=round(ind), indp=round(indp, 1), durum=durum, sev=sev))
    out.sort(key=lambda x: (-x["sev"], -x["kas"], -x["fis"]))

    wb = Workbook()
    red = PatternFill("solid", fgColor=KIRMIZI); white = Font(color="FFFFFF", bold=True)
    thin = Border(*[Side(style="thin", color="E2E8F0")] * 4)

    def tl(v): return f"{v:,.0f}".replace(",", ".")

    tFis = sum(x["fis"] for x in out); tInd = sum(x["ind"] for x in out)
    yuksek = [x for x in out if x["sev"] >= 2]

    ws = wb.active; ws.title = "Yönetici Özeti"
    ws["A1"] = "BKM KİTAP — PERSONEL KARTI ANOMALİ RAPORU"; ws["A1"].font = Font(bold=True, size=15, color=KIRMIZI)
    ws["A2"] = f"Son {gun} gün · personel kartı (Group1Id=2) · sinyal: kart çok farklı KASİYER elinde okutulmuş = kendi değil başkasına"; ws["A2"].font = Font(italic=True, color="64748B")
    kpis = [("Personel Kart (aktif)", tl(len(out))), ("Toplam Fiş", tl(tFis)), ("Toplam İndirim", tl(tInd) + " ₺"),
            ("Şüpheli Kart (≥5 kasiyer)", tl(len(yuksek)))]
    for i, (k, v) in enumerate(kpis):
        ws.cell(4, 1 + i * 2, k).font = Font(size=9, color="64748B")
        ws.cell(5, 1 + i * 2, v).font = Font(bold=True, size=13, color=KIRMIZI)
    lines = ["", "MANTIK", "• Personel kendi alışverişinde kartını 1-2 kasiyerde okutur. Kart 5+ farklı kasiyer elinde okutulduysa = başkalarının alışverişinde kullanılıyor (personel indirimi/puan kaçağı) şüphesi.",
             "• Yüksek toplam indirim + çok kasiyer birlikte → güçlü sinyal.", "", "EN ŞÜPHELİ KARTLAR"]
    for x in yuksek[:12]:
        lines.append(f"• {x['ad']} — {x['fis']} fiş, {x['kas']} farklı kasiyer, indirim {tl(x['ind'])} ₺ (%{str(x['indp']).replace('.',',')}) → {x['durum']}")
    lines += ["", "ÖNERİLER",
              "• Şüpheli kartların fiş detayını (hangi kasiyer, hangi gün) incele — kasiyerle yüz yüze görüş.",
              "• Personel kartında günlük/aylık kullanım limiti + farklı-kasiyer uyarısı.",
              "• Personel alışverişi sadece kendi mağazası/kendi kartı + amir onayı kuralı.",
              "", "Not: Çok kasiyer tek başına suç değil (personel başka kasada alışveriş yapmış olabilir); liste inceleme önceliği verir, kanıt değil."]
    rr = 7
    for ln in lines:
        cc = ws.cell(rr, 1, ln)
        if ln in ("MANTIK", "ÖNERİLER", "EN ŞÜPHELİ KARTLAR"): cc.font = Font(bold=True, size=12, color=KIRMIZI)
        elif ln.startswith("Not:"): cc.font = Font(italic=True, size=9, color="64748B")
        else: cc.font = Font(size=10)
        rr += 1
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGH": ws.column_dimensions[col].width = 15

    d = wb.create_sheet("Anomali Tablosu")
    hdr = ["Personel", "CustomersId", "Telefon", "Fiş", "Farklı Kasiyer", "Farklı Kasa", "Farklı Gün", "Ciro ₺", "İndirim ₺", "İndirim %", "Durum"]
    for j, h in enumerate(hdr, 1):
        cc = d.cell(1, j, h); cc.fill = red; cc.font = white; cc.alignment = Alignment(horizontal="center"); cc.border = thin
    for i, x in enumerate(out, 2):
        vals = [x["ad"], x["id"], x["tel"], x["fis"], x["kas"], x["kasa"], x["gun"], x["ciro"], x["ind"], x["indp"], x["durum"]]
        for j, v in enumerate(vals, 1):
            cc = d.cell(i, j, v); cc.border = thin
            if j in (4, 5, 6, 7, 8, 9): cc.number_format = "#,##0"
            if j == 10: cc.number_format = "0.0"
        if x["sev"] >= 3:
            for j in range(1, 12): d.cell(i, j).fill = PatternFill("solid", fgColor="FEE2E2")
            d.cell(i, 5).font = Font(bold=True, color="DC2626")
        elif x["sev"] == 2:
            d.cell(i, 5).font = Font(bold=True, color="DC2626")
    for j, w in enumerate([26, 12, 14, 8, 14, 11, 11, 12, 12, 10, 28], 1):
        d.column_dimensions[get_column_letter(j)].width = w
    d.freeze_panes = "A2"

    outp = R / "briefings" / "personel-kart-anomali.xlsx"
    wb.save(outp)
    print(f"Excel: {outp}")
    print(f"  {gun} gün · {len(out)} personel kart · {len(yuksek)} şüpheli (≥5 kasiyer) · toplam indirim {tl(tInd)} ₺")
    for x in yuksek[:5]:
        print(f"    {x['ad']}: {x['fis']} fiş / {x['kas']} kasiyer / indirim {tl(x['ind'])}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--gun", type=int, default=90); a = ap.parse_args()
    build(a.gun)
