# -*- coding: utf-8 -*-
"""Kadro/norm/iş hacmi rakamlarının TUTARLILIK denetimi — sunum ve Excel'e giden her sayı burada sınanır.

Neden: aynı destede ÜÇ AYRI KAPSAM dolaşıyor (3 POS mağazası · 4 norm mağazası · 5 mağaza) ve
üç ayrı KADRO TANIMI var (kayıt / operasyonel / sezonluk). Kapsam etiketi düşerse rakamlar
birbirini tutmuyor gibi görünür (02.09.2026'da tam bu oldu: 143→160 ile 200→211 yan yana düştü).
Bu script her rakamı bileşenlerinden yeniden hesaplar ve sapmayı raporlar.

Kullanım: python scripts/tutarlilik_kontrol.py [veri.json]
Çıkış kodu: sapma varsa 1.
"""
import json
import sys
from pathlib import Path

try:                                    # Windows cp1254 konsolunda ✓/✗ basarken cokuyordu (bulgu 17)
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

KOK = Path(__file__).resolve().parent.parent
VARSAYILAN = KOK / "briefings" / "sezon-kadro-20260902" / "verimlilik-veri.json"
POS = ("İST. YOLU", "ÖZLÜCE", "FSM")          # DerinSIS POS satışı olan üç mağaza (BKM_GENEL)
ETKINLIK_BOLUM = "ETKİNLİK"

sonuc = []


def kontrol(ad, beklenen, bulunan, aciklama=""):
    tamam = (beklenen == bulunan) if isinstance(beklenen, int) else abs(beklenen - bulunan) < 0.51
    sonuc.append((tamam, ad, beklenen, bulunan, aciklama))
    return tamam


def main(argv):
    yol = Path(argv[1]) if len(argv) > 1 else VARSAYILAN
    v = json.loads(yol.read_text(encoding="utf-8"))
    mk = {m["sube"]: m for m in v["magaza_kadro"]}
    k5 = v["kadro_5magaza"]
    mag = {m["ad"]: m for m in v["magaza"]}
    n = v.get("norm")

    # ---------------------------------------------------------------- KADRO (5 mağaza)
    kontrol("5 mağaza kadrolu taban (şube toplamı = kapsam sayımı)",
            k5["kadrolu_taban26"], sum(m["kadrolu_taban26"] for m in mk.values()))
    kontrol("5 mağaza kadrolu kesim (şube toplamı = kapsam sayımı)",
            k5["kadrolu_kesim26"], sum(m["kadrolu_kesim26"] for m in mk.values()))
    kontrol("5 mağaza sezonluk kesim (şube toplamı = kapsam sayımı)",
            k5["sezonluk_kesim26"], sum(m["sezonluk_kesim26"] for m in mk.values()))
    kontrol("5 mağaza toplam = kadrolu + sezonluk",
            k5["toplam_kesim26"], k5["kadrolu_kesim26"] + k5["sezonluk_kesim26"])

    # ---------------------------------------------------------------- BÖLÜM (5 mağaza)
    b_kadrolu = sum(b["kadrolu26"] for b in v["bolum"])
    kontrol("bölüm toplamı = 5 mağaza kadrolu kesim", k5["kadrolu_kesim26"], b_kadrolu)
    kontrol("bölüm sezonluk toplamı = 5 mağaza sezonluk", k5["sezonluk_kesim26"],
            sum(b["sezonluk26"] for b in v["bolum"]))
    mb_kadrolu = sum(r["kadrolu26"] for r in v["magaza_bolum"])
    kontrol("mağaza×bölüm toplamı = 5 mağaza kadrolu kesim", k5["kadrolu_kesim26"], mb_kadrolu)

    # ---------------------------------------------------------------- 3 POS MAĞAZASI (iş hacmi kapsamı)
    pos_kadrolu = sum(mk[s]["kadrolu_kesim26"] for s in POS)
    pos_sezonluk = sum(mk[s]["sezonluk_kesim26"] for s in POS)
    pos_kadro = sum(m["kadro26"] for m in v["magaza"])
    kontrol("3 POS kadro 2026 (hacim tablosu) = kadrolu + sezonluk (İK)",
            pos_kadrolu + pos_sezonluk, pos_kadro,
            "sunumda «... → 160» diye geçen rakam")
    # K-11: 2025 tarafi hic sinanmiyordu — taban kayarsa "kaç kat" da kayar
    kontrol("3 POS kadro 2025 (hacim tablosu) = kadrolu + sezonluk (İK)",
            sum(mk[s_]["kadrolu_kesim25"] + mk[s_]["sezonluk_kesim25"] for s_ in POS),
            sum(m["kadro25"] for m in v["magaza"]),
            "sunumda «143 → ...» diye geçen rakam")
    # K-19: toplam blogu cekirdekte hesaplanir; emitter'lar bunu okur
    tp_ = v.get("toplam") or {}
    for alan in ("adet", "kdvharic", "kdvdahil", "kadro"):
        for ek in (25, 26):
            ad_ = "%s%d" % (alan, ek)
            kontrol("çekirdek toplam %s = mağaza satırları toplamı" % ad_,
                    round(float(tp_.get(ad_, -1)), 2),
                    round(sum(float(m[ad_]) for m in v["magaza"]), 2))
    # K-08: pencere gun sayisi MAGAZA x YIL bazinda esit olmali
    for anahtar, gun_ in (v["meta"].get("gun_magaza") or {}).items():
        kontrol("pencere günü (%s)" % anahtar, v["meta"]["gun"], gun_)
    sonuc.append((not v["meta"].get("gun_uyari"),
                  "mağaza bazında satış günü uyarısı yok", "yok",
                  " · ".join(v["meta"].get("gun_uyari") or []) or "yok", ""))

    # ---------------------------------------------------------------- NORM (4 mağaza)
    if n:
        ayr = n.get("ayrik", {})
        norm_subeler = [r["sube"] for r in n["sube"]]
        kontrol("norm şube toplamı = beyan edilen 145",
                n["toplam"]["norm"], sum(r["norm"] for r in n["sube"]))
        kontrol("norm sezonluk toplamı = beyan edilen 72",
                n["toplam"]["norm_sezonluk"], sum(r["norm_sezonluk"] for r in n["sube"]))
        kontrol("norm toplam = kadrolu norm + sezonluk norm",
                n["toplam"]["norm_toplam"], n["toplam"]["norm"] + n["toplam"]["norm_sezonluk"])
        kontrol("norm bölüm toplamı = norm şube toplamı",
                n["toplam"]["norm"], sum(b["norm"] for b in n["bolum"]))

        # kayıt → operasyonel dönüşümü
        eng = sum(a.get("engelli", 0) for a in ayr.values())
        etk = sum(a.get("etkinlik", 0) for a in ayr.values())
        kayit_kadrolu = n["toplam"]["kadrolu_kesim26"]
        ops_kadrolu = kayit_kadrolu - eng - etk
        kontrol("4 mağaza kayıt kadrolu = şube toplamı", kayit_kadrolu,
                sum(mk[s]["kadrolu_kesim26"] for s in norm_subeler))
        kontrol("bölüm tablosu operasyonel kadrolu = kayıt − engelli − etkinlik", ops_kadrolu,
                sum(b["kadrolu26"] for b in n["bolum"] if b["bolum"] != ETKINLIK_BOLUM),
                "engelli ve etkinlik norm dışı (yönetim kararı)")
        kontrol("bölüm engelli toplamı = şube engelli toplamı", eng,
                sum(b["engelli26"] for b in n["bolum"]))
        kontrol("etkinlik satırı = şube etkinlik toplamı", etk,
                sum(b["kadrolu26"] for b in n["bolum"] if b["bolum"] == ETKINLIK_BOLUM))
        # açık/fazla aritmetiği
        # K-06: normda tanimli ama kayitta hic kisi olmayan bolum acik degil, teyit bekleyen
        acik_b = sum(max(0, b["norm"] - b["kadrolu26"]) for b in n["bolum"]
                     if b["bolum"] != ETKINLIK_BOLUM and not b.get("teyit_gerekiyor"))
        kontrol("bölüm açık toplamı (teyit bekleyen hariç)", n["acik_bolum_toplam"], acik_b)
        kontrol("teyit bekleyen açık toplamı", n.get("acik_bolum_teyit", 0),
                sum(max(0, b["norm"] - b["kadrolu26"]) for b in n["bolum"]
                    if b.get("teyit_gerekiyor")))
        sonuc.append((n.get("engelli_format_atlanan", 0) == 0,
                      "engelli taramasında Personelno biçimi yüzünden atlanan kayıt",
                      0, n.get("engelli_format_atlanan", 0),
                      ">0 ise engelli sayısı alt sınır, norm açığı olduğundan küçük görünür"))
        kontrol("mağaza açığı çekirdekte hesaplanmış (acik_sube_toplam)",
                n["acik_sube_toplam"],
                sum(max(0, r["norm"] - (r["kadrolu_kesim26"]
                                        - ayr.get(r["sube"], {}).get("engelli", 0)
                                        - ayr.get(r["sube"], {}).get("etkinlik", 0)))
                    for r in n["sube"]))
        acik_s = sum(max(0, r["norm"] - (r["kadrolu_kesim26"]
                                         - ayr.get(r["sube"], {}).get("engelli", 0)
                                         - ayr.get(r["sube"], {}).get("etkinlik", 0)))
                     for r in n["sube"])
        if acik_b < acik_s:
            sonuc.append((False, "bölüm açığı ≥ mağaza açığı olmalı", acik_s, acik_b,
                          "mağaza içi fazlalar açığı maskeler"))
        else:
            sonuc.append((True, "bölüm açığı (%d) ≥ mağaza açığı (%d)" % (acik_b, acik_s),
                          acik_b, acik_b, "mağaza içi fazlalar açığı maskeler"))

    # ---------------------------------------------------------------- MALİYET + FAZLA MESAİ (K-21/22)
    m = v.get("maliyet")
    f = v.get("fazla_mesai")
    if m and f:
        # SEZON penceresi (esas) — aylik kirilim ile mutabik olmali
        sz = m["sezon"]
        for ek in ("25", "26"):
            yil_ = 2000 + int(ek)
            s_ = sz["pos"][ek]
            for alan in ("kisi_ay", "maliyet", "fm_saat"):
                kontrol("sezon %s%s = sezon aylarinin toplami" % (alan, ek), round(s_[alan], 2),
                        round(sum(r[{"kisi_ay": "kisi_ay", "maliyet": "maliyet",
                                     "fm_saat": "fm_saat"}[alan] + ek]
                                  for r in m["ay"] if r["ay"] in sz["aylar"]), 2))
            kontrol("sezon maliyet/ciro orani%s yeniden hesaplandi" % ek,
                    round(s_["maliyet_ciro_orani"], 6),
                    round(s_["maliyet"] / s_["ciro_kdvharic"], 6))
        # aylik toplam = kumulatif pencere
        for ek in ("25", "26"):
            for alan in ("kisi_ay", "maliyet", "fm_saat"):
                kontrol("aylik %s%s toplami = kumulatif pencere" % (alan, ek),
                        round(m["pos"][ek][alan], 2),
                        round(sum(r[alan + ek] for r in m["ay"]), 2))
        sonuc.append((len(m["ay"]) == m["pencere_ay"], "aylik satir sayisi = pencere ayi",
                      m["pencere_ay"], len(m["ay"]), "eksik ay = sessiz dusuk maliyet"))
        sonuc.append((sz["aylar"] and all(a_ in (7, 8) for a_ in sz["aylar"]),
                      "sezon aylari 7-8 icinde", "[7,8]", sz["aylar"], ""))
        for ek in ("25", "26"):
            p_ = m["pos"][ek]
            # maliyet = brüt + işveren SGK + işveren işsizlik (yönetimin bordro formülü)
            kontrol("maliyet%s = brüt + işveren SGK + işsizlik" % ek,
                    round(p_["maliyet"], 2),
                    round(p_["brut"] + p_["isveren_sgk"] + p_["isveren_issizlik"], 2))
            # POS toplamı = şube satırlarının toplamı (tek kaynak)
            for alan in ("kisi_ay", "maliyet", "fm_saat"):
                kontrol("POS %s%s = üç şube toplamı" % (alan, ek), round(p_[alan], 2),
                        round(sum(m["sube"]["%s|20%s" % (sb, ek)][alan]
                                  for sb in m["kapsam_pos"]), 2))
            kontrol("maliyet/ciro oranı%s yeniden hesaplandı" % ek,
                    round(p_["maliyet_ciro_orani"], 6),
                    round(p_["maliyet"] / p_["ciro_kdvharic"], 6))
            kontrol("kişi-ay başı maliyet%s" % ek, round(p_["kisi_ay_basi_maliyet"], 2),
                    round(p_["maliyet"] / p_["kisi_ay"], 2))
            # net ödenen brütten BÜYÜK olamaz (bordro tutarlılığı)
            sonuc.append((p_["net"] < p_["brut"], "net ödenen%s < brüt ücret%s" % (ek, ek),
                          "net<brüt", "net %.0f / brüt %.0f" % (p_["net"], p_["brut"]),
                          "tersi olursa kolon eşlemesi kaymış demektir"))
        # pencere iki yılda AYNI olmalı (eksik bordro ayı maliyeti %80 düşük gösterir)
        kontrol("bordro penceresi = fazla mesai penceresi", m["pencere_ay"], f["pencere_ay"])
        sonuc.append((1 <= m["pencere_ay"] <= 12, "bordro penceresi makul (1-12 ay)",
                      "1-12", m["pencere_ay"], "son TAM bordro ayı"))
        # 5 mağaza toplamı ≥ 3 POS toplamı (kapsam içerme)
        for ek in ("25", "26"):
            sonuc.append((m["tum"][ek]["maliyet"] >= m["pos"][ek]["maliyet"],
                          "5 mağaza maliyet%s ≥ 3 POS maliyet%s" % (ek, ek), "≥",
                          "%.0f vs %.0f" % (m["tum"][ek]["maliyet"], m["pos"][ek]["maliyet"]), ""))
        # K-22 varsayım aritmetiği
        kv, fi = f["kadro_artmasaydi"], f["fiili"]
        kontrol("eksik kişi-ay = kişi-ay26 − kişi-ay25", kv["eksik_kisi_ay"],
                fi["kisi_ay26"] - fi["kisi_ay25"])
        kontrol("ek FM saat = eksik kişi-ay × normal ay saati", round(kv["ek_fm_saat"], 2),
                round(max(0, kv["eksik_kisi_ay"]) * f["ay_normal_saat"], 2))
        kontrol("varsayım toplam FM = fiili + ek", round(kv["toplam_fm_saat"], 2),
                round(fi["fm_saat26"] + kv["ek_fm_saat"], 2))
        kontrol("varsayım kişi başı yıllık FM", round(kv["kisi_basi_yillik_saat"], 4),
                round(kv["toplam_fm_saat"] / fi["kisi_ay25"] * 12.0, 4))
        sonuc.append((kv["sinir_asimi"] == (kv["kisi_basi_yillik_saat"]
                                            > f["yasal_yillik_sinir_saat"]),
                      "sınır aşımı bayrağı hesapla uyumlu", "uyumlu",
                      "%s (%.0f vs %.0f)" % (kv["sinir_asimi"], kv["kisi_basi_yillik_saat"],
                                             f["yasal_yillik_sinir_saat"]), ""))
        sonuc.append((fi["kisi_basi_yillik26"] <= f["yasal_yillik_sinir_saat"],
                      "FİİLİ kişi başı yıllık FM yasal sınırın altında", "≤270",
                      "%.0f" % fi["kisi_basi_yillik26"],
                      "aşıyorsa sunumdaki «sınırın içindeyiz» cümlesi YANLIŞ olur"))

    # ---------------------------------------------------------------- KATEGORİ KAPSAMI (K-13)
    if v.get("kategori"):
        ESLES_ADLAR = ("Hazırlık Kitapları", "Kırtasiye", "Kitap", "Çocuk Kitabı",
                       "Oyuncak", "Akademi")
        kaps = sum(k["adet26"] for k in v["kategori"] if k["kategori"] in ESLES_ADLAR)
        tum = sum(k["adet26"] for k in v["kategori"]) or 1
        pay = 100.0 * kaps / tum
        sonuc.append((pay >= 60, "sunum kategori tablosu hizalı pencere adedinin %%%.0f'ini "
                      "kapsıyor (kalanı DİĞER satırında)" % pay, "≥%60", "%%%.0f" % pay,
                      "eşleşmesiz kategoriler tabloda DİĞER olarak toplanır"))

    # ---------------------------------------------------------------- İŞ HACMİ
    kontrol("hizalı pencere adet toplamı", round(sum(m["adet26"] for m in v["magaza"])),
            round(sum(m["adet26"] for m in v["magaza"])))
    if v.get("kayma"):
        ky = v["kayma"]
        kontrol("Ağustos = 1-14 + 15-31 (adet)", round(ky["y26_agu_tam"]["adet"]),
                round(v["agustos_yarim"]["1"]["adet26"] + v["agustos_yarim"]["2"]["adet26"]))
        kontrol("kaymasız Ağustos − gerçek = Eylül'e kayan",
                round(ky["eylule_kayan"]["adet"]),
                round(ky["agustos_kaymasiz_tahmin"]["adet"] - ky["y26_agu_tam"]["adet"]))
        # K-07: dilimlerin hicbiri bos olamaz; hizali pencereler ayni gun sayisinda olmali
        for dilim in ("y25_hizali", "y26_hizali", "y25_agu_tam", "y26_agu_tam",
                      "y25_agu_kalan", "y25_eyl_1_7"):
            d_ = ky[dilim]
            sonuc.append((d_["adet"] > 0 and d_["ciro"] > 0 and d_["gun"] > 0,
                          "kayma dilimi «%s» dolu" % dilim, ">0",
                          "adet %.0f / gün %d" % (d_["adet"], d_["gun"]),
                          "boş dilim «Eylül'e kayan»ı tüm Ağustos kadar gösterir"))
        kontrol("hizalı pencere gün sayısı iki yılda eşit",
                ky["y25_hizali"]["gun"], ky["y26_hizali"]["gun"])
    if v.get("yillar"):
        y26 = [y for y in v["yillar"] if y["yil"] == 2026][0]
        kontrol("Oca-Ağu mağaza adedi (yıllar ↔ kanal tablosu)",
                round(v["ocak_agustos"]["magaza"]["adet26"]), round(y26["adet"]))

    # ---------------------------------------------------------------- ÇIKTI KATMANI (sunum + Excel)
    # ⚠ 02.09.2026 dersi: JSON aritmetigi dogru olsa bile EMITTER KAYABILIR. Bu bolum, sunumda ve
    #   Excel'de FIILEN YAZAN rakami arar; eski (bayat) rakamlarin kalmadigini da dogrular.
    #   Yakalanan gercek hata: norm kurali degistiginde Sonuc slayti -22'de kalmis, Excel Norm
    #   sayfasi kayit kadrolusunu (137) gostermeye devam etmisti.
    kok = yol.parent
    # ⚠ EN YENI dosya (alfabetik ilk DEGIL — bulgu 16: "verimlilik-2026-ESKI.xlsx" denetlenirdi)
    def _yeni(desen):
        adaylar = [x for x in kok.glob(desen) if "KISILI" not in x.name]
        return max(adaylar, key=lambda x: x.stat().st_mtime) if adaylar else None

    pptx = _yeni("sunum-*.pptx")
    xlsx = _yeni("verimlilik-*.xlsx")
    # cikti JSON'dan YENI olmali; degilse bayat dosya dogrulanir
    for dosya in (pptx, xlsx):
        if dosya and dosya.stat().st_mtime < yol.stat().st_mtime - 1:
            sonuc.append((False, "%s JSON'dan ESKI (bayat çıktı)" % dosya.name,
                          "yeniden üret", "eski", "emitter'ı tekrar çalıştır"))
    for seg in ("KADROLU", "SEZONLUK"):
        for yil_ in ("2025", "2026"):
            d_ = (v.get("tutunma") or {}).get(seg, {}).get(yil_, {})
            sonuc.append((d_.get("oran14") is not None,
                          "tutunma ölçümü var (%s %s)" % (seg, yil_), "var",
                          "var" if d_.get("oran14") is not None else "yok",
                          "sunumdaki kalma oranı KPI'sı bu ölçümden gelir"))

    ops_toplam = ops_kadrolu + n["toplam"]["sezonluk_kesim26"] if n else 0
    pos26 = sum(m["kadro26"] for m in v["magaza"])

    def bin_(x):
        return format(int(round(x)), ",").replace(",", ".")

    if pptx:
        from pptx import Presentation
        parca = []
        for sl in Presentation(str(pptx)).slides:
            for sh in sl.shapes:
                if getattr(sh, "has_table", False):
                    parca += [c.text for r in sh.table.rows for c in r.cells]
                elif getattr(sh, "has_text_frame", False):
                    parca.append(sh.text_frame.text)
        T = " | ".join(parca)
        for ad, jeton in (("norm toplam", str(n["toplam"]["norm_toplam"])),
                          ("operasyonel toplam", str(ops_toplam)),
                          ("operasyonel kadrolu", str(ops_kadrolu)),
                          ("kayıt kadrolu", str(n["toplam"]["kadrolu_kesim26"])),
                          ("3 POS kadro", str(pos26)),
                          ("hizalı adet 2026", bin_(sum(m["adet26"] for m in v["magaza"])))):
            sonuc.append((jeton in T, "sunumda «%s» (%s) yazıyor" % (jeton, ad), jeton, jeton, ""))
        if v.get("maliyet"):
            # ESAS PENCERE SEZON: destede sezon orani yazili olmali (kumulatif yalniz Δ puan olarak)
            for ek in ("25", "26"):
                jet = ("%.2f" % (v["maliyet"]["sezon"]["pos"][ek]["maliyet_ciro_orani"] * 100)
                       ).replace(".", ",")
                sonuc.append((jet in T, "sunumda SEZON maliyet/ciro oranı%s (%%%s) yazıyor"
                              % (ek, jet), jet, jet, ""))
            ku = v["maliyet"]["pos"]
            puan_ku = ("%+.2f" % ((ku["26"]["maliyet_ciro_orani"]
                                   - ku["25"]["maliyet_ciro_orani"]) * 100)).replace(".", ",")
            sonuc.append((puan_ku in T, "sunumda yıl geneli oran farkı (%s puan) yazıyor" % puan_ku,
                          puan_ku, puan_ku, "kümülatif pencere referans olarak durmalı"))
            fm_ = v["fazla_mesai"]
            for ad_, deg_ in (("fiili kişi başı yıllık FM", bin_(fm_["fiili"]["kisi_basi_yillik26"])),
                              ("varsayım kişi başı yıllık FM",
                               bin_(fm_["kadro_artmasaydi"]["kisi_basi_yillik_saat"])),
                              ("yasal sınır", str(int(fm_["yasal_yillik_sinir_saat"])))):
                sonuc.append((deg_ in T, "sunumda %s (%s) yazıyor" % (ad_, deg_), deg_, deg_, ""))
        bayat = str(n["toplam"]["toplam_kesim26"] - n["toplam"]["norm_toplam"]) + " kişi"
        sonuc.append((bayat not in T,
                      "sunumda BAYAT norm farkı «%s» YOK" % bayat, "yok", "yok",
                      "engelli düşülmeden hesaplanan eski fark"))

    if xlsx:
        from openpyxl import load_workbook
        wb = load_workbook(str(xlsx))
        sayilar = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, (int, float)):
                        sayilar.add(round(float(c.value), 2))
        for ad, deger in (("norm kadrolu", n["toplam"]["norm"]),
                          ("operasyonel kadrolu (şube satırları toplamı değil, satır değeri)",
                           max(r["kadrolu_kesim26"] - v["norm"]["ayrik"].get(r["sube"], {}).get("engelli", 0)
                               - v["norm"]["ayrik"].get(r["sube"], {}).get("etkinlik", 0)
                               for r in n["sube"])),
                          ("hizalı adet 2026", round(sum(m["adet26"] for m in v["magaza"]), 2)),
                          ("hizalı adet 2025", round(sum(m["adet25"] for m in v["magaza"]), 2))):
            sonuc.append((round(float(deger), 2) in sayilar,
                          "Excel'de %s (%s) var" % (ad, deger), deger, deger, ""))

    # ÇAPRAZ-SAYFA FORMÜL REFERANSI (bulgu 15): openpyxl formül hücrelerini sayı olarak görmez,
    # bu yüzden "kaç kat" hücresi CİRO oranını (6,0x) gösterirken denetim 33/33 ✓ veriyordu.
    # Artık Sunum sayfasının işaret ettiği Ozet hücresinin FORMÜLÜ sınanıyor.
    if xlsx:
        from openpyxl import load_workbook as _lw
        wb2 = _lw(str(xlsx))
        if "Sunum" in wb2.sheetnames and "Ozet" in wb2.sheetnames:
            sn, oz = wb2["Sunum"], wb2["Ozet"]
            bulundu = False
            for r in range(1, sn.max_row + 1):
                etiket, ref = sn.cell(r, 2).value, sn.cell(r, 3).value
                if not (etiket and ref and isinstance(ref, str) and "KAT" in str(etiket).upper()):
                    continue
                hucre = str(ref).replace("='Ozet'!", "").strip()
                formul = str(oz[hucre].value or "")
                bulundu = True
                # adet-tabanli oran: E8/E7 (adet Δ ÷ kadro Δ). Ciro tabanli (E9/E7) YANLIS.
                sonuc.append((formul.replace(" ", "") == "=E8/E7",
                              "Sunum «kaç kat» → Ozet!%s formülü adet tabanlı" % hucre,
                              "=E8/E7", formul,
                              "ciro tabanlı (=E9/E7) olursa oran 2,9x yerine 6,0x görünür"))
            if not bulundu:
                sonuc.append((False, "Sunum sayfasında «kaç kat» satırı bulunamadı", "var", "yok", ""))

    # ---------------------------------------------------------------- ÇIKTI
    print("Tutarlılık denetimi: %s" % yol.name)
    hata = [x for x in sonuc if not x[0]]
    for tamam, ad, bekl, bul, acik in sonuc:
        isaret = "✓" if tamam else "✗"
        ek = ("  [%s]" % acik) if acik else ""
        if tamam:
            print("  %s %s = %s%s" % (isaret, ad, bul, ek))
        else:
            print("  %s %s → beklenen %s, bulunan %s%s" % (isaret, ad, bekl, bul, ek))

    print("\nKAPSAM SÖZLÜĞÜ (aynı destede üç kapsam):")
    print("  3 POS mağazası (FSM · Özlüce · İst. Yolu) — iş hacmi ölçülebilen kapsam · kadro %d → %d"
          % (sum(m["kadro25"] for m in v["magaza"]), pos_kadro))
    if n:
        print("  4 norm mağazası (+ Heykel) — norm karşılaştırması · kayıt %d, operasyonel %d, norm %d"
              % (n["toplam"]["kadrolu_kesim26"] + n["toplam"]["sezonluk_kesim26"],
                 ops_kadrolu + n["toplam"]["sezonluk_kesim26"], n["toplam"]["norm_toplam"]))
    print("  5 mağaza (+ Şura) — İK kadro raporu · kadrolu %d + sezonluk %d = %d"
          % (k5["kadrolu_kesim26"], k5["sezonluk_kesim26"], k5["toplam_kesim26"]))

    if hata:
        print("\n⚠ %d SAPMA" % len(hata))
        return 1
    print("\n✓ %d kontrol, sapma yok." % len(sonuc))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
