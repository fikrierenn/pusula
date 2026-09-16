# -*- coding: utf-8 -*-
"""DEFTER GRUBU STOK YETERLİLİĞİ — ÜRÜN DETAYLI, FORMÜLLÜ Excel.

GMY isteği 14.09.2026 (2. tur): _"ürün detaylı, senin çıkarımların değil de her şey
Excel'de görünecek şekilde formüllü olarak hazırlar mısın"_.

TASARIM İLKESİ — ÇIKARIM EXCEL'DE, PYTHON'DA DEĞİL:
  · Python YALNIZ ÖLÇÜLEN ham kolonları çeker (satış adetleri, snapshot stokları,
    bugünkü stok, depo stoğu, birim maliyet/fiyat). Hiçbir oran/açık/kapak
    Python'da hesaplanmaz.
  · Tüm türetmeler Excel FORMÜLÜ olarak yazılır ve `Parametreler` sayfasındaki
    hücrelere bağlanır. Büyüme çarpanını değiştiren sonucu anında görür.
  · `Özet` sayfasındaki her rakam SUMIFS/SUMPRODUCT ile `Ürün Detay`tan gelir —
    tek bir sabit sayı yazılmaz. Rapordaki iddialar böylece Excel'de denetlenir.

OKUL HİZALAMASI (kritik): okul açılışı 2025 → 08.09.2025, 2026 → 14.09.2026.
Takvim ayı kıyası Ağustos'ta −%20 gösterir, hizalı kıyas +%11'dir; yön ters.
Pencereler çıpaya göre kurulur, ay adına göre DEĞİL. Çıpalar Parametreler'de.

⚠ pyodbc kullanılıyor (pymssql DEĞİL): DerinSIS varchar kolonları CP1254;
pymssql Türkçe'yi bozuyor (coding-discipline § rapor scripti).

⚠ Merkez depo stoğu ERP defterinden OKUNMAZ (kullanıcı direktifi 03.09.2026) —
WMS hücresel stok (`depo.stok_adres_palet_vw`) kullanılır.

Kullanım:
    python scripts/defter_stok_yeterlilik_excel.py [--cikti yol.xlsx] [--grup Defterler]
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import os
import re
import sys

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

MEKAN_AD = {1: "FSM", 4477: "Özlüce", 4478: "İst.Yolu"}

# Okul açılış çıpaları — hizalı pencereler bunlardan türetilir
OKUL_2025 = dt.date(2025, 9, 8)
OKUL_2026 = dt.date(2026, 9, 14)
ON_GUN = 45          # ön-sezon penceresi: T−45..T−1
SEZON_GUN = 30       # sezon penceresi: T+0..T+30

KIRMIZI = PatternFill("solid", fgColor="E30622")
GRI = PatternFill("solid", fgColor="EFEFEF")
SARI = PatternFill("solid", fgColor="FFF3C4")
BASLIK_YAZI = Font(bold=True, color="FFFFFF", size=10)
KALIN = Font(bold=True)
SOLUK = Font(italic=True, color="808080")
INCE = Side(style="thin", color="BFBFBF")


def env_oku(yol: str) -> dict[str, str]:
    """`.env` dosyasını okur. Kimlik YALNIZ burada durur (tek kimlik yolu)."""
    if not os.path.exists(yol):
        sys.exit(f".env bulunamadi: {yol}")
    env: dict[str, str] = {}
    with open(yol, encoding="utf-8") as f:
        for ln in f:
            if ln.lstrip().startswith("#"):
                continue
            m = re.match(r"\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.+?)\s*$", ln)
            if m:
                env[m.group(1)] = m.group(2).strip().strip('"')
    return env


def baglan(env: dict[str, str]) -> pyodbc.Connection:
    host, port = env["MSSQL_HOST"], env.get("MSSQL_PORT", "1433")
    # ODBC connection-string'e env degeri gomulurken whitelist guard (injection)
    if not re.fullmatch(r"[A-Za-z0-9._\-]+", host) or not re.fullmatch(r"\d+", port):
        sys.exit("Gecersiz MSSQL_HOST/MSSQL_PORT (.env)")
    cn = pyodbc.connect(
        "Driver={ODBC Driver 18 for SQL Server};"
        f"Server={host},{port};Database=DerinSISBkm;"
        f"UID={env['MSSQL_USER']};PWD={env['MSSQL_PASSWORD']};"
        "TrustServerCertificate=yes;Timeout=30",
        timeout=30,
    )
    cn.timeout = 900
    return cn


def iso(g: dt.date) -> str:
    """Yerel sorgularda tarih ISO-bitişik yazılır (yyyy-MM-dd YASAK, DMY belirsiz)."""
    return g.strftime("%Y%m%d")


# =============================================================================
# HAM ÖLÇÜM — hiçbir oran/türetme YOK. Sadece sayılan şeyler.
# Satış tabanı: irsHrk ehTip 100 − 101 + 4 − 5 (kanonik şube formülü)
# Adet işareti: satış NEGATİF, iade POZİTİF → net adet = -SUM(ehAdetN)
# =============================================================================
SQL_DETAY = """
WITH evren AS (
    -- Hücre = mağaza × ürün. Herhangi bir pencerede satışı VEYA stoğu olan her hücre.
    SELECT ehMekan, ehstkID AS stkID FROM dbo.irsHrk WITH (NOLOCK)
      WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
        AND ehTrhS >= ? AND ehTrhS < ?
    UNION
    SELECT ehMekan, ehstkID FROM dbo.irsHrk WITH (NOLOCK)
      WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
        AND ehTrhS >= ? AND ehTrhS < ?
    UNION
    SELECT ehMekan, stkID FROM bkm.StokAyBakiyeMekanBazli WITH (NOLOCK)
      WHERE Donem IN (?, ?) AND Kaynak = 'irsHrk' AND ehMekan IN (1,4477,4478) AND Stok <> 0
    UNION
    SELECT ehMekan, ehstkID FROM dbo.stokSonAltDepo_vw WITH (NOLOCK)
      WHERE ehMekan IN (1,4477,4478) AND stok <> 0
),
gy_on AS (   -- geçen yıl ön-sezon (T−45..T−1)
    SELECT ehMekan, ehstkID, -SUM(ehAdetN) AS adet, SUM(ehTutarN) AS tutar
    FROM dbo.irsHrk WITH (NOLOCK)
    WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
      AND ehTrhS >= ? AND ehTrhS < ?
    GROUP BY ehMekan, ehstkID
),
gy_sezon AS (   -- geçen yıl sezon (T+0..T+30) — beklentinin TABANI
    SELECT ehMekan, ehstkID, -SUM(ehAdetN) AS adet, SUM(ehTutarN) AS tutar
    FROM dbo.irsHrk WITH (NOLOCK)
    WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
      AND ehTrhS >= ? AND ehTrhS < ?
    GROUP BY ehMekan, ehstkID
),
by_on_urun AS (   -- bu yıl ön-sezon, ÜRÜN bazında (üç şube toplamı)
    -- 'Canlı mı' testi ürün düzeyinde kurulur: ürün HERHANGİ bir şubede satıyorsa
    -- hâlâ koleksiyondadır. Şube bazlı bakmak "bu şubede satmamış"ı "bırakılmış"
    -- sanır ve açığı sessizce küçültür (ölçüldü: 1.584 → 1.318 satır).
    SELECT ehstkID, -SUM(ehAdetN) AS adet
    FROM dbo.irsHrk WITH (NOLOCK)
    WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
      AND ehTrhS >= ? AND ehTrhS < ?
    GROUP BY ehstkID
),
by_on AS (   -- bu yıl ön-sezon, MAĞAZA bazında (T−45..T−1)
    SELECT ehMekan, ehstkID, -SUM(ehAdetN) AS adet, SUM(ehTutarN) AS tutar,
           MAX(ehTrhS) AS son_satis
    FROM dbo.irsHrk WITH (NOLOCK)
    WHERE ehTip IN (4,5,100,101) AND ehMekan IN (1,4477,4478)
      AND ehTrhS >= ? AND ehTrhS < ?
    GROUP BY ehMekan, ehstkID
),
giris25 AS (SELECT ehMekan, stkID, Stok FROM bkm.StokAyBakiyeMekanBazli WITH (NOLOCK)
            WHERE Donem = ? AND Kaynak = 'irsHrk' AND ehMekan IN (1,4477,4478)),
giris26 AS (SELECT ehMekan, stkID, Stok FROM bkm.StokAyBakiyeMekanBazli WITH (NOLOCK)
            WHERE Donem = ? AND Kaynak = 'irsHrk' AND ehMekan IN (1,4477,4478)),
bugun AS (SELECT ehMekan, ehstkID, SUM(stok) AS stok FROM dbo.stokSonAltDepo_vw WITH (NOLOCK)
          WHERE ehMekan IN (1,4477,4478) GROUP BY ehMekan, ehstkID),
depo AS (   -- merkez depo = WMS hücresel stok (ERP defteri DEĞİL)
    SELECT stkID,
           SUM(CASE WHEN adrsAlanTipID = 0 THEN Stok ELSE 0 END) AS raf,
           SUM(CASE WHEN adrsAlanTipID = 1 THEN Stok ELSE 0 END) AS giris_alani,
           SUM(CASE WHEN adrsAlanTipID = 2 THEN Stok ELSE 0 END) AS cikis_alani
    FROM depo.stok_adres_palet_vw WITH (NOLOCK)
    GROUP BY stkID
)
SELECT
    e.ehMekan,
    e.stkID,
    u.stkKod,
    LEFT(ISNULL(u.stkAd, '(ad yok)'), 90)       AS UrunAd,
    ISNULL(u.Kat2, '(alt grup yok)')            AS AltGrup,
    ISNULL(u.mrkAd, '(marka yok)')              AS Marka,
    ISNULL(u.FirmaAd, '(tedarikçi yok)')        AS Tedarikci,
    ISNULL(go.adet, 0)                          AS GyOnAdet,
    ISNULL(gs.adet, 0)                          AS GySezonAdet,
    ISNULL(gs.tutar, 0)                         AS GySezonTutar,
    ISNULL(bo.adet, 0)                          AS ByOnAdet,
    ISNULL(bou.adet, 0)                         AS ByOnUrunAdet,
    ISNULL(bo.tutar, 0)                         AS ByOnTutar,
    bo.son_satis                                AS SonSatis,
    ISNULL(g25.Stok, 0)                         AS Giris2025,
    ISNULL(g26.Stok, 0)                         AS Giris2026,
    ISNULL(bg.stok, 0)                          AS BugunStok,
    ISNULL(d.raf, 0)                            AS DepoRaf,
    ISNULL(d.giris_alani, 0)                    AS DepoGiris,
    ISNULL(d.cikis_alani, 0)                    AS DepoCikis,
    ISNULL(u.SonAlis, 0)                        AS BirimMaliyet,
    ISNULL(u.SatisFiyat, 0)                     AS SatisFiyat
FROM evren e
JOIN bkm.UrunBilgi u ON u.stkID = e.stkID AND u.Kat1 = ?
LEFT JOIN gy_on    go  ON go.ehMekan  = e.ehMekan AND go.ehstkID  = e.stkID
LEFT JOIN gy_sezon gs  ON gs.ehMekan  = e.ehMekan AND gs.ehstkID  = e.stkID
LEFT JOIN by_on    bo  ON bo.ehMekan  = e.ehMekan AND bo.ehstkID  = e.stkID
LEFT JOIN by_on_urun bou ON bou.ehstkID = e.stkID
LEFT JOIN giris25  g25 ON g25.ehMekan = e.ehMekan AND g25.stkID   = e.stkID
LEFT JOIN giris26  g26 ON g26.ehMekan = e.ehMekan AND g26.stkID   = e.stkID
LEFT JOIN bugun    bg  ON bg.ehMekan  = e.ehMekan AND bg.ehstkID  = e.stkID
LEFT JOIN depo     d   ON d.stkID     = e.stkID
ORDER BY e.ehMekan, e.stkID
"""

# Ürün Detay kolon düzeni — (başlık, tip, biçim, genişlik)
#   tip 'ham'    = veritabanından gelen ÖLÇÜM
#   tip 'formul' = Excel'de hesaplanan ÇIKARIM (şablon {r} ile satır numarası)
KOLONLAR = [
    ("Mağaza",                 "ham",    None,          10),
    ("stkID",                  "ham",    "0",           9),
    ("Stok kodu",              "ham",    None,          14),
    ("Ürün",                   "ham",    None,          46),
    ("Alt grup",               "ham",    None,          22),
    ("Marka",                  "ham",    None,          18),
    ("Tedarikçi",              "ham",    None,          22),
    # --- ölçülen satış ---
    ("GY ön-sezon adet",       "ham",    "#,##0",       11),
    ("GY sezon adet",          "ham",    "#,##0",       11),
    ("GY sezon ₺",             "ham",    "#,##0",       12),
    ("BY ön-sezon adet",       "ham",    "#,##0",       11),
    ("BY ön-sezon TÜM ŞUBE",   "ham",    "#,##0",       12),
    ("BY ön-sezon ₺",          "ham",    "#,##0",       12),
    ("Son satış",              "ham",    "DD.MM.YYYY",  11),
    # --- ölçülen stok ---
    ("Giriş stoğu 31.08.2025", "ham",    "#,##0",       12),
    ("Giriş stoğu 31.08.2026", "ham",    "#,##0",       12),
    ("Bugün mağaza stoğu",     "ham",    "#,##0",       12),
    ("Depo RAF",               "ham",    "#,##0",       10),
    ("Depo GİRİŞ",             "ham",    "#,##0",       10),
    ("Depo ÇIKIŞ",             "ham",    "#,##0",       10),
    ("Birim maliyet ₺",        "ham",    "#,##0.00",    12),
    ("Satış fiyatı ₺",         "ham",    "#,##0.00",    12),
    # --- FORMÜLLER (hepsi Parametreler'e bağlı) ---
    ("Beklenen kalan sezon",   "formul", "#,##0.0",     13),
    ("Açık adet",              "formul", "#,##0.0",     11),
    ("Kapak (kat)",            "formul", "0.00",        10),
    ("Depoda toplam",          "formul", "#,##0",       11),
    ("Depodan karşılanır",     "formul", "#,##0.0",     13),
    ("Satın alma gereken",     "formul", "#,##0.0",     13),
    ("Açık maliyeti ₺",        "formul", "#,##0",       12),
    ("Açık satış kaybı ₺",     "formul", "#,##0",       13),
    ("Aksiyon",                "formul", None,          20),
    ("Canlı mı",               "formul", None,          14),
    ("Zirvede sattı mı",       "formul", None,          14),
    ("Kendi büyümesi",         "formul", "0.0%",        12),
    ("Giriş stoğu değişimi",   "formul", "#,##0",       13),
    ("Atıl stok ₺",            "formul", "#,##0",       12),
]

# Formül şablonları — kolon harfleri KOLONLAR sırasına göre hesaplanır (aşağıda)
def formul_sablonlari(K: dict[str, str]) -> dict[str, str]:
    """K: başlık → kolon harfi. Formüller okunur kalsın diye adlandırılmış."""
    gy_sezon = K["GY sezon adet"]
    gy_on = K["GY ön-sezon adet"]
    by_on = K["BY ön-sezon adet"]
    by_on_urun = K["BY ön-sezon TÜM ŞUBE"]
    bugun = K["Bugün mağaza stoğu"]
    raf, dgiris, maliyet, fiyat = K["Depo RAF"], K["Depo GİRİŞ"], K["Birim maliyet ₺"], K["Satış fiyatı ₺"]
    g25, g26 = K["Giriş stoğu 31.08.2025"], K["Giriş stoğu 31.08.2026"]
    beklenen, acik = K["Beklenen kalan sezon"], K["Açık adet"]
    depo_top, karsilanir = K["Depoda toplam"], K["Depodan karşılanır"]
    return {
        # beklenen = geçen yılın AYNI SKU-mağaza sezon adedi × büyüme çarpanı
        "Beklenen kalan sezon": f"={gy_sezon}{{r}}*Buyume",
        "Açık adet":            f"=MAX(0,{beklenen}{{r}}-{bugun}{{r}})",
        "Kapak (kat)":          f'=IFERROR({bugun}{{r}}/{beklenen}{{r}},"")',
        "Depoda toplam":        f"={raf}{{r}}+{dgiris}{{r}}",
        "Depodan karşılanır":   f"=MIN({acik}{{r}},MAX(0,{depo_top}{{r}}))",
        "Satın alma gereken":   f"={acik}{{r}}-{karsilanir}{{r}}",
        "Açık maliyeti ₺":      f"={acik}{{r}}*{maliyet}{{r}}",
        "Açık satış kaybı ₺":   f"={acik}{{r}}*{fiyat}{{r}}",
        "Aksiyon": (
            f'=IF({acik}{{r}}<=0,"YETERLİ",'
            f'IF({depo_top}{{r}}>={acik}{{r}},"TRANSFERLE KAPANIR",'
            f'IF({depo_top}{{r}}>0,"KISMİ TRANSFER","SATIN ALMA GEREKİR")))'
        ),
        # canlı = depoda stok var VEYA bu yıl ön-sezonda satmış (bırakılmış tasarım değil)
        "Canlı mı":             f'=IF(OR({depo_top}{{r}}>0,{by_on_urun}{{r}}>0),"CANLI","BIRAKILMIS?")',
        "Zirvede sattı mı":     f'=IF({by_on}{{r}}>0,"SATTI","SATMADI")',
        "Kendi büyümesi":       f'=IFERROR({by_on}{{r}}/{gy_on}{{r}}-1,"")',
        "Giriş stoğu değişimi": f"={g26}{{r}}-{g25}{{r}}",
        # atıl = zirvede hiç satmamış ama rafta duran stoğun maliyeti
        "Atıl stok ₺":          f"=IF({by_on}{{r}}<=0,MAX(0,{bugun}{{r}})*{maliyet}{{r}},0)",
    }


# =============================================================================
# FORMÜL REFERANS DENETİMİ
#
# Neden gerekli: openpyxl formülü YAZAR, HESAPLAMAZ. Yanlış bir kolon harfi
# (kolon sırası değişince olur) Excel'de hata vermez — tipler uyuşursa DEĞER
# SESSİZCE KAYAR. Dapper pozisyonel record vakasının Excel ikizi
# (sql-server-conventions § Dapper POZİSYONEL RECORD).
#
# Aşağıdaki tablo NİYETİN İKİNCİ BEYANIDIR: her formül kolonunun hangi
# başlıklara bakması gerektiği. Kolon sırası kayarsa iki beyan ayrışır ve
# denetim KIRIK der. Çıkış 0 geçti · 1 KIRIK · 2 KOŞAMADI.
# =============================================================================
BEKLENEN_REF = {
    "Beklenen kalan sezon":   {"GY sezon adet"},
    "Açık adet":              {"Beklenen kalan sezon", "Bugün mağaza stoğu"},
    "Kapak (kat)":            {"Bugün mağaza stoğu", "Beklenen kalan sezon"},
    "Depoda toplam":          {"Depo RAF", "Depo GİRİŞ"},
    "Depodan karşılanır":     {"Açık adet", "Depoda toplam"},
    "Satın alma gereken":     {"Açık adet", "Depodan karşılanır"},
    "Açık maliyeti ₺":        {"Açık adet", "Birim maliyet ₺"},
    "Açık satış kaybı ₺":     {"Açık adet", "Satış fiyatı ₺"},
    "Aksiyon":                {"Açık adet", "Depoda toplam"},
    "Canlı mı":               {"Depoda toplam", "BY ön-sezon TÜM ŞUBE"},
    "Zirvede sattı mı":       {"BY ön-sezon adet"},
    "Kendi büyümesi":         {"BY ön-sezon adet", "GY ön-sezon adet"},
    "Giriş stoğu değişimi":   {"Giriş stoğu 31.08.2026", "Giriş stoğu 31.08.2025"},
    "Atıl stok ₺":            {"BY ön-sezon adet", "Bugün mağaza stoğu", "Birim maliyet ₺"},
}
# Parametreye bağlı OLMASI GEREKEN formüller — bağ koparsa çarpan oynatılamaz
BUYUME_BEKLENEN = {"Beklenen kalan sezon"}


def denetle(yol: str) -> int:
    from openpyxl import load_workbook

    if not os.path.exists(yol):
        print(f"KOSAMADI: dosya yok — {yol}")
        return 2
    try:
        wb = load_workbook(yol, data_only=False)
    except Exception as e:  # bozuk dosya = koşamadı, yeşil DEĞİL
        print(f"KOSAMADI: dosya acilamadi — {e}")
        return 2
    if "Ürün Detay" not in wb.sheetnames:
        print("KOSAMADI: 'Ürün Detay' sayfasi yok")
        return 2

    wd = wb["Ürün Detay"]
    basliklar = [c.value for c in wd[1]]
    harf_basi = {get_column_letter(i): b for i, b in enumerate(basliklar, start=1)}
    basi_harf = {b: get_column_letter(i) for i, b in enumerate(basliklar, start=1)}
    son = wd.max_row
    if son < 2:
        print("KOSAMADI: 'Ürün Detay' bos (nufus 0 — 'ihlal yok' DEGIL 'bakamadim')")
        return 2

    kirik: list[str] = []
    uyari: list[str] = []
    hucre_re = re.compile(r"(?<![A-Z0-9$!])\$?([A-Z]{1,3})\$?(\d+)")

    # --- 1) Ürün Detay formül kolonları
    for baslik, bekl in BEKLENEN_REF.items():
        if baslik not in basi_harf:
            kirik.append(f"'{baslik}' kolonu dosyada YOK (BEKLENEN_REF ile sema ayrismis)")
            continue
        h = basi_harf[baslik]
        f = wd[f"{h}2"].value
        if not isinstance(f, str) or not f.startswith("="):
            kirik.append(f"'{baslik}' ({h}2) formul DEGIL: {f!r}")
            continue
        gorulen = {harf_basi.get(m.group(1)) for m in hucre_re.finditer(f)
                   if m.group(2) == "2"}
        gorulen.discard(None)
        if gorulen != bekl:
            kirik.append(
                f"'{baslik}' ({h}2) YANLIS REFERANS\n"
                f"      beklenen: {sorted(bekl)}\n"
                f"      gorulen : {sorted(gorulen)}\n"
                f"      formul  : {f}"
            )
        if baslik in BUYUME_BEKLENEN and "Buyume" not in f:
            kirik.append(f"'{baslik}' ({h}2) 'Buyume' parametresine BAGLI DEGIL: {f}")
        # son satirda da aynı formül olmalı (araya elle satır girilmiş olabilir)
        fs = wd[f"{h}{son}"].value
        if not isinstance(fs, str) or not fs.startswith("="):
            kirik.append(f"'{baslik}' son satirda ({h}{son}) formul yok: {fs!r}")

    # --- 2) Buyume defined name
    if "Buyume" not in wb.defined_names:
        kirik.append("'Buyume' adlandirilmis alani YOK — Ozet/Kirilim formulleri kirilir")
    else:
        at = wb.defined_names["Buyume"].attr_text
        if "Parametreler" not in at:
            kirik.append(f"'Buyume' Parametreler sayfasina bakmiyor: {at}")

    # --- 3) Özet/Mağaza/Kırılım: Ürün Detay'a bakan aralıklar geçerli mi
    aralik_re = re.compile(r"'Ürün Detay'!\$([A-Z]{1,3})\$(\d+):\$([A-Z]{1,3})\$(\d+)")
    formul_say = 0
    for ad in wb.sheetnames:
        if ad in ("Ürün Detay", "Parametreler", "Okuma Kılavuzu"):
            continue
        ws = wb[ad]
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value.startswith("="):
                    continue
                formul_say += 1
                for m in aralik_re.finditer(c.value):
                    h1, r1, h2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                    if h1 != h2:
                        kirik.append(f"{ad}!{c.coordinate}: aralik iki KOLONA yayilmis ({h1}:{h2}) — SUMIFS bozulur")
                    if h1 not in harf_basi:
                        kirik.append(f"{ad}!{c.coordinate}: '{h1}' kolonu Ürün Detay'da YOK")
                    if (r1, r2) != (2, son):
                        kirik.append(
                            f"{ad}!{c.coordinate}: aralik {h1}{r1}:{h2}{r2} ama veri 2..{son} "
                            f"— {'EKSIK SATIR (sessiz alt-tahmin)' if r2 < son else 'FAZLA SATIR'}"
                        )
    if formul_say == 0:
        print("KOSAMADI: ozet sayfalarinda hic formul bulunamadi")
        return 2

    # --- 4) Sabit sayı kaçağı: Özet'te formül yerine düz rakam kalmış mı
    wo = wb["Özet"] if "Özet" in wb.sheetnames else None
    if wo is not None:
        for row in wo.iter_rows(min_col=2, max_col=2):
            for c in row:
                if isinstance(c.value, (int, float)):
                    uyari.append(f"Özet!{c.coordinate} SABIT SAYI ({c.value}) — formul olmali")

    # --- 5) BEKLENEN DEĞERLER: dosyanın kendi HAM kolonlarından bağımsız yeniden hesap
    # Excel'in formül motoru bu ortamda YOK (Excel COM / LibreOffice kurulu değil).
    # Bu blok formüllerin VERMESİ GEREKEN sayıları yazar; dosya açıldığında Özet
    # sayfası bunlarla aynı değilse formül hatalıdır. "Hesaplanmadı" ≠ "doğru".
    wbd = load_workbook(yol, read_only=True, data_only=True)
    ham = wbd["Ürün Detay"]
    it = ham.iter_rows(values_only=True)
    J = {b: i for i, b in enumerate(list(next(it)))}
    try:
        g = float(wbd["Parametreler"]["B4"].value or 0)
    except Exception:
        g = 0.0
    if g <= 0:
        print("KOSAMADI: Parametreler!B4 (Buyume) okunamadi")
        return 2
    b = dict.fromkeys(
        ["g25", "g26", "c25", "c26", "gyon", "byon", "gysez", "bugun", "bekl",
         "acik", "acikTL", "karsi", "satin", "satmaz", "atilTL", "nAcik", "birak"], 0.0)
    for row in it:
        g25, g26 = row[J["Giris stogu 31.08.2025".replace("Giris stogu", "Giriş stoğu")]], row[J["Giriş stoğu 31.08.2026"]]
        gyon, byon = row[J["GY ön-sezon adet"]], row[J["BY ön-sezon adet"]]
        byu = row[J["BY ön-sezon TÜM ŞUBE"]]
        gysez, bug = row[J["GY sezon adet"]], row[J["Bugün mağaza stoğu"]]
        depo = row[J["Depo RAF"]] + row[J["Depo GİRİŞ"]]
        mal = row[J["Birim maliyet ₺"]]
        bekl = gysez * g
        acik = max(0.0, bekl - bug)
        canli = depo > 0 or byu > 0
        b["g25"] += g25 if g25 > 0 else 0
        b["g26"] += g26 if g26 > 0 else 0
        b["c25"] += 1 if g25 > 0 else 0
        b["c26"] += 1 if g26 > 0 else 0
        b["gyon"] += gyon
        b["byon"] += byon
        b["gysez"] += gysez
        b["bugun"] += bug if bug > 0 else 0
        b["bekl"] += bekl
        if acik > 0 and canli:
            k = min(acik, max(0.0, depo))
            b["nAcik"] += 1
            b["acik"] += acik
            b["acikTL"] += acik * mal
            b["karsi"] += k
            b["satin"] += acik - k
        if acik > 0 and not canli:
            b["birak"] += 1
        if byon <= 0 and bug > 0:
            b["satmaz"] += bug
            b["atilTL"] += bug * mal

    print(f"Denetlenen: {yol}")
    print(f"  Ürün Detay satiri : {son - 1}")
    print(f"  ham kolon         : {len(basi_harf) - len(BEKLENEN_REF)}")
    print(f"  denetlenen formul : {len(BEKLENEN_REF)} kolon + {formul_say} ozet/kirilim formulu")
    for u in uyari:
        print(f"  UYARI: {u}")
    if kirik:
        print(f"\nKIRIK ({len(kirik)}):")
        for k in kirik:
            print(f"  - {k}")
        return 1
    print("")
    print(f"  FORMULLERIN VERMESI GEREKEN DEGERLER (buyume={g:.2f})")
    d25 = b["g25"] / max(b["c25"], 1)
    d26 = b["g26"] / max(b["c26"], 1)
    print(f"    giris stogu 2025 / 2026       : {b['g25']:>12,.0f} / {b['g26']:>12,.0f}  "
          f"({b['g26'] / max(b['g25'], 1) - 1:+.1%})")
    print(f"    stoklu cesit-magaza 25 / 26   : {b['c25']:>12,.0f} / {b['c26']:>12,.0f}")
    print(f"    cesit basi derinlik 25 / 26   : {d25:>12.1f} / {d26:>12.1f}  "
          f"({d26 / max(d25, 1e-9) - 1:+.1%})")
    print(f"    on-sezon adet GY / BY         : {b['gyon']:>12,.0f} / {b['byon']:>12,.0f}  "
          f"({b['byon'] / max(b['gyon'], 1) - 1:+.1%})")
    print(f"    GY sezon (beklenti tabani)    : {b['gysez']:>12,.0f}")
    print(f"    beklenen kalan sezon          : {b['bekl']:>12,.0f}")
    print(f"    bugun magaza stogu            : {b['bugun']:>12,.0f}")
    print(f"    magaza kapagi (kat)           : {b['bugun'] / max(b['bekl'], 1):>12.2f}")
    print(f"    acik satir (canli)            : {b['nAcik']:>12,.0f}")
    print(f"    acik adet / maliyet TL        : {b['acik']:>12,.0f} / {b['acikTL']:>12,.0f}")
    print(f"    depodan karsilanir / satinalma: {b['karsi']:>12,.0f} / {b['satin']:>12,.0f}")
    print(f"    birakilmis gorunen acik satir : {b['birak']:>12,.0f}")
    print(f"    zirvede satmayan stok / TL    : {b['satmaz']:>12,.0f} / {b['atilTL']:>12,.0f}")
    print("")
    print("  UYARI: Excel formulleri BU ORTAMDA HESAPLANMADI (Excel COM / LibreOffice yok).")
    print("    Yukaridakiler dosyanin HAM kolonlarindan BAGIMSIZ yeniden hesaptir.")
    print("    Dosya acilinca Ozet sayfasi bunlarla AYNI olmali; degilse formul hatalidir.")
    print("\nGECTI: tum formul referanslari beklenen kolonlara bakiyor.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cikti", default=None)
    ap.add_argument("--grup", default="Defterler", help="bkm.UrunBilgi.Kat1 grup adı")
    ap.add_argument("--denetle", metavar="XLSX", default=None,
                    help="Üretilmiş dosyanın formül referanslarını denetle (DB'ye bağlanmaz). "
                         "Çıkış 0 geçti · 1 KIRIK · 2 KOŞAMADI.")
    a = ap.parse_args()

    if a.denetle:
        return denetle(a.denetle)

    cikti = a.cikti or os.path.join(
        KOK, "ciktilar",
        f"{a.grup.lower()}-stok-yeterlilik-formullu-{dt.date.today():%Y-%m-%d}.xlsx",
    )
    os.makedirs(os.path.dirname(cikti), exist_ok=True)

    gy_on_b, gy_on_s = OKUL_2025 - dt.timedelta(days=ON_GUN), OKUL_2025
    gy_sz_b, gy_sz_s = OKUL_2025, OKUL_2025 + dt.timedelta(days=SEZON_GUN + 1)
    by_on_b, by_on_s = OKUL_2026 - dt.timedelta(days=ON_GUN), OKUL_2026
    snap25, snap26 = iso(dt.date(2025, 8, 31)), iso(dt.date(2026, 8, 31))

    env = env_oku(os.path.join(KOK, ".env"))
    cn = baglan(env)
    try:
        with cn.cursor() as cur:
            satirlar = cur.execute(
                SQL_DETAY,
                iso(gy_on_b), iso(gy_sz_s),      # evren: geçen yıl tüm pencere
                iso(by_on_b), iso(by_on_s),      # evren: bu yıl ön-sezon
                snap25, snap26,                  # evren: iki snapshot
                iso(gy_on_b), iso(gy_on_s),      # gy_on
                iso(gy_sz_b), iso(gy_sz_s),      # gy_sezon
                iso(by_on_b), iso(by_on_s),      # by_on_urun (ürün bazlı)
                iso(by_on_b), iso(by_on_s),      # by_on (mağaza bazlı)
                snap25, snap26,                  # giris25 / giris26
                a.grup,
            ).fetchall()
    finally:
        cn.close()

    if not satirlar:
        # Sessizlik kanıt değil: boş dönmek "veri yok" değil "ölçemedim" olabilir.
        sys.exit(f"HATA: '{a.grup}' grubu icin hic satir donmedi — grup adi dogru mu?")

    K = {b: get_column_letter(i) for i, (b, *_) in enumerate(KOLONLAR, start=1)}
    SAB = formul_sablonlari(K)

    wb = Workbook()

    # ----------------------------------------------------------------- PARAMETRELER
    wp = wb.active
    wp.title = "Parametreler"
    wp["A1"] = "OYNATILABİLİR PARAMETRELER — değiştir, tüm sayfalar yeniden hesaplanır"
    wp["A1"].font = Font(bold=True, size=13)
    wp["A3"], wp["B3"] = "Parametre", "Değer"
    wp["C3"] = "Ne işe yarar"
    for h in ("A3", "B3", "C3"):
        wp[h].fill, wp[h].font = KIRMIZI, BASLIK_YAZI

    wp["A4"] = "Büyüme çarpanı"
    wp["B4"] = 1.11
    wp["B4"].fill, wp["B4"].font = SARI, KALIN
    wp["B4"].number_format = "0.00"
    wp["C4"] = ("Beklenen kalan sezon = geçen yılın aynı ürün-mağaza sezon adedi × BU SAYI. "
                "Ölçüldü: okula hizalı ön-sezon adedi +%11,0 (1,11). "
                "Denemek için: 1,00 (büyüme yok) · 1,26 (şirket geneli adet büyümesi) · 1,50 (üst uç).")

    wp["A5"], wp["B5"] = "Okul açılışı — geçen yıl", OKUL_2025
    wp["A6"], wp["B6"] = "Okul açılışı — bu yıl", OKUL_2026
    for h in ("B5", "B6"):
        wp[h].number_format = "DD.MM.YYYY"
    wp["C5"] = wp["C6"] = "Pencereler bu çıpalara göre kuruldu (bilgi — değiştirmek sorguyu yeniden koşturmayı gerektirir)."

    wp["A7"], wp["B7"] = "Ön-sezon penceresi (gün)", ON_GUN
    wp["C7"] = "T−45 .. T−1 (okuldan önceki 45 gün)"
    wp["A8"], wp["B8"] = "Sezon penceresi (gün)", SEZON_GUN
    wp["C8"] = "T+0 .. T+30 (okul günü ve sonraki 30 gün) — beklentinin tabanı"

    wp["A10"] = "ÖLÇÜLEN PENCERELER (sorgu bunlarla koştu)"
    wp["A10"].font = KALIN
    for i, (ad, b, s) in enumerate(
        [
            ("Geçen yıl ön-sezon", gy_on_b, gy_on_s - dt.timedelta(days=1)),
            ("Geçen yıl sezon", gy_sz_b, gy_sz_s - dt.timedelta(days=1)),
            ("Bu yıl ön-sezon", by_on_b, by_on_s - dt.timedelta(days=1)),
        ],
        start=11,
    ):
        wp.cell(row=i, column=1, value=ad)
        wp.cell(row=i, column=2, value=f"{b:%d.%m.%Y} – {s:%d.%m.%Y}")

    wp["A15"] = "Kesim anı (stok fotoğrafı)"
    wp["B15"] = dt.datetime.now().replace(microsecond=0)
    wp["B15"].number_format = "DD.MM.YYYY HH:MM"
    wp["C15"] = ("Mağaza stoğu BU AN okundu. Kasa satışı ERP'ye saatte bir akar ve "
                 "sezon günlerinde stok saat saat düşer — bu dosya bir FOTOĞRAFTIR.")
    wp["A16"], wp["B16"] = "Ürün grubu (Kat1)", a.grup
    wp["A17"], wp["B17"] = "Satır sayısı (mağaza × ürün)", len(satirlar)
    wp["A18"] = "Satış tabanı"
    wp["B18"] = "irsHrk ehTip 100 − 101 + 4 − 5 (kanonik şube formülü), KDV hariç net"
    wp["A19"] = "Mağaza stoğu kaynağı"
    wp["B19"] = "dbo.stokSonAltDepo_vw (anlık) · bkm.StokAyBakiyeMekanBazli (31.08 snapshot)"
    wp["A20"] = "Merkez depo kaynağı"
    wp["B20"] = "depo.stok_adres_palet_vw (WMS) — ERP defteri KULLANILMADI (senkron sorunu)"

    wp.column_dimensions["A"].width = 32
    wp.column_dimensions["B"].width = 30
    wp.column_dimensions["C"].width = 110
    for r in range(4, 21):
        wp.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")

    wb.defined_names.add(DefinedName("Buyume", attr_text="Parametreler!$B$4"))

    # ----------------------------------------------------------------- ÜRÜN DETAY
    wd = wb.create_sheet("Ürün Detay")
    wd.append([b for b, *_ in KOLONLAR])
    for i, (b, tip, *_) in enumerate(KOLONLAR, start=1):
        c = wd.cell(row=1, column=i)
        c.fill = KIRMIZI if tip == "ham" else PatternFill("solid", fgColor="7A1020")
        c.font = BASLIK_YAZI
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=INCE)
    wd.row_dimensions[1].height = 34

    for r in satirlar:
        wd.append([
            MEKAN_AD.get(r.ehMekan, r.ehMekan), r.stkID, r.stkKod, r.UrunAd,
            r.AltGrup, r.Marka, r.Tedarikci,
            float(r.GyOnAdet), float(r.GySezonAdet), float(r.GySezonTutar),
            float(r.ByOnAdet), float(r.ByOnUrunAdet), float(r.ByOnTutar),
            r.SonSatis.date() if r.SonSatis else None,
            float(r.Giris2025), float(r.Giris2026), float(r.BugunStok),
            float(r.DepoRaf), float(r.DepoGiris), float(r.DepoCikis),
            float(r.BirimMaliyet), float(r.SatisFiyat),
        ] + [None] * sum(1 for _, t, *_ in KOLONLAR if t == "formul"))

    son = wd.max_row
    for i, (b, tip, bic, _) in enumerate(KOLONLAR, start=1):
        harf = get_column_letter(i)
        if tip == "formul":
            sab = SAB[b]
            for rr in range(2, son + 1):
                wd.cell(row=rr, column=i, value=sab.format(r=rr))
        if bic:
            for rr in range(2, son + 1):
                wd.cell(row=rr, column=i).number_format = bic
        wd.column_dimensions[harf].width = KOLONLAR[i - 1][3]

    wd.freeze_panes = "E2"
    wd.auto_filter.ref = f"A1:{get_column_letter(len(KOLONLAR))}{son}"

    # ----------------------------------------------------------------- ÖZET (tamamı formül)
    wo = wb.create_sheet("Özet", 1)
    D = "'Ürün Detay'!"

    def rng(baslik: str) -> str:
        return f"{D}${K[baslik]}$2:${K[baslik]}${son}"

    acik = rng("Açık adet")
    canli = rng("Canlı mı")
    magaza = rng("Mağaza")
    aksiyon = rng("Aksiyon")

    wo["A1"] = f"{a.grup.upper()} — STOK YETERLİLİĞİ ÖZETİ"
    wo["A1"].font = Font(bold=True, size=14)
    wo["A2"] = ("Bu sayfadaki HER SAYI 'Ürün Detay' sayfasından formülle gelir — sabit "
                "yazılmış rakam yoktur. Parametreler!B4'ü değiştirin, hepsi güncellenir.")
    wo["A2"].font = SOLUK

    sat = 4

    def blok(baslik: str, satir_listesi):
        nonlocal sat
        wo.cell(row=sat, column=1, value=baslik).font = KALIN
        wo.cell(row=sat, column=1).fill = GRI
        sat += 1
        for etiket, formul, bic in satir_listesi:
            wo.cell(row=sat, column=1, value=etiket)
            c = wo.cell(row=sat, column=2, value=formul)
            if bic:
                c.number_format = bic
            sat += 1
        sat += 1

    blok("1. SEZONA NE KADAR STOKLA GİRDİK (31.08 snapshot)", [
        ("2025 giriş stoğu (adet)",
         f'=SUMIF({rng("Giriş stoğu 31.08.2025")},">0")', "#,##0"),
        ("2026 giriş stoğu (adet)",
         f'=SUMIF({rng("Giriş stoğu 31.08.2026")},">0")', "#,##0"),
        ("Değişim %", "=IFERROR(B6/B5-1,\"\")", "0.0%"),
        ("2025 stoklu çeşit-mağaza",
         f'=COUNTIF({rng("Giriş stoğu 31.08.2025")},">0")', "#,##0"),
        ("2026 stoklu çeşit-mağaza",
         f'=COUNTIF({rng("Giriş stoğu 31.08.2026")},">0")', "#,##0"),
        ("2025 çeşit başına derinlik", "=IFERROR(B5/B8,\"\")", "0.0"),
        ("2026 çeşit başına derinlik", "=IFERROR(B6/B9,\"\")", "0.0"),
        ("Derinlik değişimi %", "=IFERROR(B11/B10-1,\"\")", "0.0%"),
    ])

    blok("2. OKULA HİZALI SATIŞ (ön-sezon T−45..T−1)", [
        ("Geçen yıl adet", f'=SUM({rng("GY ön-sezon adet")})', "#,##0"),
        ("Bu yıl adet", f'=SUM({rng("BY ön-sezon adet")})', "#,##0"),
        ("Adet değişimi %", "=IFERROR(B16/B15-1,\"\")", "0.0%"),
        ("Bu yıl ₺", f'=SUM({rng("BY ön-sezon ₺")})', "#,##0"),
        ("Geçen yıl sezon adedi (T+0..T+30) — beklentinin tabanı",
         f'=SUM({rng("GY sezon adet")})', "#,##0"),
    ])

    blok("3. BUGÜN — STOK vs BEKLENEN", [
        ("Bugün mağaza stoğu (adet)",
         f'=SUMIF({rng("Bugün mağaza stoğu")},">0")', "#,##0"),
        ("Beklenen kalan sezon (adet)",
         f'=SUM({rng("Beklenen kalan sezon")})', "#,##0"),
        ("Mağaza kapağı (kat)", "=IFERROR(B22/B23,\"\")", "0.00"),
        ("Merkez depo RAF+GİRİŞ (adet, ürün bazında tekil)",
         f'=SUMPRODUCT(({magaza}="{MEKAN_AD[1]}")*{rng("Depoda toplam")})', "#,##0"),
        ("Mağaza + depo toplam kapağı", "=IFERROR((B22+B25)/B23,\"\")", "0.00"),
    ])

    blok("4. AÇIK (yalnız CANLI ürünler: depoda stok var ya da bu yıl satmış)", [
        ("Açık olan ürün-mağaza sayısı",
         f'=COUNTIFS({canli},"CANLI",{acik},">0")', "#,##0"),
        ("Açık adet",
         f'=SUMIFS({acik},{canli},"CANLI")', "#,##0"),
        ("Açık maliyeti ₺",
         f'=SUMIFS({rng("Açık maliyeti ₺")},{canli},"CANLI")', "#,##0"),
        ("Açık satış kaybı ₺ (liste fiyatı)",
         f'=SUMIFS({rng("Açık satış kaybı ₺")},{canli},"CANLI")', "#,##0"),
        ("Depodan karşılanabilir adet",
         f'=SUMIFS({rng("Depodan karşılanır")},{canli},"CANLI")', "#,##0"),
        ("Satın alma gereken adet",
         f'=SUMIFS({rng("Satın alma gereken")},{canli},"CANLI")', "#,##0"),
        ("— TRANSFERLE KAPANIR (satır)",
         f'=COUNTIFS({canli},"CANLI",{aksiyon},"TRANSFERLE KAPANIR")', "#,##0"),
        ("— KISMİ TRANSFER (satır)",
         f'=COUNTIFS({canli},"CANLI",{aksiyon},"KISMİ TRANSFER")', "#,##0"),
        ("— SATIN ALMA GEREKİR (satır)",
         f'=COUNTIFS({canli},"CANLI",{aksiyon},"SATIN ALMA GEREKİR")', "#,##0"),
        ("Bırakılmış görünen ürün-mağaza (açığı var ama canlı değil)",
         f'=COUNTIFS({canli},"BIRAKILMIS?",{acik},">0")', "#,##0"),
    ])

    blok("5. AÇIĞIN TERSİ — ZİRVEDE HİÇ SATMAYAN STOK", [
        ("Satmayan ürün-mağaza sayısı",
         f'=COUNTIFS({rng("Zirvede sattı mı")},"SATMADI",{rng("Bugün mağaza stoğu")},">0")', "#,##0"),
        ("Satmayan stok adedi",
         f'=SUMIFS({rng("Bugün mağaza stoğu")},{rng("Zirvede sattı mı")},"SATMADI",'
         f'{rng("Bugün mağaza stoğu")},">0")', "#,##0"),
        ("Satmayan stok maliyeti ₺",
         f'=SUM({rng("Atıl stok ₺")})', "#,##0"),
        ("Mağaza stoğunun yüzdesi", "=IFERROR(B42/B22,\"\")", "0.0%"),
    ])

    wo.column_dimensions["A"].width = 58
    wo.column_dimensions["B"].width = 18
    wo["A2"].alignment = Alignment(wrap_text=True)

    # ----------------------------------------------------------------- MAĞAZA ÖZET
    wm = wb.create_sheet("Mağaza Özet", 2)
    basliklar = ["Mağaza", "Giriş 2025", "Giriş 2026", "Giriş değişimi %",
                 "GY ön-sezon", "BY ön-sezon", "Ön-sezon değişimi %",
                 "GY sezon (taban)", "Beklenen kalan", "Bugün stok", "Kapak",
                 "Açık adet (canlı)", "Açık ₺ (canlı)", "Depodan karşılanır",
                 "Satın alma gereken", "Satmayan stok", "Atıl ₺"]
    wm.append(basliklar)
    for i in range(1, len(basliklar) + 1):
        c = wm.cell(row=1, column=i)
        c.fill, c.font = KIRMIZI, BASLIK_YAZI
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wm.row_dimensions[1].height = 34

    for j, ad in enumerate(("FSM", "Özlüce", "İst.Yolu"), start=2):
        m = f'{magaza},"{ad}"'
        wm.cell(row=j, column=1, value=ad)
        wm.cell(row=j, column=2, value=f'=SUMIFS({rng("Giriş stoğu 31.08.2025")},{m},{rng("Giriş stoğu 31.08.2025")},">0")')
        wm.cell(row=j, column=3, value=f'=SUMIFS({rng("Giriş stoğu 31.08.2026")},{m},{rng("Giriş stoğu 31.08.2026")},">0")')
        wm.cell(row=j, column=4, value=f"=IFERROR(C{j}/B{j}-1,\"\")")
        wm.cell(row=j, column=5, value=f'=SUMIFS({rng("GY ön-sezon adet")},{m})')
        wm.cell(row=j, column=6, value=f'=SUMIFS({rng("BY ön-sezon adet")},{m})')
        wm.cell(row=j, column=7, value=f"=IFERROR(F{j}/E{j}-1,\"\")")
        wm.cell(row=j, column=8, value=f'=SUMIFS({rng("GY sezon adet")},{m})')
        wm.cell(row=j, column=9, value=f"=H{j}*Buyume")
        wm.cell(row=j, column=10, value=f'=SUMIFS({rng("Bugün mağaza stoğu")},{m},{rng("Bugün mağaza stoğu")},">0")')
        wm.cell(row=j, column=11, value=f"=IFERROR(J{j}/I{j},\"\")")
        wm.cell(row=j, column=12, value=f'=SUMIFS({acik},{m},{canli},"CANLI")')
        wm.cell(row=j, column=13, value=f'=SUMIFS({rng("Açık maliyeti ₺")},{m},{canli},"CANLI")')
        wm.cell(row=j, column=14, value=f'=SUMIFS({rng("Depodan karşılanır")},{m},{canli},"CANLI")')
        wm.cell(row=j, column=15, value=f'=SUMIFS({rng("Satın alma gereken")},{m},{canli},"CANLI")')
        wm.cell(row=j, column=16, value=f'=SUMIFS({rng("Bugün mağaza stoğu")},{m},{rng("Zirvede sattı mı")},"SATMADI",{rng("Bugün mağaza stoğu")},">0")')
        wm.cell(row=j, column=17, value=f'=SUMIFS({rng("Atıl stok ₺")},{m})')

    wm.cell(row=5, column=1, value="TOPLAM").font = KALIN
    for i in range(2, len(basliklar) + 1):
        h = get_column_letter(i)
        if i in (4, 7, 11):
            continue
        wm.cell(row=5, column=i, value=f"=SUM({h}2:{h}4)").font = KALIN
    wm.cell(row=5, column=4, value="=IFERROR(C5/B5-1,\"\")").font = KALIN
    wm.cell(row=5, column=7, value="=IFERROR(F5/E5-1,\"\")").font = KALIN
    wm.cell(row=5, column=11, value="=IFERROR(J5/I5,\"\")").font = KALIN

    for i in range(2, len(basliklar) + 1):
        h = get_column_letter(i)
        bic = "0.0%" if i in (4, 7) else ("0.00" if i == 11 else "#,##0")
        for rr in range(2, 6):
            wm.cell(row=rr, column=i).number_format = bic
        wm.column_dimensions[h].width = 15
    wm.column_dimensions["A"].width = 12

    # ------------------------------------------------- KIRILIM (alt grup / marka / tedarikçi)
    for sayfa_ad, kolon_ad in (("Alt Grup Kırılımı", "Alt grup"),
                               ("Marka Kırılımı", "Marka"),
                               ("Tedarikçi Kırılımı", "Tedarikçi")):
        degerler = sorted({getattr(r, {"Alt grup": "AltGrup", "Marka": "Marka",
                                       "Tedarikçi": "Tedarikci"}[kolon_ad]) for r in satirlar})
        wk = wb.create_sheet(sayfa_ad)
        bas = [kolon_ad, "GY sezon adet", "Beklenen kalan", "Bugün stok", "Kapak",
               "Açık adet (canlı)", "Açık ₺ (canlı)", "Depodan karşılanır",
               "Satın alma gereken", "Satmayan stok", "Atıl ₺"]
        wk.append(bas)
        for i in range(1, len(bas) + 1):
            c = wk.cell(row=1, column=i)
            c.fill, c.font = KIRMIZI, BASLIK_YAZI
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wk.row_dimensions[1].height = 30
        kr = rng(kolon_ad)
        for j, d in enumerate(degerler, start=2):
            # Excel formülünde çift tırnak kaçışı
            g = str(d).replace('"', '""')
            f = f'{kr},"{g}"'
            wk.cell(row=j, column=1, value=d)
            wk.cell(row=j, column=2, value=f'=SUMIFS({rng("GY sezon adet")},{f})')
            wk.cell(row=j, column=3, value=f"=B{j}*Buyume")
            wk.cell(row=j, column=4, value=f'=SUMIFS({rng("Bugün mağaza stoğu")},{f},{rng("Bugün mağaza stoğu")},">0")')
            wk.cell(row=j, column=5, value=f"=IFERROR(D{j}/C{j},\"\")")
            wk.cell(row=j, column=6, value=f'=SUMIFS({acik},{f},{canli},"CANLI")')
            wk.cell(row=j, column=7, value=f'=SUMIFS({rng("Açık maliyeti ₺")},{f},{canli},"CANLI")')
            wk.cell(row=j, column=8, value=f'=SUMIFS({rng("Depodan karşılanır")},{f},{canli},"CANLI")')
            wk.cell(row=j, column=9, value=f'=SUMIFS({rng("Satın alma gereken")},{f},{canli},"CANLI")')
            wk.cell(row=j, column=10, value=f'=SUMIFS({rng("Bugün mağaza stoğu")},{f},{rng("Zirvede sattı mı")},"SATMADI",{rng("Bugün mağaza stoğu")},">0")')
            wk.cell(row=j, column=11, value=f'=SUMIFS({rng("Atıl stok ₺")},{f})')
        for i in range(2, len(bas) + 1):
            h = get_column_letter(i)
            bic = "0.00" if i == 5 else "#,##0"
            for rr in range(2, len(degerler) + 2):
                wk.cell(row=rr, column=i).number_format = bic
            wk.column_dimensions[h].width = 15
        wk.column_dimensions["A"].width = 32
        wk.freeze_panes = "B2"
        wk.auto_filter.ref = f"A1:{get_column_letter(len(bas))}{len(degerler) + 1}"

    # ----------------------------------------------------------------- KILAVUZ
    wg = wb.create_sheet("Okuma Kılavuzu")
    wg["A1"] = "KOLON SÖZLÜĞÜ VE SINIRLAR"
    wg["A1"].font = Font(bold=True, size=13)
    wg["A3"], wg["B3"], wg["C3"] = "Kolon", "Tip", "Tanım"
    for h in ("A3", "B3", "C3"):
        wg[h].fill, wg[h].font = KIRMIZI, BASLIK_YAZI

    tanim = {
        "Mağaza": "1 FSM · 4477 Özlüce · 4478 İst.Yolu. Merkez depo ayrı kolonlarda.",
        "stkID": "DerinSIS ürün anahtarı. Ürün eşleşmesi HER ZAMAN stkID üstünden (stkKod barkod DEĞİL).",
        "Stok kodu": "urn.stkKod — BARKOD DEĞİL, muhasebe kodu. Eşleştirmede kullanılmaz.",
        "GY ön-sezon adet": "Geçen yıl, okuldan önceki 45 günde bu mağazada satılan net adet.",
        "GY sezon adet": "Geçen yıl, okul günü + sonraki 30 günde satılan net adet. BEKLENTİNİN TABANI.",
        "GY sezon ₺": "Aynı pencerenin net cirosu (KDV hariç).",
        "BY ön-sezon adet": "Bu yıl, okuldan önceki 45 günde satılan net adet.",
        "BY ön-sezon TÜM ŞUBE": "Aynı ürünün ÜÇ ŞUBE TOPLAMINDA bu yıl ön-sezon adedi. "
            "'Canlı mı' testi bunu kullanır: ürün başka şubede satıyorsa hâlâ koleksiyondadır. "
            "Şube bazlı bakmak açığı sessizce küçültür (ölçüldü: 1.584 → 1.318 satır).",
        "BY ön-sezon ₺": "Aynı pencerenin net cirosu (KDV hariç).",
        "Son satış": "Bu yıl ön-sezonda bu mağazada en son satış tarihi. Boş = bu pencerede hiç satmadı.",
        "Giriş stoğu 31.08.2025": "Geçen yıl sezona girerken ay-sonu bakiyesi (snapshot).",
        "Giriş stoğu 31.08.2026": "Bu yıl sezona girerken ay-sonu bakiyesi (snapshot).",
        "Bugün mağaza stoğu": "Kesim anı mağaza stoğu (dbo.stokSonAltDepo_vw). Negatif olabilir — defter hatası.",
        "Depo RAF": "Merkez depo raf alanı (WMS). ÜRÜN BAZINDA — üç mağaza satırında AYNI değer tekrar eder, toplamayın.",
        "Depo GİRİŞ": "Merkez depo giriş alanı (WMS), mal kabul edilmiş henüz raflanmamış.",
        "Depo ÇIKIŞ": "Merkez depo çıkış alanı — SEVKE HAZIRLANMIŞ mal. Karşılanabilirlik hesabına KATILMAZ.",
        "Birim maliyet ₺": "bkm.UrunBilgi.SonAlis — son alış birim maliyeti. 0 = hiç alış faturası yok.",
        "Satış fiyatı ₺": "bkm.UrunBilgi.SatisFiyat — liste/etiket fiyatı (KDV dahil).",
        "Beklenen kalan sezon": "= GY sezon adet × Buyume. Parametreler!B4'e bağlı.",
        "Açık adet": "= MAX(0, Beklenen − Bugün stok). Bugünkü stokla kapatılamayan kısım.",
        "Kapak (kat)": "= Bugün stok ÷ Beklenen. 1'in altı = sezonu kendi stoğuyla çıkarmıyor.",
        "Depoda toplam": "= Depo RAF + Depo GİRİŞ (ÇIKIŞ hariç).",
        "Depodan karşılanır": "= MIN(Açık, Depoda toplam). Bugün transferle kapanabilecek kısım.",
        "Satın alma gereken": "= Açık − Depodan karşılanır. Depo da boş, sipariş gerekiyor.",
        "Açık maliyeti ₺": "= Açık × Birim maliyet. Kaçan malın maliyet değeri.",
        "Açık satış kaybı ₺": "= Açık × Satış fiyatı. Kaçan cironun liste-fiyatı karşılığı (ÜST sınır).",
        "Aksiyon": "YETERLİ / TRANSFERLE KAPANIR / KISMİ TRANSFER / SATIN ALMA GEREKİR.",
        "Canlı mı": "CANLI = depoda stoğu var ya da bu yıl satmış. BIRAKILMIS? = ikisi de yok, muhtemelen "
                    "geçen yılın tasarımı — açığı gerçek açık sayılmaz.",
        "Zirvede sattı mı": "SATMADI = bu yıl ön-sezonun 45 gününde hiç satmamış. Atıl stok göstergesi.",
        "Kendi büyümesi": "= BY ön-sezon ÷ GY ön-sezon − 1. Ürünün kendi büyümesi; grup ortalamasıyla kıyaslayın.",
        "Giriş stoğu değişimi": "= 2026 giriş − 2025 giriş. Negatif = bu yıl daha ince girdik.",
        "Atıl stok ₺": "Bu yıl hiç satmamış ürünün rafta duran stoğunun maliyeti.",
    }
    r = 4
    for b, tip, *_ in KOLONLAR:
        wg.cell(row=r, column=1, value=b)
        wg.cell(row=r, column=2, value="ÖLÇÜM" if tip == "ham" else "FORMÜL")
        wg.cell(row=r, column=2).font = KALIN if tip == "formul" else Font()
        wg.cell(row=r, column=3, value=tanim.get(b, ""))
        wg.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    r += 1
    wg.cell(row=r, column=1, value="SINIRLAR — karar verirken okunması gerekenler").font = Font(bold=True, size=12)
    r += 1
    for s in [
        "Kayıp satış ALT SINIRDIR: stok bitince satış kesilir, o yüzden geçen yılın adedi "
        "gerçek talebin altındadır (sağdan sansürlü ölçüm). Literatürdeki düzeltme EM tabanlı "
        "sansürlü-talep tahminidir; uygulanmadı.",
        "Tek yıldan ürün bazında tahmin GÜRÜLTÜLÜDÜR — defter tasarımı yıllık döner. "
        "'Açık adet' bir ÖNCELİK SIRASIDIR, sipariş emri değil.",
        "Depo kolonları ÜRÜN bazındadır, mağaza bazında değil. Üç mağaza aynı depo stoğunu "
        "paylaşır; 'Depodan karşılanır' kolonları toplanınca depo stoğu birden fazla kez sayılabilir.",
        "Merkez deponun GEÇEN YIL ne taşıdığı ÖLÇÜLEMEDİ — WMS ay-sonu snapshot'ı yalnız "
        "31.08.2026 için var. Depo tarafında yıl karşılaştırması yapılamaz.",
        "Kasa satışı ERP'ye saatte bir akar; kesim günü verisi EKSİKTİR. Bu yüzden bu yılın "
        "sezon (okul günü sonrası) satışı hiç kıyaslanmadı — yalnız ön-sezon kıyaslandı.",
        "Birim maliyet bugünün son alış fiyatıdır; iki yıla da aynı fiyat uygulanınca ₺ "
        "karşılaştırması enflasyondan arınır ama 'gerçek o yılın maliyeti' değildir.",
        "Kapsam seçimi: defterde dokuz mekan var; bu dosya üç mağaza + merkez depoyu kapsar. "
        "İade Deposu (4480) gibi mekanlar HARİÇTİR.",
        "Negatif mağaza stoğu satırları ERP defteri hatasıdır (fiziksel olarak imkânsız); "
        "toplamlarda '>0' süzgeci kullanıldı, ama satır bazında görünür bırakıldı.",
    ]:
        wg.cell(row=r, column=1, value="⚠")
        c = wg.cell(row=r, column=3, value=s)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1

    wg.column_dimensions["A"].width = 26
    wg.column_dimensions["B"].width = 10
    wg.column_dimensions["C"].width = 118

    wb.save(cikti)
    print(f"Yazildi: {cikti}")
    print(f"  urun-magaza satiri : {len(satirlar)}")
    print(f"  ham kolon          : {sum(1 for _, t, *_ in KOLONLAR if t == 'ham')}")
    print(f"  formul kolon       : {sum(1 for _, t, *_ in KOLONLAR if t == 'formul')}")
    print("  Ozet/Magaza/Kirilim sayfalarindaki TUM rakamlar formuldur (sabit sayi yok).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
