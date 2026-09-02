# -*- coding: utf-8 -*-
"""Sunum yerlesim denetcisi — GOZLE DEGIL OLCEREK. Cakisan/tasan kutulari bulur.

Neden: gorsel QA'de slayt slayt bakmak kaciriyor (02.09.2026: KPI etiket-deger binmesi ve
kart ust bandi cakismasi gozden kacti). Bu script her slaytta metin tasiyan kutularin
sinirlarini okur, ikili cakismalari ve yasak bolge ihlallerini raporlar.

Yasak bolgeler (bkm-sunum yerlesim guvenli bolge kurali):
  baslik       y < 1.45      (baslik placeholder alani — icerik girmez)
  logo         x 0.20-2.05 · y 6.15-7.05
  alt bant     y > 7.30
Cakisma esigi: iki kutunun kesisim alani, kucuk kutunun %12'sinden fazlaysa RAPORLANIR
(kart + kart-ici metin gibi kasitli ic-ice yerlesimler haric — kapsayan kutu metinsizse atlanir).

Kullanim: python scripts/sunum_yerlesim_denetle.py [dosya.pptx]
Cikis kodu: ihlal varsa 1 (CI/otomasyon icin).
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

KOK = Path(__file__).resolve().parent.parent
VARSAYILAN = KOK / "briefings" / "sezon-kadro-20260902" / "sunum-kadro-sezon2026.pptx"

BASLIK_ALT = 1.45
LOGO = (0.20, 2.05, 6.15, 7.05)     # x1, x2, y1, y2
ALT_BANT = 7.30
KESISIM_ESIGI = 0.12


def inc(v):
    return Emu(v).inches if v is not None else 0.0


def kutu(sh):
    return (inc(sh.left), inc(sh.top), inc(sh.left) + inc(sh.width), inc(sh.top) + inc(sh.height))


def metin(sh):
    try:
        if sh.has_text_frame:
            return " ".join(p.text for p in sh.text_frame.paragraphs).strip()
    except Exception:
        pass
    return ""


def kesisim(a, b):
    x = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    y = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    return x * y


def alan(a):
    return max(0.0, a[2] - a[0]) * max(0.0, a[3] - a[1])


def denetle(yol):
    pr = Presentation(str(yol))
    ihlal = []
    for i, sl in enumerate(pr.slides, start=1):
        sekiller = []
        for sh in sl.shapes:
            k = kutu(sh)
            if alan(k) <= 0:
                continue
            m = metin(sh)
            tur = ("tablo" if getattr(sh, "has_table", False) else
                   "grafik" if getattr(sh, "has_chart", False) else
                   "metin" if m else "sekil")
            sekiller.append({"ad": sh.shape_type, "kutu": k, "metin": m, "tur": tur,
                             "ph": sh.is_placeholder})

        # 1) yasak bolge ihlalleri — yalniz METIN/TABLO/GRAFIK icin (dekoratif sekil serbest)
        for s_ in sekiller:
            if s_["tur"] == "sekil":
                continue
            x1, y1, x2, y2 = s_["kutu"]
            ozet = (s_["metin"][:42] + "…") if len(s_["metin"]) > 42 else s_["metin"]
            if y1 < BASLIK_ALT and not s_["ph"]:
                ihlal.append((i, "BASLIK ALANI", "y1=%.2f < %.2f" % (y1, BASLIK_ALT), ozet))
            if y2 > ALT_BANT:
                ihlal.append((i, "ALT BANT", "y2=%.2f > %.2f" % (y2, ALT_BANT), ozet))
            if (x1 < LOGO[1] and x2 > LOGO[0] and y1 < LOGO[3] and y2 > LOGO[2]):
                ihlal.append((i, "LOGO BOLGESI", "kutu=(%.2f,%.2f)-(%.2f,%.2f)" % (x1, y1, x2, y2), ozet))

        # 2) metin tasiyan kutular arasi cakisma
        metinli = [s_ for s_ in sekiller if s_["tur"] in ("metin", "tablo", "grafik")]
        for a in range(len(metinli)):
            for b in range(a + 1, len(metinli)):
                ka, kb = metinli[a]["kutu"], metinli[b]["kutu"]
                ks = kesisim(ka, kb)
                if ks <= 0:
                    continue
                kucuk = min(alan(ka), alan(kb))
                if kucuk <= 0 or ks / kucuk < KESISIM_ESIGI:
                    continue
                ihlal.append((i, "CAKISMA",
                              "%%%.0f kesisim" % (100 * ks / kucuk),
                              "«%s» ↔ «%s»" % (metinli[a]["metin"][:26], metinli[b]["metin"][:26])))
    return pr, ihlal


def main(argv):
    yol = Path(argv[1]) if len(argv) > 1 else VARSAYILAN
    pr, ihlal = denetle(yol)
    n = len(pr.slides._sldIdLst)
    print("Denetlenen: %s (%d slayt)" % (yol.name, n))
    if not ihlal:
        print("✓ İhlal yok — çakışma, başlık/logo/alt bant taşması bulunamadı.")
        return 0
    print("⚠ %d ihlal:" % len(ihlal))
    for slayt, tip, detay, ozet in ihlal:
        print("  Slayt %-2d %-14s %-22s %s" % (slayt, tip, detay, ozet))
    return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
