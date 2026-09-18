"""İŞTEN ÇIKMIŞ PERSONEL LİSTESİ — PDKS kapsamındaki hayalet satırların kaynağı.

ÖLÇÜT (İK uygulaması, kullanıcı bildirdi 18.09.2026):
    çıkışta kart numarası personel numarasıyla AYNI yapılır + takip kapatılır.
        Per_AuswNr = CONVERT(varchar, Per_PersNr)  AND  Per_ZeitAktiv = 0

⚠ `Per_ZeitAktiv` TEK BAŞINA ÖLÇÜT DEĞİLDİR — o alan "zaman takibi açık mı"
  demektir, "çalışıyor mu" demez. Kart basmayan AKTİF personel de 0 taşır.
  Gecoweb "Personel Bilgileri" ekranında İŞTEN ÇIKIŞ TARİHİ ALANI YOKTUR;
  oradaki "ZT A.Bitiş Tarihi" takibin bittiği tarihtir, istihdamın değil.

İki grup ayrı sayfada:
  A · dönem ÖNCESİ çıkmış  — dönemde hiç kart okutması yok, satırların tamamı hayalet
  B · dönem İÇİNDE çıkmış  — okutmaları son iş gününde kesiliyor, o güne kadarı GERÇEK

Kullanım:
    python vardiya/zei32_cikmis_personel.py --bas 31.08.2026 --bit 30.09.2026
"""

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import pyodbc
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8")
        except Exception:
            pass

REPO = Path(__file__).resolve().parent.parent
PANEL = (r"Driver={ODBC Driver 18 for SQL Server};Server=BT-FIKRI\SQLEXPRESS;"
         r"Database=BkmPanel;Trusted_Connection=yes;TrustServerCertificate=yes;"
         r"Login Timeout=10")

BASLIK_A = ["PersNr", "Ad Soyad", "Şube", "Bölüm", "Görev", "Kart No",
            "ZT Başlama", "ZT A.Bitiş", "Rapor satırı", "DESIZ satırı"]
BASLIK_B = ["PersNr", "Ad Soyad", "Şube", "Bölüm", "Görev", "Kart No",
            "İlk okutma", "Son okutma", "Çalışılan gün", "Rapor satırı",
            "Son okutmadan sonraki hayalet satır"]


def cift_atlama(cn, ic_sql: str):
    """LIVE201 → [PDKS] çift atlama. Tırnak iki kez katlanır."""
    l1 = "SELECT * FROM OPENQUERY([PDKS], '" + ic_sql.replace("'", "''") + "') x"
    sql = "SELECT * FROM OPENQUERY(LIVE201, '" + l1.replace("'", "''") + "') y"
    cur = cn.cursor()
    cur.execute(sql)
    kol = [d[0] for d in cur.description]
    sat = cur.fetchall()
    cur.close()
    return kol, sat


def tarih_oku(metin: str) -> dt.date:
    m = re.fullmatch(r"(\d{2})\.(\d{2})\.(\d{4})", metin.strip())
    if not m:
        sys.exit(f"Tarih dd.MM.yyyy olmalı: {metin!r}")
    return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))


def veri_cek(cn, bas: dt.date, bit: dt.date):
    b8, t8 = bas.strftime("%Y%m%d"), bit.strftime("%Y%m%d")
    ic = f"""
        SELECT  p.Per_PersNr,
                AdSoyad = LTRIM(RTRIM(p.Per_Vorname)) + ' ' + LTRIM(RTRIM(p.Per_Name)),
                Sube   = LTRIM(RTRIM(ISNULL(p.Per_Grp2, ''))),
                Bolum  = LTRIM(RTRIM(ISNULL(p.Per_Grp3, ''))),
                Gorev  = LTRIM(RTRIM(ISNULL(p.Per_Grp4, ''))),
                KartNo = LTRIM(RTRIM(ISNULL(p.Per_AuswNr, ''))),
                p.Per_ZeitAktivVonDatum, p.Per_ZeitAktivBisDatum,
                ilkOk = (SELECT MIN(l.TLe_Datum) FROM TTagLes l
                         WHERE l.TLe_PersNr = p.Per_PersNr
                           AND l.TLe_Datum >= '{b8}' AND l.TLe_Datum <= '{t8}'
                           AND l.TLe_VonZeit IS NOT NULL),
                sonOk = (SELECT MAX(l.TLe_Datum) FROM TTagLes l
                         WHERE l.TLe_PersNr = p.Per_PersNr
                           AND l.TLe_Datum >= '{b8}' AND l.TLe_Datum <= '{t8}'
                           AND l.TLe_VonZeit IS NOT NULL),
                gun   = (SELECT COUNT(DISTINCT l.TLe_Datum) FROM TTagLes l
                         WHERE l.TLe_PersNr = p.Per_PersNr
                           AND l.TLe_Datum >= '{b8}' AND l.TLe_Datum <= '{t8}'
                           AND l.TLe_VonZeit IS NOT NULL),
                satir = (SELECT COUNT(*) FROM TTagMoS s
                         WHERE s.TMS_PersNr = p.Per_PersNr
                           AND s.TMS_Datum >= '{b8}' AND s.TMS_Datum <= '{t8}'),
                desiz = (SELECT COUNT(*) FROM TTagMoS s
                         WHERE s.TMS_PersNr = p.Per_PersNr
                           AND s.TMS_Datum >= '{b8}' AND s.TMS_Datum <= '{t8}'
                           AND s.TMS_LetzteAbwArt = 'DESIZ'),
                hayalet = (SELECT COUNT(*) FROM TTagMoS s
                           WHERE s.TMS_PersNr = p.Per_PersNr
                             AND s.TMS_Datum >= '{b8}' AND s.TMS_Datum <= '{t8}'
                             AND s.TMS_Datum > ISNULL(
                                   (SELECT MAX(l3.TLe_Datum) FROM TTagLes l3
                                    WHERE l3.TLe_PersNr = p.Per_PersNr
                                      AND l3.TLe_Datum >= '{b8}' AND l3.TLe_Datum <= '{t8}'
                                      AND l3.TLe_VonZeit IS NOT NULL), '17530101'))
        FROM    TPerTab p
        WHERE   p.Per_ZeitAktiv = 0
            AND LTRIM(RTRIM(ISNULL(p.Per_AuswNr, ''))) = CONVERT(varchar, p.Per_PersNr)
            AND EXISTS (SELECT 1 FROM TTagMoS s
                        WHERE s.TMS_PersNr = p.Per_PersNr
                          AND s.TMS_Datum >= '{b8}' AND s.TMS_Datum <= '{t8}')
    """
    _, sat = cift_atlama(cn, ic)
    if not sat:
        sys.exit("KOŞAMADI: PDKS boş döndü — ölçüte uyan kimse yok mu, "
                 "yoksa linked server mi düştü? (boş sonuç 'temiz' demek DEĞİL)")
    return sat


def _g(v):
    return v.date() if isinstance(v, dt.datetime) else v


def sayfa_yaz(ws, basliklar, satirlar, genisler, not_metni):
    ws.append([not_metni])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(basliklar))
    h = ws.cell(row=1, column=1)
    h.font = Font(bold=True, size=11)
    h.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 46
    ws.append(basliklar)
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for s in satirlar:
        ws.append(s)
    ws.freeze_panes = "A3"
    for i, g in enumerate(genisler, 1):
        ws.column_dimensions[get_column_letter(i)].width = g


def main() -> int:
    ap = argparse.ArgumentParser(description="İşten çıkmış personel listesi (Excel)")
    ap.add_argument("--bas", required=True, help="dd.MM.yyyy")
    ap.add_argument("--bit", required=True, help="dd.MM.yyyy")
    ap.add_argument("--cikti", default=None)
    a = ap.parse_args()
    bas, bit = tarih_oku(a.bas), tarih_oku(a.bit)

    cn = pyodbc.connect(PANEL)
    cn.timeout = 300
    try:
        ham = veri_cek(cn, bas, bit)
    finally:
        cn.close()

    once, icinde = [], []
    for r in ham:
        (pn, ad, sube, bolum, gorev, kart, ztvon, ztbis,
         ilk, son, gun, satir, desiz, hayalet) = r
        if gun == 0:
            once.append([pn, ad, sube, bolum, gorev, kart,
                         _g(ztvon), _g(ztbis), satir, desiz])
        else:
            icinde.append([pn, ad, sube, bolum, gorev, kart,
                           _g(ilk), _g(son), gun, satir, hayalet])
    once.sort(key=lambda x: (x[2], x[1]))
    icinde.sort(key=lambda x: (x[7], x[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "A · Dönem öncesi çıkmış"
    sayfa_yaz(
        ws, BASLIK_A, once,
        [9, 30, 20, 22, 26, 10, 13, 13, 13, 13],
        f"A · DÖNEM ÖNCESİ ÇIKMIŞ — {len(once)} kişi · {sum(r[8] for r in once)} rapor satırı "
        f"(bunun {sum(r[9] for r in once)}'i DESIZ). Dönemde HİÇ kart okutması yok; "
        f"satırların tamamı hayalet. Ölçüt: kart no = personel no VE zaman takibi kapalı. "
        f"Dönem {bas:%d.%m.%Y}–{bit:%d.%m.%Y}.")

    ws2 = wb.create_sheet("B · Dönem içinde çıkmış")
    sayfa_yaz(
        ws2, BASLIK_B, icinde,
        [9, 30, 20, 22, 26, 10, 13, 13, 14, 13, 30],
        f"B · DÖNEM İÇİNDE ÇIKMIŞ — {len(icinde)} kişi · {sum(r[9] for r in icinde)} rapor satırı, "
        f"bunun {sum(r[10] for r in icinde)}'i son okutmadan SONRAKİ hayalet gün. "
        f"Çalıştıkları günler GERÇEKTİR, silinmez. Son okutma tarihleri birbirinden farklı "
        f"— her biri kendi son iş günü, ölçütü doğrulayan desen budur.")

    etiket = f"{bas:%Y%m%d}_{bit:%Y%m%d}"
    yol = Path(a.cikti) if a.cikti else REPO / "vardiya" / f"zei32_cikmis_personel_{etiket}.xlsx"
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb.save(yol)

    print(f"[OK] {yol}")
    print(f"     A · dönem öncesi çıkmış : {len(once):>3} kişi · "
          f"{sum(r[8] for r in once):>4} satır · DESIZ {sum(r[9] for r in once)}")
    print(f"     B · dönem içinde çıkmış : {len(icinde):>3} kişi · "
          f"{sum(r[9] for r in icinde):>4} satır · hayalet {sum(r[10] for r in icinde)}")
    print(f"     TOPLAM                  : {len(once)+len(icinde):>3} kişi · "
          f"{sum(r[8] for r in once)+sum(r[9] for r in icinde):>4} satır")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
